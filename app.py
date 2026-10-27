# ==============================================================================
# APLIKASI ENTERPRISE GENERATOR RAB & LPJ DWP KPPBC TMC MALANG (V9.0 PRODUCTION)
# Arsitektur: Clean Native Container UI + True Relational CRUD + Mail Merge Engine
#             + Split-Alignment Accounting (Rp Rata Kiri & Angka Rata Kanan)
#             + Smart Tiered Iuran DWP + Triple-Layer Auto-History + Google Sheets
# Aturan Koding Mutlak: 0 Backslash Character di Seluruh Baris Kode Python
# ==============================================================================

import os
import io
import json
import base64
import shutil
import tempfile
import subprocess
from datetime import date, datetime

import pandas as pd
import streamlit as st
from PIL import Image

import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

try:
    import gspread
    from google.oauth2.service_account import Credentials
    HAS_GSPREAD = True
except Exception:
    HAS_GSPREAD = False


# ==============================================================================
# BAGIAN 1: KONSTANTA, ASET BAWAAN & KONFIGURASI HALAMAN V9.0
# ==============================================================================

st.set_page_config(
    page_title="DWP BC Malang — RAB & LPJ Suite V9.0",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded",
)

NL = chr(10)
TAB = chr(9)
BS = chr(92)
BULLET_CHAR = chr(8226)

DEFAULT_SPREADSHEET_ID = "1_I2f5Gb_F2hMCIIpwwjLF_LpWeCSlWc_E60rYh9hTxc"
LOCAL_HISTORY_JSON_FILE = "database_riwayat_dwp_malang.json"
LOCAL_GCP_SA_FILE = ".gcp_service_account_dwp.json"

SHEET_DB_RAB = "DATABASE_RAB"
SHEET_DB_LPJ = "DATABASE_LPJ"
SHEET_REKAP_DANA = "REKAP_ITEM_DANA"

COLS_DB_RAB = [
    "ID_RAB",
    "WAKTU_SIMPAN",
    "NAMA_KEGIATAN",
    "LOKASI_ACARA",
    "SUMBER_DANA",
    "KET_TABEL_DANA",
    "TOTAL_RAB",
    "CATATAN_NOTE",
    "STATUS_KONVERSI",
    "PAYLOAD_JSON",
]

COLS_DB_LPJ = [
    "ID_LPJ",
    "ID_RAB_REF",
    "WAKTU_SIMPAN",
    "TGL_LAPORAN",
    "NAMA_KEGIATAN",
    "TEMA_KEGIATAN",
    "KONTEKS_HUT",
    "TGL_ACARA",
    "WAKTU_ACARA",
    "TEMPAT_ACARA",
    "SUMBER_DANA",
    "KET_TABEL_DANA",
    "TOTAL_DANA",
    "PENANDATANGAN",
    "PAYLOAD_JSON",
]

COLS_REKAP_DANA = [
    "ID_DOKUMEN",
    "JENIS_DOK",
    "TGL_ACARA",
    "NAMA_KEGIATAN",
    "NO_ITEM",
    "KEBUTUHAN_KEGIATAN",
    "JUMLAH",
    "HARGA_SATUAN",
    "HARGA_TOTAL",
    "KETERANGAN",
]

TEMPLATE_LPJ_FILENAME = "LPJ_DWP GABUNGAN DI BLITAR_2026-1.docx"
TEMPLATE_RAB_FILENAME = "RAB DWP GABUNGAN DI JEMBER.docx"

EMBEDDED_LOGO1_B64 = "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDAAMCAgMCAgMDAwMEAwMEBQgFBQQEBQoHBwYIDAoMDAsKCwsNDhIQDQ4RDgsLEBYQERMUFRUVDA8XGBYUGBIUFRT/2wBDAQMEBAUEBQkFBQkUDQsNFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBT/wAARCABIAEoDASIAAhEBAxEB/8QAHwAAAQUBAQEBAQEAAAAAAAAAAAECAwQFBgcICQoL/8QAtRAAAgEDAwIEAwUFBAQAAAF9AQIDAAQRBRIhMUEGE1FhByJxFDKBkaEII0KxwRVS0fAkM2JyggkKFhcYGRolJicoKSo0NTY3ODk6Q0RFRkdISUpTVFVWV1hZWmNkZWZnaGlqc3R1dnd4eXqDhIWGh4iJipKTlJWWl5iZmqKjpKWmp6ipqrKztLW2t7i5usLDxMXGx8jJytLT1NXW19jZ2uHi4+Tl5ufo6erx8vP09fb3+Pn6/8QAHwEAAwEBAQEBAQEBAQAAAAAAAAECAwQFBgcICQoL/8QAtREAAgECBAQDBAcFBAQAAQJ3AAECAxEEBSExBhJBUQdhcRMiMoEIFEKRobHBCSMzUvAVYnLRChYkNOEl8RcYGRomJygpKjU2Nzg5OkNERUZHSElKU1RVVldYWVpjZGVmZ2hpanN0dXZ3eHl6goOEhYaHiImKkpOUlZaXmJmaoqOkpaanqKmqsrO0tba3uLm6wsPExcbHyMnK0tPU1dbX2Nna4uPk5ebn6Onq8vP09fb3+Pn6/9oADAMBAAIRAxEAPwD9UQRimlgi56UN8or88f2yP2qtR8V6/J8MPh7Lcyxmc2moX1mWkkvJCdhto9h/1f8AfH8fTp/rMalRUqftKhw4vFQwkOeZ6v8AHr9v3wv8Op7jR/Cluni7Xo/keSOTFnbye8n/AC0/7Z18wH4v/tGftB6lbRaNc6rbQTgz26aKP7PtwE/dn/SP+Wn3x/y07V1Hw9/Zd8L/AAv8If8ACXfFq+Nvq1vL9og0TzI/+Wcn+rk8vzPM8zy6j8Y/tjeLNaa3i8OQQeFrKCP/AJZeXcSSf+Q//adfAYviCpWn7PAwPXyLgzPuKf3k/cpnO/8ADJnx18YahcXerat/plv/AMvOr6xJJJJ/1z/1lTx+D/2nvhhYy67Bq+u/Y7CNyY49X+2RpH/17ySSf+i6w9D8UfE/xXPqeoaZqfiPUX+/fXFnI/6eXT/D/wC0L8QfDaW8Nv4ku5LO3k/497n95/8AbK83+0s3WvPA+1Xg3dtYbEfvEesfDH/gon4h8LarHofxR8PeY0fySXllAba8Q/8ATS3f/wC1191eCPHmhfEbQLfW/DmoQatp0/3Li3fINfEEHj34YftPT22m/EHSY9A1qODy7fUorjy/3n/TOT/2nJXkmn3vj39iH4nJdW5uL3whqFx+7Ikj+z6vbx/+i5K+ly3iCGIn9XxPuTPzTM8szbhev7PH0/3Z+r4GBSba5L4beP8ASPih4S07xHoc/wBp0++iEiE8Ov8AsP6Guv3V9nc2hJVVzI+fP2xPjM3wY+Dt/PYT+X4g1WQafYSJjMbyffk/7Zx5P12V8gfscfDK50jTNY+MWoXcdtZ6Rb3EdlHc/wDLxJ5cnmSSSf8APOtH/gpb4qn1z4s+HPDcRzbaXphuPL/6aTyY/wDaUdaPx0j0T4dfBTwJ4L0MfZpruKK9v7eP935h8uP95JH/ANNJMf6z/nnXwfEmLqJwwkPtlcN5V/rHxFTw7/hwPGfiN8RNV+J3ie41zWZP3kn7uOOP/V28f/POOt/RfgnqmsfCfUPHialax2toxX7Fj55B9a9O/Zm/ZU0z4xeFp9f1jV57e28+SBLez2B/+2nFdl4n+BY/Zx8Qafrnh3VLvWdBlk8jUvD0582S4jk/dufLQfvDXz8qX1eh7SHwH9V5hxPgMBL+y8rn7OpT8v8AyT/gnsv7JVlo+l/A3w+9qLdHu0eS4k/56Tbzn9a+VfiF8Hf+FhftL+KvDfhqe1so4/8ATZN4/dp+7j8z/wAiSV9Fxfs6Pos2o3fhTxHNpui6hpdwLDTmfMdpeSGOQSx/9+64rxR8M9O8Ea/pnhK31bUv+Eg15zd674otbd5LhI/7iGP/AFfmSf8AounWxdaFNU6lP4D8vy3N1gsdXxmFrv2lS/8An+Gx8a+IdJn8P69qGlzbPtFpPJA/l/8ATOvpP4Z6xb/tO/Du5+Gfi6++zXNhHFcWeoxf8fEnl/8AXT/tn/38r3nVf2N/hbJ4ZlWLTGtpvL8wX8d1J5mf7+c18RfDjVdP8EfGjSrgXTyaZZ6n5YvI5PL/ANH8zy/MqcbQmqakvjPvcXjMDxzlGIoOl+8prsepfsJ/Ei++E3xk174Ua7NH5F5cXEcR3/u4r23/ANZ/38jT/wAhx1+jW4etfk7+1Da2vwm/aXsvEvhu3+zR/wCh69HFFH5cf2jzP3n/AH88v/yJX6lW+pxT28Uinh1DD8RX6NlGK+t4OE2fyNha7wc6mEn9h2Pzd+P3h/xB4p/bw1VNJ02fW30uXT9Q+xRyxRyS28dvbySRx+ZXsXx4/Z78W/HnxZo3iTQ/7O0u0uNMt0+zai8kdwkn7yT5/Lj96539rKxt/BP7WPgrxTe6rqHhrTtQ07y/7V0tPMljuY/MjwY/LPmR+XJHX1F4M16HVrGJoNaj1+VHjze2/lx+Z/y0+5/1zr4jiCusPj6ftNj6XhjHV8oxdfEYf+IeA/CP4BfGn4IahI+l6l4auLK6OZ9PuLq4eNz/AM9OIP3desWXhzx54v8AF2n65r+r6doVvYf8unh+V7j7X/108yMf+i69tltY76GN23AD1rPltpryyMkUgj/54munHYGo0vZfw/jPpcVnNfMKjr4jl9o+tiZf3refDGLhHj2ddlcB8TvAt1440i0h03WJtI1SzkjlilEHmW9xJ/00H/LRK6+MCxgiz1cyQD/v5Uum6Y8lxLJ5nlRf6vy464Kjni+Sh7M8yE/YT9pA+YPHfgD9oL4jadJpF34k8NWVg48qdNNknjL/AO/+7rzBf2DPiRZGOaTUvDUSRnfg3s//AMj19+2tvHNM5ljEs0D7PMo165KxpD5iR+Z/y0k/gr2IYSnTwlSeIPp8JxbmWCp/V8JyQT/uI/Nn/godqtzd+MfDFhLo8+niysJR9uzmO83mMYj+n/tSvvb4RTXVz8J/BU1wJPtEmiWTyZ67jAhP61+eH7XeveMviD8eLb4f6tfWF7HaX/2fSI7b93/x8SR+X5n/AE0/1dfqDZaPDaWVvAifJFGqL9AMV7+RU/Z4Y/GJf7Tja9SHc8E/bc+Fl34++ES6vofmL4g8M3A1Wzlg/wBbsT/WBP0f/tnXAfs8ftEHxv4YsrrxR460KXVY45Lf+xf7Pkt72ST/AL+fvP8AtnHX2ZIisqr7cV+dP7RPwY8Wfsu+N5vif8N7mW38PXb4vY4IxIlgX4MZjfjyP+eeP9X/AN+64c+yr+0KOm56laVTB11i6X/b595aVq0VugT/AJd3/wBXWoljMI/KWcC3A/gT56+Vf2efj1ZfEPTUt7XWr/Wtct/9I1e91mCCzSD/AJZ/u44/8/8APSve/D3iq3vVtIobt7b7XJJ5cdz+7kk/65+Z/rK+Oy3MauGqfU8XA9+nUpYiHtKZ0+o6a0rRiFBtjGY/n+49XpLEb98cjxyDr/t1WuLFInjKSOPMfp5n36zvEWp2nhmzju7i9is4Hnjj3zv+7y/7tB/KvqZ+zw/PUnTKNYbNMSR3fe718/8A7S3xs8IfD7QtQ0LxcupxSa3pdyLQW0f7yT935flxyf6uOT95Vb48/tOeFvhNJrehatc3EPiWTSxeae9vb+Zvkk8yOP8A65/vI/8AlpXwfb6j8T/2v/F+haBcXc+t3FnH5aSyRRxx2cf7vzLiSSOuL97mb9nD4DwsxzKnQXsqHv1D0n9gT4U3XxI+MUnjDVt91p3h/wD0j7Tcnf8AaLyT/V//ABz/AL91+oPle9cF8F/hLo/wW8A6b4Z0dN0NvGPOuZP9ZcSf8tJH969Cwa+6ow9jBQQ8twv1SjyMXAqleafBqllLb3cMdzbypseKQb0cUUVqeq1c+Ifjv/wTutNUvZNb+G93BpNz5nmf2LeD/Rv+2cnWP/c6V5f4Y8YfFP4XfHXQvFfxptNdi0nSbe4tzc/YvNtk8yOT/VyW/wC7+/5dFFeNmGEoyjztany+NpRwcuajoe4v/wAFKfhvczEf2X4i/wBHfj/Ro/n/APIlcD8cv2qtG/aT+FOt+EfBmh+JpfEUj29xHY/2eZJJI45I5P8Al3MlFFeEnKrpNnFHH168eWbOH+GH7BnxI+J+ppq3jm+k8OafIMu95J9ovZfpGf8AV/8AbT/v3X3/APCT4MeGPgv4dj0fw1pyW0R5muJBvuLh/wC/JJ1Joor7KnRhRXuI9jAYOjRXNBanoVFFFaHvH//Z"
EMBEDDED_LOGO2_B64 = "iVBORw0KGgoAAAANSUhEUgAAAFcAAABECAYAAAD0iqlZAAAABmJLR0QA/wD/AP+gvaeTAAAACXBIWXMAAA7EAAAOxAGVKw4bAAAgAElEQVR4nNW8d5xTVf7//zq3pLdJJskkmd6YwtB7laKAgmDHrqjoomtfV3cVEQu2xYK6il3Xiqigwke6MnSYYWCA6X0mmUnvyc3NPb8/9ut+XEXFXfHz+L0ej/tHknPe7/d55txzT72EUor/v+iZJ87z6bI+X3jd1XTT/3UspyLm/zqA71RPiKzCvyb3p35/kBDF2XPONtTszXzvJ9NcQxSPlhJHlBJyeqL8dfrd4S47gyi+/3noK4RfdZPplv3P3ue+9/OOpp/KN/V8WAdcbUBaVP5UmsziRXc77lze9s5TN9V88zQZ/f3fBu95hQ/d/eDvWt7fzdnH980xv/DUZa/Jzrh0YG4OcSRfIPwLD8hve0hxaa9l9LnPV40o1AbDPd0ny/vgxURTWzltB89Gce7cOapPnyNnnizd1TPXf1QxOI+vHF42tEW5aO/rz835Vl5NCgHgjJfv+/u7ttaG15/MWZAUye9S7t/Fyeer7JOdWmuD3Zq7aERllv+dFvS9L4z+uHT41c/ojOnM0ko9EqkBuPo/WnGy/G4vLhk7rjIPkhdFuZkIY+qz95EfA1KO8bRs/Pr9vnCkB0XFBlJUWjDp7eazDxubyJhMj/+FweXZxdachZ/ce33xlyslIj/d5T6tcJOEkGU32Bf7cf62UcPLMhR8As31OzyHPp6yXqfPnq9RCWjtqEUw3Au3vyvUFu9ffzI7GdmKoZFwKyKhdrQ0H8Cg4kE5WwHND9PxDKU6dfUTHC9Qr68Dfn8niossmofWFO1Y8SgyxGiHL9umZUeOmjjbtPKSA7PfIObTWX5yOnsL//PCiHOVeZd/kkqFuFh4ANmODHR2Hk87HJmsq68TuUVWZFh5bN9Vh/37a5e89Ap9+WR2tr1T/HhSpr+nIF8DCTIQIS/c+cfX7bPraPSHaaUbCP/msEl1Fqu5TKVIIxr0IstqhafPHe13BURDRqk+05gPntVhy7bdjd9KpcM2vfB88nSU/7TV3MlXfW7Y3WL7QEyGuYC3B0p5Godrd4ElCdZkkEOvZdHV0YyGhgapKPfI0ycDmwIhymXEHk5qhvY5G7Fn77fweJrx9aa1fPmZuWNTJM39qECv0tTcHdVjDBqhgUhpKOQsokEfsrNz1Dyv0Ks5DgHPAKLBAKZPO3OQuWvHc6eLwWmpuQ8QwtmWzv1kxLip5xIphkTCh2jUDTHlh0wmQavhoFYrwCt4hFPR8IzkF8XRs6n7h3Y+e/WGubbcyg9crh1qg7ERSqUfEmWhV41CX2uVGHDaPr3wzzcvPFkM298jN7O6C1YhHUNKSMJuyUM4mIJOlYX2djeMhmxQmQEtrmDqmosUk0XTX/f/1hzIHQtvXJtbkjvVkalvPbpv66HznZ+9WDUYLfUu0GMPAhcNhnQhQNcAuBhAx2IwhVkgq9R9/P07F74y74J5s+WyRIOcoZEDNUfLKypKGJYlKp1Wk5FMBsFzIpJJPxIJHwhikMsoeBnAMiJEyoHwGeh1xluu/MO60h8Gd24VGXLZH+/eYTS6dSp1PSOkWtDrDGJQ0RnobChMGISqv5x55+3P/jCffA8p/qD54sZcu47Ewh7EohGE/TEwRAkhxoHjdFDJMqA22sHqjAgKSWox67qO1tWxJqMmyDOxXjBKm8vpdgzyvDVj0cU4cWI58GAlpAcfAgUASoHXAOamSpD7jwDxYeDzp2BqPzvuRk1mcYE/LGWRcx6E6hLbtK/KysZNpWIEkUBXioi+YDTcK8ZDPmjUgg9MIpISKSinZiKCPlMiaoVSbeDSaWumVj0IHKMneo2WtrW1EIYFjYRCcGRbiZwH9HolkkKQimIEHJ+GSsECJIVYNEiCkRSNRNTQ6MrTsy95iD/Zvz/yRqJfosW5TH7um7PnVDKtHU3wuGIdG8Y7z1hdSLtOlueZC8moBQ+s2n/syB4UF9kISyXqHvDCYc+HxxUmRpMDvoEw5dUGEhJFGo5FEYtHyEB/H+wOM1UpeUBKobOjBUqeiasUNJgWBSrnJR9PhKgkpcESFr6QlOULhWVVQ0YQT0CQiVAbho+ZQo41dkpp0rmE+3IZjV3qJvNbnm5rnzJ+VEZpiZ2PBNKZMXUcpgo7CIQsMZ1CWqKQiBxpwiORAhJJCiElIujtRaezFUa9AXKep3q9HkFfD0JBUJMpA9F4CslEFCkxDoWSBcszoABC8RQNBATwjA0kkf76p26tQ6/Q4GwZObbolSLi9rJgGAtYrr39p8ACwIZqnBjS0puOhNNsw4keqlLyoOk0vL4IIokkTbkH4PWGkfZ5aUJMg5HxMOhUVJdnRTQWQLezD1qNHIW5BijlrFKp4JVqpRw8x9jSKQHJZBJpIQWrnYVMoYVEFJCQIKzcQFubmnDi6M5373to22oOAD4w09A1t5Jhe54u3C8ETdZUnJBoUITdAljMZhBGAggDwsjAMDLIKA9JFCEKQDIWgpIXkWm2EIVcBZ1ORQMBNQwGLTFkaBEOB5FMixBESqlACSdyYFgGlFGD4RVQKczU7j3wsw8VpQpk+vRrUVu3HjpTPnKL8n70IPu+Nrto9M0nztmQW1R1roxLg2FS0KpV1O8PEoVGjVgsieyibDCEw/FjzZTnCWHSKWh1KjBSGESnoJIUI0QEYskEaIoBK8ohyTgwkAAxBSqmkBYZtHR1QUjJqUpvha8vgKYT+z/664stiwDgX0G+5aDdu917Jte0/K2mMNeiycyUAzSKluZesCwgSRKSQhoSIwOnUINhFBBBQKQ0srPNJEPH0fr6WqTSBOk0hdGko4lEBIQjYHkZOFCA46jEyMCwDHgVgSrNIuqND+hv31SN234aViKYFf9me5O0a+8xNqdQj3Ak5Rg66ufwAub+DY+S4pFz3d5+IuMoFEo1AuEQzeB4eIJuwqlkVBQkGPRaJOMJerTmMChJQKUCBlcUIRZL0QFXJ5QygrA3AldaAEPS4BkWDEtA0xLSaRm0GcUIiAQeVz+NBPu+2Lyq5crl5J+9hH+rARP+Mb5l9wx+atDq2ZNrUcoYSmHXGaFUcSCEIhKLIy5QMDIFWE4JUWKgUWjh8wVoR2s/+vrakF9QAq0uEwxD0d7RCqVai5QoQaQSGI5FNCGAYRgIaQHpOA+LUtOQS2n850BNuuyZMY3H9zCzZl0LlT6FffuPWtY+/UjeBXff3/lTecaH0fdtKEYlyhCW5xGOxsEp5Ygkw5BrZDQY90GMUWh5EzKzHFBwDMRUGO3tx+Dpd8Gg5zFmxHBQMYx4LAghGQZD0lDKePAch3RKRDzBgZGbEItHxDxpx6NfrQg9vOcxKn0Xw7/BvbCfaM6YueAf9pxCLpnoQX9vJyAFYc5UQatRQkiJiCckSDEZxDSPpJAmx90+qtMZSNXgEbSkuAhpicXhow1IJARiNVtoMiUQjuMoJQDhWBCGA6UUVGLBcDxSjKJq00VEf9YaGvwhoPQFArvAufxPi2+Z8/CR+uNEqclBv7sNjvxhmm/3OWtfnrf0vM1fLP/mZHD9cyaMUCTkjFKhAENSaGtrRBopeDwDKBlUSsREmqplaoJ0jLY2dSEaCWLCxFGwGOXw+Xpw/MhhuJ1qGI0KkhJCNC0moFYqIKk1YBiWREJR6g2kITESzFmlSVewcmzsjD33AXj0uxj+rZ+bup9wS4Ozu+edM8Wq5H2oq/kaVitBJNSNkN8Ps1kPDhqwRAebtRiBYAR6gwpKjRbpFAeXK4ABdxRpcOBkChCWQSKVwIhRw9He2YG0SOELhGGx2CGmCCyWHHi8CdTt2/HRnS/svkz1vWCuGwwrN+qGD86df/EZPT1d0Bpl2LN3G+w5FugzDdBoMhEMxNOH6o682G5y3P3N8ofE7/I+P4Rkm66/qdZq1psyjVo0N9ZDJueQTMaRlWWBTKmAQqnGwQP7YDHoEA0GIYki1Go5jHoVVCoGajWFWgV4vV2IxLzo7+uFIIgw6M0gkEFIsVCqbYjRDBQNGoWuHh9Sri9uWXBfx0snhQsAFcFlY25b3bN75HA7w3MDcDoPwWJioJJLyNDp4XMn4HcnkGUshiCIYGRp1B8/RmIRSouLh0BvdKDXOQC3z084uYxaskxobGlCcXEhQpEocnOK0N7Wg1AgiaLiSlCZBuG4CDYd+PSGS/92ZQ+l8cl5ZMaca+5dp1DqVVl2B8ZMGI/33/8HGDkDhZLD5DMmoqO7DXWHj2HE8Ak4fOjw0Ucff3g2jVLn5ytlVb3Jq7aNHz3M1NfViGxbJhgiIsuciZ7eLsSSCXg8HqLT62kw4EV/XztGjhgCvUqHI0cPw+3sRk6OGXabHqmUH9m5JsKzEhVSCaQpC57VIRajiEUowGVAbS5BW5cf+/fva77sgT1lBab/BXrSEdqWV+f8sS9hWDlimJ1tatoBpSyKgK8LGrkSHFFBTHIwqB0I+IMgMgFqjQZKhRGxGODxxpEUJKh0OhiMBnT3dYOTsYhEQtDqTSgsGITuLhcqykdg997DUOjNCEXTGFI1HF2dPR1dXa52W3bZ5IM1R7lF1/0BVpsdz6x6Ho88shw33nITgiEPmlobUFBcAPeAF1dffgOJhKI0EQuGbltyde2Xn68dS5JhhUaRhlpBEPA5oVPLoFXL4XT1QJ+RAbenHzK5HFQSYDIo4ezrBkcZFBTmQ6tUwOXsQMDfC5lcQFqKguUoJEmCkKKApADLGSDnjWDlJhhsFajeddBZtXrbmJk+2vt9juyyZct+BNcw9MoD8q54x8adXbMslgxereQQDAZBwMBksMCgM0KnMYJCAiUiFHI5hKSIYDACjuWRaTaDZwkZ8DgxbGgFwmE/7A4b4vEYjtbVY9z4ydi79yB8/jC6+twwmCwYGPBjy7ZqQ3nFsAJKOea6xX9AT58T/mAIH69Zg1tvvxULL12IceNGo7O7A2fNPgtbtmyDmCSkqnII/D6P3NM/kL99y2Zu5IghqD9agwMH9oJjJCSFGPwBN/rdTlgsGRDTSTAsoNHIkRIisNst0KgU6O5sR3d3G5RKDg6HCXq9CnI5C4VCBqVKA16uBifTwaDPgkKRgUicpTu2bjvwIq2d+cK3tO+HHH92buHQTGJvW2BfbbSWz1ErRMJBJDxhEQ3EIWc1kKgIf8BFQ9EIMeisyC8oA03LaEenE4mUQMzWTCiVPIKxMA7W1NBZs+eTltYe6vcmYHcUkT5XiFaOnEo6nV7a0d4Nk9FG+gf89Oab7yC+QAw2ew5WPPEUguEQHn50Of7xwT9oSWkhzFYj5s47B9POmI4v1m0hJUWl9IXnV5IP3nsDV192Ad5c/TzOP/cs6na1E7ezDQ5HBkwZGurzOokjOwtd3W3geZZazUbicnXQDL0OTJoQn8cDKZ2kBr0KLBFJMNgPa5YBopSCWqenLKsmsSiPRJKjvV3uQAl/ZPnRO6KrLqc0fTJ+Pwv3lRuJauw5pluckTFLkQqpUskADCoViUWigARKRZEoVITqdQZiMGZRMcWS7s4B6vUHYLPZSGlZCWqPHEJjcxMdN2kqurpc0BvshFIl9QcEwitN1OlPIxhLA5TDhRcsJB5viL7y6lu48MLLyM7de1BUXEqbm5sxeEglmTBhHE1RAZYsM+YvmItLLr4M+bnlpK72CI2EfDhzxhRCUmH0u9pow5GDOHv2ZLS31oNIYWIwKChoHEWF2aS3rwuJRIQm4xFSWJhNmxpOIJVIoaJ0EBy2LPh9AxhwdiAlxkAYkQRCQSpXa8HyGYjHZGBYXerRqzW3Pjtq3bZxx2nzT/H7WbiPl5F8cd7447POGqNIRFxIRPqRk2UEQwUIiThEIQaNQQ0hKaLf5cNAvx86bSZycnKRTqfR09cDmYJF8aAy+EIRnDjRDo0+C12dPpgthThc3wFe50Be0RBMmDAJoUgcd991L665djHkMhV27d2Hmto6VFZWQqlW4MMP38f4SROgVMlx8x+XYMbU2bBYHOTyy6+k2XYL/B4XBpwdmH3WJDTW18Bk4GE0cDhxbD8IiUOjBKqqSiDRJDhGRGdXG3RaJcpKB4GKFLX7DyLg96KyYhAqBhWA5SSIUgrO/n7wCh3AGBBNyKEz5KGpaQBfrftiyyOfuWeZuZND/Nlh5L0NtOO9h3VL3O7hbzis2ZAxLJJCHD63E5IYh0rBouVwMxLJFDJ0ZlRUVIBn5Ojr60UwEoZCIYMxwwCffwDd3QOwWMwIRyny8wvR0enBrNnnoMuTQjCaAiUcGpva0NHZg75+Fy66cCG0GUZk5+RBq9fBYslEXl4eBrwD0Ot18PuDeOvdN2HUZ8GoN+LokcNY+syzGFxejFEjhiArOxf9fS3Iyy+BQq2DQqaEkPSiraMVCjkwbGg59IYKnDhejxMN9TBo9RhUVoRk3AG/px/ffPMNGIYiO9cBfyQCY6YavJJBJA4ISKFvIBxVL3Zf9lNgf7HmfqeGddMX7WvV/33E0AJZzb6vkWNTIZ0KIp4IIy+/mHK8gqgUWgR9YXS0dlGWZZGX7yBagxauARdAGPQN+KkENTFmFtLnVn2Eu+6+j5QNnghdVjE2bdtDNRodGTJsJByOfGzeshlDqoZDqVRDoVDgoksuxtq1a6FQK6goCsjKsmLbtm2ort6FhRdeQbq7e2lJcT6ZOmUS9u3ZhRnTJ0Cn5uhdd9xA+rqbMG5sBT7/9C0KKUhkXALpdBhKOYOhQyoABnTA1U+c3b1QsDwspkzkZtuoXMaQo0ePQKvXIBBOwOmPoaRsLOTqfLr2s687dow/Om7PlXTg57id8mR53XIytskw43VzprYiy8wQnglDp1PB6QnRAXeQpJJpyIkCDCFUr9GC40G8AS8EUYDb46XeQALFZSPIqHFz6FcbdqGwcDiROD1du3E32V97jK58+hmseuEFUlIyiN5//wOEZXhs+J+vsWTJEmg0GjQ1NmLcxIk0KcRRe7AGAJCRacSEMZPI8uWP0BHDqkgoHMHOb7fhhuuuop+vfR/l5XnkzVefhdHAIsehoaYMFtU7NxApFUJOdib1elwky26jeq0OapmKCLE4ejo7KBVFlJUUwGg2knAkAqXehI6+AO1zx6KFOZq3Dpy94f6byI9HlP8xXAAYTwl382PlR7V6WVla6KVun4fItRYqSTJiNdup2WAi6aRAU8kkYvEQGXD3wWa3Q2J4KlIZNIZsZGSWoK0zhI6uEAnHWDp++lyMmTANq55/Houuu44ADF2zZg0efvhRIooi7r77btrY3IyNGzfC09+PO/50J9555x1MnDgRM2bMwD133kuSCYHqtDqyY8cOOmx4FdlVvZ02NxyFjE+hr6sBPZ31KC+3YUh5DsyZHGlqOAT3QBc1ZxpIIBimPMshK9NKskxmSqQ0vO5+uJ3dCMfCiCYF4igopr0DMfrn81tm+SrollPl9auXed54eEi7SpnKUymC8Ae9KCgfi3AYkHNyMBQIef1gIUKnU4PlKA7V1sDqyIXBmI1oSgl/iIfFVo5jJ5wwWgpRUjkGTe3dqD10CIRl8eX69ejt6QEATJs+Ha3t7ejr64OYTMKekwOZnENHa/s/g6EAKAOHPQdOpwt5eXkQhCTOmTsHerUMdXX7cNcdi1F7YBuy7WpQ0YuRQwvQ0XYYTY2HkWU1w+7Igc/jgxCLg4gSbBYLdColerpbEBcSkKlUkGtMaO0NQBE5cdk1T9APTxvc3Z8sjvR271XptSHkFOahu5/CORABz/Aw6QxgRBEpIY5kPAyf343svFzEU2kYMvOQkVmEtu4E5Cp7SkrLPudU+Rc8+/fXmLLKoaje+Q0OHTqIP/3pHjgcDuzfvxfvv/ceVr/6KhobG2EymeB2u3G4pgZWqxUlJSWIRuLIMBiw89tdmD5zJqLRKM46czbefusNCKk4eruaMX/uNGcq1Hxw7PhR84K+drh6jsNmVSHLosHhwwcBwqOwsBh5NgdaG5qQikWhUcuRjIdQWloMTziMYCINf4xDY8POd+5dGbvmVFn9bG/hZEoKaag0GoiiDy0trfAnDaBEC5s9F7l2O5LRCMJ+D1IpDXRGPSRK4XcHQWRRKHQUNbVHhCHZ782sv9u/d2Dx9RNHjh5vP3f+AtTVHcE33+5CPB7H8RP1mD9/Pt55511g0XX46qsvodVq8c477+Cuu+6CRqMBz8lx4MABEIAMGTqYFhUVwdnXj7KyClx19XVIJiKgUhKPPPjHtz/f3vTXj1aG3snIdFw+ZOhIpBIDiEVDKC2twMFDh+Fx+2E1ZEGt1sGU5QCoiH27W5BOUyQkiqhIYMgqQZKYf3Iv28n0q5fWQ6EIysoqSFVVFdEbDBCCx/5+21V7L9y4aSvWrvuCrP10Hb7dtRfBUIzY7DmEEo7YHTng5Wo0t7Slbr2Uu/Kiu3zVD1Eqjh0/9ZWRo8aizzWAtZ+tw6atW3HNNddAq9Vi1JjR5PP1n5OvvvoS3d3dqKmpIQaDATu2f4tYNIH33/8Qx48fx/bt26HT6XDRhReSUCiExuZW8vEnn+L4iWZs2rxdWvqHc17MUFLadf/Gq/vaD3wmSQQ+bwgHDhxCWgSZMnk6hCTF/oO1pLbmKOnpdiHg7hu4f5zrJmevK11aOgg5efkkJ68ADrv5RyvUvwncGPnnzkGaTiZ7OrsQDIZRkFeMxv/BrrP6UevzepGbn4lps8dh/BmjoDBkoLHDhzRnA68uRlNLf5c58Nl485ita76z+dS1Vzyx8asv6VtvvIaZ06ejp7OHvPn623jq8b/h0kuuQI4jF/PnzcOyh5aisCgf3T2dePGlVbj22mthtWZh+bJHkE6IGFRYihuvX4z7778Pt9+6BH09HWhtbYZMJmufcfXKXgB4jFLpelJ30dbPn3s4HJMnZ865Asm0Cr0uLxieQ8XgCpSUlcDp7sXb7+6U7ToP74FJxRUqJRFFCX6/H3r58T0AkIic2oa+U0p0c+9+0z2XTNmy/mYyP57wCz6/hxqNmTQwEMN9Y+G99l5UXn3pLCSEVgqVE3HeA0GhpHrHSNrer/PWHu578OnQ7vLZf6U137d7jELgGfKOSadBW2MzPnz3I+px+TF5/HSMGjaeXrHwajq4cgjCAT/On38e3b17J8ACap0eL/79FUycMA0txztw1RVXIx4I0Q2ff4bezib66afvY+LEkdBZdLd83592BZXue0V4cLLu46FbdzZupsp8ITM7H2kmBcKlaFaukbq8HRg9OVOvSSPBcdGAJKWpICapzqDFxh1x07IVxrtzRy5/NdyNX9ymekpw5c+P8DtybEV55z39mccTsAqCgI72Lsh4JW4dhcD696YwPEcxfGQZ+gZa0OHsxLr/+aTB2b1i8rY7V+ff9fD6h3tWnnwpJ9S+8bZLF17gDwZ9SCYTePWVN3D1lYvw2qtvYmDAg8bGRkiShOKyfIweMxK7du3C/PnzkEql0dreBkqBXGsesh12bNu8CQcP7cG582Zj7Wcf+zYGjp50lcJxFW267YFPZ3363JODQm0vzu/p7Y3FBAG79uzGeRdeAIstizw6DWRIhbWT5XikUmls3/4NzNmVf50x9cEnF10/pFmbg1/sCZwS3JVPcNKhrR897g0E4cjJgz07F61tnejq6sHNK6F48plv09F4ApFwAokYhcWcg3tuo3+6/Ca66xOCyM/Z/sfG5tDaT9e+9+QTTwKQUFdXiynTJkFIxTF1yiSsePxR6PRaMFTC4UMH8dDSB6DgZQh4PZg0YSIuvvgi5ORlY/+BA2B4BsNGjCBWWxZUGtljO5a9lfg53y/voB1n3Ui/sJiynCqVDYkEB47TQKI8HvWAaWru4b/+eiuyrNkoLipHjr0Ie3YeSK1XHnnhVLidcptbcACrP1vzkXNw1XCAsoQyLOwOB1pmlj/7cSOqcvIKUH+8lUhUiUQMdOEsHD9V2698+PpSrVaDY0ePkdWrV+OBvzyAO2+/Aw888FdUV3+L5cuXwZFtg92RhQMH92PFikcgSSmMGT0SnT2dEKUUDtcfxmtvvwkJFGqtJmpb9daPduL8lI7Vt3v6+oIoLhmKrzdVQ2/IwsHr8YLCOHzk6LETSEd7L6jIorSoAmLE9dLxW+jPVpjvdMpdsadYKlUz5KxQOHpYo81gCwoGwRsIYsiEaUO3f8MOCQQjUKtNVK5Uoa3NlfzkArhO1TYN0MAtS+7Y73K5xwwbOhpbtm5CPB7B+i/WYOz4oXji8cewZMmNJDs7m5rNWWCIHN6BEJ5+6hnU1dZBpVIgmojgvAvnYtyEcXQg6ln1mEIm/bLnf6qIeJ4Mh8jarMxMuD2HUFLEQtTPuL7YZkMokqRymRpWazZaOtztgQe33nWqdn9VV2zSB/TYzg2vfhiNJVFeOZyIaUp6nF0YPmY4s2XzDpQWDyVmcz5YosG4k+4R/2ll5+bdW1xcjEsvXYg9+3ajq7uDlJWXkgkTxuPuu+/A0gcewZVX3ITZsxbgrDPnYP/eani8bogQseHrjdi8dRP5dudODB05QvxwwqYnf41vjT5LZ7cVoLO7n1RUDIOz34NEKgWP34c+l5OMGjUKYkpMN777+nVPcvSU/7Rf3c/1vRi67pvNH3/T2NQmFZeWU0pEaPUKqtbpEAwlqUqWCYe9VH7+fmQsfZMopjbuyDwVu/fdd8duuVzWpddrwRICv99Da2v206NHavD6a6uRToHmOBRgAJw3fyb6eruQSkbBy1g6cfIETJs+jd795z/h5ddf/KrjjMOBU/I5959b+jcf5/4gk8kQDAapSq2A0ZgBk9kElmcwbNhgWr1rh3Bs3dPXL9tLd/waVqc8/P2IEHYQoM0WETy4Ayy/C9fXBdQPjp04x9rj6kdhSSH5csNXmLfgEqpWZeHlBQ/ka88beUdDG1/+ycY9s08pGKK88+iRI39bv349vJ5+qDUsFHIJCiUDCgF9Pb3Izy+G2+mjSoWOJARKeZWWyBRypJHC2Amj6TX3Xn1e17d9637JV8fBTtuLb9/aOipbLTwAABARSURBVDu9ftwuMnHNOXPnlLY2H0J762FMmz4aep2ctne0onrT7h2pl3DHExm07l5CmHX9UJ2wnFqbe8pw/7p45iKN2b5KxXkiad/Gj+604NUHN6DXVIQFglV+jy13eMmAJ0SMRjsGDx1NLbJ9t62vzn48lpCJtz7wmlHO/vLtNPuPG+RVsl0bcnKzs0NBD+QyEdHoAMswtCAY9pC0kIJOZwSReKSSSKjV+nqjPVvf6xqARES09DRuLc1ce8sjT7C/6Oua88nC65e88P5rK+/fe+vNf8jsdnYXb/jqI1x5+QKIKZ+wd9vWbe+Pwz3XFSFAV2P6uIsuumb7/uCIATcrXfz6Rsf4X9gl9KvgnnXLRvkk1bsNM6cPzxdivbSn+ygtztcMHK5e/7HwMJ6d+RpGbe0f9sHI0ZMYhpVh3ZcbkpOmzpYDWgyiD1cVX0qPnZKjH+gvImE814/9JLegbIFOo4VMpoSr1w0T8Vz6xxVffPSf2ASAG68ZvPOmGxZPPLxnO+1uO44zz5xCxFQYG778yHNbEtP1V2LkR7XMLdnF06s0+hxZY4sXWY4q+sm7z97+6ueRVafi45Tb3E0vzEmONXwwp6WlUzJkWGGx5hBWlmHNr5ryR/Kw+ag5D9npREp0dfcBooDxYyrkUsoJq1nCqg2mpf8phMc4Kg0ZTNYVFuQg06SHTquCzWpM/2nUwb3/qc37lxKuonzQ8IG+JpSVmEkqMUCS8QEYdDxkLEJPiXj9QHTGm0UVs0byMr1MrtCB5RTYf/Cbg6cKFviVD7RZf6ENrQ2HPm/v6EFaIuju6kNBQQHOX3ChuqZ/3t90KoOsrakJvv4eDB7kQDzSiUi4FY5s63gA+JgQ2a8FAQBtDYdTyUQcXR2t6O1uB5VEcu+TLvWvtSMQQpIiYeebUDa4vEDV1XEERPJjzMhi9HScwPGjBzFu3JjCRYsvGT24cjjSKRnCEQnO/hAkIk9//V71KXfDgP+gtzBm1d5rj584EVCpjcgvLEVXRzuaGo5DTIooKylDtsWC9pZjEJIDJNvOg0E/cvP0pvuXEsXro80H6nHGj444/ZIO1SQi8XiIUqShUsqQoVfGr9qP3p/LQ/DjI6rP3Fpx1UOLzU9+tReFRj0HtTwNc4Yc+TlGxCNuqGQcssxZcPa4sXf3EUSjLCoqxhKNxg5vT+O71V20+tfE/avhnkNo6NbB2ycertnXEwpFkJebC61CBYNGCyUvI0adFixNQRQ8UMji0KgTRM7HmPnLwYwbP7Zy3a15K79vL50gpHyM5dmpt+0w/JRPzXbs5FkkM0162CxmMAwjb5sI/U+lv+22i//CsLjm+99d9twcHWcYsdKQkZl3ogZMOOCEI0tPOJKAmAhCybPIy81GwBdEMkGRl1MGsykfdTWt0pGDe9b/kd9//a9l9R8dldJeRU9cdvmuwT1N6z9w9faB51iaZbHB5/FS0DRsViMYxKnP0wZJ9FEGIdw4ChhSXgqNKfNaYev6Id/ZYhWUzpoxb/oY5ZpNYupD9mT+nn4XTH9/PwRBQCQSgauvJ+jbhZPOrSqxa5TNWrzs3Nnjcr7/fXGr8qtBg4oz2lpOQAlATIYgSTEa9LrAQQIHCjkjA6EsrKZs8KwO9XUtCY1r0/Wm5dsu4Jad+uDhO/3H59DMpTR0xX2+y6+Nbqva+fWXn4iJaD0n54IZVhOVq+XwB4JIUx7hqIDMTENsqRyswcST0lI7e9/H7+ySlq75V/MwdoxFqqxUjfzTXY886w3/+Ha+9HWEZCpOMNo1UGUQRFO+xIeA8MN0R54n2TfdsuKrieMruOFDVf+atHlrheXKUWMHTTKb5Kg7ADy6CAMZOgdC7iSQVsCUkQ0hwYNSDdQqW/LY0cbG3Z+99fjAsPXWMx/sePPSn9iu9Ev61cs8P1TrDfRYYDG55Gzra+zqwdB98gbO0g+zvX3uqBmyWDwIY4YZy/76xmahHkUXpbwkEPVg/NR89fO9z718Oy66AgAUpBpGk44Uzq1YsnrpMd3eCWTRuov+t0AuJ2Qkg+P6E+2QUgLavK2Jcunfp/zWXUwGuSfM3jVvdpUpGjpOA93bWgGgsI3oH8oa94LJ4qWtjY24556Fud84PjxoapcSuVnjlF1t+6mi2IJQIN21+/3335K/hA+jV6B1CaWp/5bNb3KCcumCkU+dH532YTiimvfkHWgW0mFnfX0Lvtm+n+7be1isKcIdI+cjL8uiQl62GkRyI0NvuSgPJB8AEgE3OOqDnPOQvIJhV87ek/Ph2R+l/9VE9N+EZCjqEaKiF0odhwkzyo8+Sv73Ns38Gxleb7MeqCi1GHs7j1Ex6iUzRkMEgOvfxrPDh+VrGxp3obDIDLfbY77sIirs/fLlD1QqBUyZFiQECbOmDT9x2QXZPdxL0++sv2j8gUcvHjbsv+Xym8DNsajDRYNGna/M/vOb77bN2pvtKMll0imq4hmkAkcPH3qGOo2t8MV9TrBCGK6248gz6/lLr8afASDpzk5q+UIE+uMYXlWBwWNmXDB986xP3f9vaSl+L9IaFYnLWQ4cI0e4ny36l/MzSclLZHL1vLlnaVKxCIqyC2Ex5OCTlUhdXkfURsPwK5IxgmxbNqxWPSqrTAkA6L6ZPlG99wOppesQ6hr2wCnIZ+0Wpq4uHnPG9aMnjhqcebCu/b/l8pvAfcXy7Qq3s8OfaTCQ8pISVq9Uw2G1kMsvuZg47KNHfvEoWXf7SoxMRf0wa3iSY1bDkaEgDvuoawFAG7XVGlUFCLqjxO/pRUWFhZSOyJ+389biawCALgPhCVXKWTnCnhDMOmE3AJxLCPvK0NKPJo0frNQqKfE6e0l2po3IqUq6thNHT/hQdt7sS1mfK0QIgM72BrJn72ex3r+TufP7squteRIZMbYI9jwT1Bk8seUaiS/YB6OydvmNbb+8o+aX9JvAbXmYCpeZ11zyzY7t6UgoCjnHI+gJoLOpHYXZuaSwfPa8L45PWenv7YaWS2NsZRHkYgjjq0bIAKBgl25N2A1kW3PgGWiH13sMQ4friduUtUK6kfDKx6kU9Ha5xXga6QTQ1x7UA8Cti8deOHnalKHdbYcQ9HYgN8uI9sZmHNtfG5pCafff3rEWR/0pkgiJaG1oQmamBr1uYdgJ4/B1g4bYM0urTCQQd8GenwWBxtHn7qEHard0fNpX/aumLH9Kv9mpdc0f6JZKffuDoVCcFpdUQa3JpHVHGmjDiTZKwIMjLNzOPvj6nTRDo0DY56RSMkKaiUBGbH5xa83efdHigmxqMnLo7DhIqeSko0aXWGavRAUA+N0tNYyogM3kAEeNxQCwN61/R6OgRMbFISVDVK+V0WjAQy29O58AgLa4cWhD/TEa8HipUq5AJBqgixZfSBzZVlJTtx89zg7a0esCq9AgmqC0sak55lvdMuntZfRnl4dOVb/pKwHeW1y3wpDcdPf+muMpXmUio8ZNJ4zMQHiFgSSSBMUl5SQSiZBYPEK0GjV6e1rFEiqjAHBfxpbz3ANNktUsI3I+RsKhLmLMkJGKG+XvAUCOAV9w0ELOaUlFvupgeipRTJkynAsFncSRZSIyTiDx6ACkeLfT/XrncwAwwXTiaHGhjSQTQWI0aLF1y3awjII0t/QQXqaDRpdDpp5xAZwDQN2RE53PX1I/7mnfj7ff/6f6TeE+BSqNWux9ZlHOlpFbd1S3+CJEsuUNhkydBV9ERE5ROfrcfpxoagMFiN954HAXIaSNEDL6PmwPDezYByYJmz0TqVQCkXAYfh8pAABHN7ZKSSU8riD98JPN+U9/Cz63wET0ej06OrqQTqeglBOcp665cwyQ2NRLyO0vYUci4UI46gLDSigrqyR1h1tRXjEVgwfPJip1Ga097BHffnnNa7cv7BrUUETrf0sep+1NIeJSwl25DjOvvnnixUeO+yZOnjA5jyGizO/twvvvbcGMaVbo9BZBrzQf7HH1GSjbwKQIipRy8DIe0OmKYDCPxcdraxN/e+e4auMyUtGmv6WeimEIvuqXKzNad1kG3/luZKAWLHXB7e9CIBiHTmbtZyWVj+cNksFmUH+2eWu+o4BFbn4l0qIOVZUzcexYU6zPua9Fw3VsqHkg/frzlLaeDgan9TUs38kZJWTACvOxR+Y2ObubdZMmViEWHSBWu40GwymwXAoGfRgmIwiLFPV4BHR3y4nbp0dzG5FeuFhbuWfiBx2v3n1VtDA/gwwfFHnu3Xc3njVx/JhyLukldquMGjIlotfrqUJmIz5Pknr9ISg0IPmDtDhcX0u7e+JobxUwNH/EumPvrr1s8X4kDKe58L8L3JwuYr9sxbjqKZPG5Q8dnI1jR6shk6cQCDoBGUVKCECIesDQJHRqEwzafFDkIRTPBPgifLv3eGrvoYPdi264smD48ALS0rA3HQv3MVajmug4DhwTQSjWingiBIaqkEz885xxll0LT6gB4XgCxYUToZRVovkEK25et3vd/Z9uutj6M1vufwv918PfU1FLFlw7x+y95MAx1du8XCxTG8zIyADkWhaCFIHekAONTAGkAEZUQ0rriWuARTAsCtu3r3nv07dffwL6cmfs2kue++SzDVcInnj5JedvrHxhtfTA9Mnnjcg0qaDW5cFgZaDkNfD7gwhH/BAogaOgGMFQGOGYgOrqauo8ln6tqerlh61c+WmvVb9Lzf1OfoGQ2EbM+GC3funQURMnGkxaMhCKwGTMglxSwtnlpf3tbjcbDW3Ojda9N/Sl5J4M+r+d+Tc/e/HBXZu23/Xa3z/RfffdRfKxBfe+lzOzvn/gMrXWNMJmKdASCqRTCegyAF/4BI4c2x9CPPb+PVOwMnkGbfm9yvu71NzvlCGjNGM+tvi0ZEeopa3e7TOUtvf1tw8qCVVL7kTvTLH2q4xlYp0PiJoopXjx3/N3+fqQYv79bS1rkvvaAby64mPyxsT1sCiGYVCLK2O0MauwwBkg13PKAfbGq2MLNNl0x8+80uH0iFL6+1w2qG4onD7inEyc3bZr7tpNH86UPnhphGsrBXeqNqY+OHru5X++YvWppr94Im48XrtUevuNCd4nH8bNX/8DE68qgUP4f3fs6b5+l2bB9gVRPdV/5z6bzVSq5AW+t/cEcuxacQh5fbLqLLrvVO0QgZB4MAKFWX1KQd8jEGbQa9PX2XIqz1FrjZDxGdKh/W2xnFzVE/MvWPHoL1v47/S7vMvxqeBNz1KFvDIspHhRphdF3uZtavv0ml8DFgCojNJTBQsAT8qotMi8bcH+/XVvdXcmY6GQiNHjRmg6uwMPKewk55ct/Hc67TX3vMvzbMPKJrXnF9gEW7b/yUN1b3y99ks0Hdz83886naqSAiFte5H10hqMKCpYeEdBYfn0np5DX958y7pzT6vj09nmhKMgq1bd9I+blkw6dNvDsP1u7fsvXLX7rrj8liVTQ4KwZNLp9HNaCzFprPaKxx5bcN8eAfz/NdAfXoNzULjwkiHVI0ci93T5+P8AJGUIUpjGozEAAAAASUVORK5CYII="
EMBEDDED_SIG_B64 = "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDAAMCAgMCAgMDAwMEAwMEBQgFBQQEBQoHBwYIDAoMDAsKCwsNDhIQDQ4RDgsLEBYQERMUFRUVDA8XGBYUGBIUFRT/2wBDAQMEBAUEBQkFBQkUDQsNFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBT/wAARCAA0AH8DASIAAhEBAxEB/8QAHwAAAQUBAQEBAQEAAAAAAAAAAAECAwQFBgcICQoL/8QAtRAAAgEDAwIEAwUFBAQAAAF9AQIDAAQRBRIhMUEGE1FhByJxFDKBkaEII0KxwRVS0fAkM2JyggkKFhcYGRolJicoKSo0NTY3ODk6Q0RFRkdISUpTVFVWV1hZWmNkZWZnaGlqc3R1dnd4eXqDhIWGh4iJipKTlJWWl5iZmqKjpKWmp6ipqrKztLW2t7i5usLDxMXGx8jJytLT1NXW19jZ2uHi4+Tl5ufo6erx8vP09fb3+Pn6/8QAHwEAAwEBAQEBAQEBAQAAAAAAAAECAwQFBgcICQoL/8QAtREAAgECBAQDBAcFBAQAAQJ3AAECAxEEBSExBhJBUQdhcRMiMoEIFEKRobHBCSMzUvAVYnLRChYkNOEl8RcYGRomJygpKjU2Nzg5OkNERUZHSElKU1RVVldYWVpjZGVmZ2hpanN0dXZ3eHl6goOEhYaHiImKkpOUlZaXmJmaoqOkpaanqKmqsrO0tba3uLm6wsPExcbHyMnK0tPU1dbX2Nna4uPk5ebn6Onq8vP09fb3+Pn6/9oADAMBAAIRAxEAPwD9U6KKKACiiigAooooAKKbwgrxqT4zXfjrzoPhXpMXi6RfkPiS5ufs2hLJj+G4HmPcdTxbRyDKbJJI6APTdd1zTfC2k3Wpatf2um6bZx+ZcXt5cJHFbx/35JH4H415V4U/aW0zx/8AFC38IeGPD+uX9lJpJ1mbxJPp8lpp6QPJ5dvs80B5BLskMcgTZJ5Z8suPMMfPeNfB8Phax07W/F0p+LHxEvL82/hy01GP7Pp1veufMj+z2X7xLeO3Ebyvcv5txHHHL+8k4jOl4X1nwn8BjdeHbrWLvxt8UNWA1XV0020Nxq2q3BHl/aJII/8Aj3t+I4ozJ5cEf7uPzO9AHteparbaVYz3l9NFa21ujyyz3D7EjROrmvm7xj8a9Q+IFlbNomtXngbwZqsMn9n6laWZufEviMpz/wASqzMcmyLy/wB59oeOR/L/AHnlxx+XcVy0z+MP25jrOi6laXnw3+EenamLWe5tbiG7v/EdxA7b7cP+8t0gjl+ST/j4jkki8tJMeZXpVj4t+CnwK0bxp4gk1rT9ETTr4aRr2tarey3l7c3kdulwlvJczF7i7kSKQYjDyY/eJ/BgAFj9nH4LS/DS78WeJLuHVdP1DxO9tK2katrtxqtzCkEZjjkuLiSSTzLiQP8AvPLPlpsjjj3iPzZPe6K8j+Nfx/8ADXwVtdNt9TF5rHiHV5jbaL4b0ePztQ1OT0jTPCf35Hwg9eRkA9cory3xZ8atD+FHw40nxT8S7mPwYlz9ngms/NlvXS8kT/j2j8qPfO+d/wDq05EZfGK6fwB4+8P/ABO8I6Z4n8OanFqeialH5ltdxsV3jvn0/wB2gDq6KKKACvmr4sfH/wAZpoXiPU/hL4a0zxRp3hm3uJ9S1zWb2W3s7hoP9bZ2YiDvcznZInmfJFHINm95BIidt8S9VvvF/iSz+HOh3lxbPqFt9t17UbGQrLYaXzHhJE/1VxcSbkj5HyRXMkZ3xVW+P2q6H8Gv2XfG80Edpomj6L4buLKytk/dwR/6OYreCPHq/lxj60AdnpPxO8P6l8OdK8cXOp2+keG9QsLfUY73UJUt444J40ePzS5+T74GDXB6h+1X4Gu72PTPA2oRfFHxFPDJLbaL4Nuo72TAMY33E4k8i2j+cfvLiSPOPk3nivjD9n79mPx7+0Z4G8G65rmtXmleGo9Cs9Nt7rxVo1nfXEUcflof7Hs3Hl2SeVGI/tkhkuJPLjk8tM19S+A/A3wo/YV8J/2F4esr+91zVwZls7aKTUNd16SJP+eaAeZs+iRRmT+DzOQDqI/gzrnxOlF78WNXh1LTHKSx+BtJZ00eDnOLmT/Wah3H7zy4CP8Al3zzW34/+OHgv4WJNpc96txqtna/aBoGkbXube32nEsqA+Xb24xzcXBjgTPMgrnH8MfEz4wGZvEmpXfwq8MSARRaJ4evEfWbj94ebm9GUtxsP+rs8yd/tHHl1zXh/wCHeheOPEr+EPDmgWln8JfDd/52r3UWCPEmsxn/AFcj/wDLxHbyJvuJJPM8yeOOPP7u4SgCr4G+H/jn4+eLR8TfGGo6t4B8OajZ/YtG8IWEgt9UtrAnzPMuLwfvIJLnO+SO38uSMRW6eZ+7kFa/ibw/a61qL/Cb4b29loPh6P8Af+NtU06TZPbxynzPscbj94by4/5aSH95FFJ5nEkkFUf2lv2iPEPhjXP+FbfCLQL3xf8AFu+t/PaJIz9i0i3fP+kXE8n7tPuSeWndwP8Ack8y8afsO/Ezxj8JbfwhZfEa38JWmqao9zrllpKSSRyxSRyfaJLm4I+0andySbPMkkkgjk5/dx+XGKAPavD/AO0t8NNG1eTwj4cmjHhTw1a/Z9T8RWixReHdEEaRxxW8l5JJHH5n7yNBHH5nbOyvMfCv7QX7OOoeEbvwYNGu7nwheX0d1p8Ov6LPqMOvSXl25+2RQSCS4kj+2P5f2iWOOPzTHHG/QDE8Q/8ABMHQL6y+GGlad4z1SLSvCdy8t2dZT+0ftAMsciRQW0j/AGS35En/AC7yB9/7zzOd/rnj/wDZI8AXPw6isX17WfBsGmalH4jvPFdnqMceo3FzbpJ/pN5e3Ik8zy97yZf/AFflp5fliOgDZ+O/7SGn/CbS9QsbGa11bxRbQ/aZ7eSV/s2lxc/6ReyIP3cfaOP/AFlw/wC7iQnJTlv2bP2eNT8P69rHxU+IF/da98SvEvypNqcMaSaRZ/8ALO3jjTIiPR3jjc7OI98nl+ZIz4FfCzQvGTaN4kstJu7PwFp9y+q+H9O1WSR7nV7yTOdd1DzPnllkx/o4k/1ceJDhzHHb/UlAHzF8RNWuJP2orTU5vDt54kg8BeC5dXgstMjjkvZLi/vfIBijeSMb0t7C4xzk/aJNnPXd/ZY8N6xo+k+PdevfD954QsvFviabxLp+h6kqve2aXFvbiQ3CxHakkkkckvlA7k8zDkmtvwxPa337VHxFZJI3ntPCvh6znwf9W/2jWJPLP/AJI3/EV7LQAVwXxR8cW3w48Om+eOXUL+5kFppmkQNi41S8fmO3iHqdhyT8iRiSST5I3Iyf2hviD4j+Gfwf8TeJPCPhm68WeKrO3Q6dpFnbS3Mk8skscY/dxfvJEQv5j7P4IzXB/CTwN4i0bSbP4k/Fie91vxfYaWUsbGOzFzLpNt5aeb5UdvH+8vLgoDKYx38qP92nzgHo3wh8I33hXw5NdeIrqK/8X6rcHUNcvYPM8iS5cD93bBzlLeONY4o/9iMb/wB5vNX/AIifDPR/ilb6Naa600tppmq22sR2kMxRLie3YyRiUD/WR7/LfYe8aVkeC/iX4i8d3VlNZ+BdX0XQXj3z3vif/QLlPl+5HbASSOc9fM8vHbfXfaZFfR2apqc8V1dAfPJb25hR/wDgBd/50Aed/Hrx3f8AgTwlp9toTWv/AAlPiXVbbQdI+2DFvHcT5zLJ6pFHHPPs/wCWnleXnLipfhX8FdI+GLX+rrNc654t1P8Aeax4n1b95qGoScHr0ijGPkt49kceeErd+J/w1034qeFZtC1Ga7sv3iXNtqOnS+Td2Vyn+ruIJP4JEPfFebX/AOzn4j8UWM2meLfjJ4z1zQJJMyadYpp+lvdx55gnntreOQoQSn7t485oAv8AxeuPG3jPXIvBvhRb3w5oElt5+u+NY2LSW0Hz/wCiafGh8w3kmB+9xsjQ5+eQpHWn4a+Ifws+Gfh+w8M2WpaZ4I0+wH2Oz03Vz/ZWfaOO48vfVb4hfstfD34naBovhrW9HuE8JaPC8Vp4f0m/udOsc/u8eZHbvH5mzy/kyeN78c5rd8CfAP4b/C+7huvCXgXw/wCG72O3Nqb3TNMgiuJY+Mo8oTzH+4OpoAtav8TNA0vU4LHddXl9dgfZYrG1kk+08/wPjy5MfX5O+K6vSLyfUbJbiW0n053GPs1wYzIn18t3T8jWh5Sb9+0b/Wn0AFc/4n8JaV4v0WXStZtI9Q0254uLO4G6KdcH5JI+jp/sHiugooAZHGIxT6K88+NPgK++J3wz17wjputT+GZNYhWzuNRtbbzJEt5JALhIweA8kfmR+Z/yz8zf25APnTwb+1prtpD8QfHsvw+8Wa/8MRr1zPH4js54Jba20+3ijtPtFtA8kcs8BNs9xJsj/d/aH/1nl5r3yH4yaL4t1CPSfA17YeNdTKRz3L6dfr9isrd4y8ctxdRb/KEnHloN8j+ZkJ5Ykkj7fQPDmneEdA03RNJto7LStOt47OztoxiOCKNBGkf4Ck8K+E9F8FaTFpOgaPY6HpkX3LLTbZLeKP6IgxQBuUUUUAFfL37Jf7R3ib4/fEP436J4gstLsrTwR4g/svTn02KVJJovMuFzKZJHBbEKcoEGc8ehRQB9Q0UUUAFFFFABRRRQAUUUUAFFFFABRRRQB//Z"

DEFAULT_UCAPAN_DONATUR = (
    "Kami mengucapkan terima kasih yang sebesar-besarnya kepada para donatur "
    "yang telah menyisihkan sedikit rezekinya untuk kegiatan ini sehingga dapat "
    "berjalan dengan sangat baik. Semoga bantuan dari Bapak/Ibu dan Keluarga "
    "dihitung Allah SWT sebagai amal kebaikan yang bernilai pahala, bermanfaat "
    "untuk sesama, dan menjadi berkah untuk kita semua. Selain itu, semoga "
    "dengan kegiatan ini semakin memperkuat tali silaturahmi dan kekeluargaan kita."
)

PRESET_RAB_JEMBER = {
    "id_rab": "RAB-2026-JEMBER",
    "baris_1_judul": "RANCANGAN ANGGARAN BIAYA",
    "nama_kegiatan_rab": "PELAKSANAAN PERTEMUAN RUTIN PAGUYUBAN DHARMA WANITA PERSATUAN DJBC DI LINGKUNGAN KANWIL DJBC JATIM II",
    "lokasi_acara_rab": "(DI KPPBC TMP C JEMBER)",
    "sumber_dana_rab": "Kas Dharma Wanita KPPBC TMC Malang",
    "ket_tabel_rab": "",
    "merge_ket_rab": False,
    "catatan_note_rab": NL.join([
        "Iuran DWP Gabungan untuk 5 orang Rp 300.000 + (2 orang x Rp 75.000)",
        "E-Toll (PP) dari Tol Singosari – Tol Probolinggo Timur dan sebaliknya Rp 112.000 x 2",
    ]),
    "status_konversi": "BELUM DIKONVERSI",
    "items_rab": [
        {"NO": 1, "KEBUTUHAN KEGIATAN": "Iuran DWP Gabungan di Jember", "JUMLAH": 7, "HARGA SATUAN": 0, "HARGA TOTAL": 450000, "KETERANGAN": ""},
        {"NO": 2, "KEBUTUHAN KEGIATAN": "Doorprize", "JUMLAH": 7, "HARGA SATUAN": 50000, "HARGA TOTAL": 350000, "KETERANGAN": ""},
        {"NO": 3, "KEBUTUHAN KEGIATAN": "Mobil (Bensin)", "JUMLAH": 2, "HARGA SATUAN": 0, "HARGA TOTAL": 0, "KETERANGAN": ""},
        {"NO": 4, "KEBUTUHAN KEGIATAN": "Sopir", "JUMLAH": 2, "HARGA SATUAN": 0, "HARGA TOTAL": 0, "KETERANGAN": ""},
        {"NO": 5, "KEBUTUHAN KEGIATAN": "Snack", "JUMLAH": 10, "HARGA SATUAN": 20000, "HARGA TOTAL": 200000, "KETERANGAN": ""},
        {"NO": 6, "KEBUTUHAN KEGIATAN": "Air Mineral 700ml (isi 12 botol)", "JUMLAH": 2, "HARGA SATUAN": 23000, "HARGA TOTAL": 46000, "KETERANGAN": ""},
        {"NO": 7, "KEBUTUHAN KEGIATAN": "Sarapan pagi", "JUMLAH": 10, "HARGA SATUAN": 30000, "HARGA TOTAL": 300000, "KETERANGAN": ""},
        {"NO": 8, "KEBUTUHAN KEGIATAN": "Makan Malam", "JUMLAH": 10, "HARGA SATUAN": 30000, "HARGA TOTAL": 300000, "KETERANGAN": ""},
        {"NO": 9, "KEBUTUHAN KEGIATAN": "E-Toll (Pulang-Pergi)", "JUMLAH": 2, "HARGA SATUAN": 224000, "HARGA TOTAL": 448000, "KETERANGAN": ""},
    ],
}

# PERBAIKAN MUTLAK V9.0: LPJ Blitar dan RAB Jember adalah 2 contoh kegiatan berbeda & TIDAK saling terkait!
PRESET_LPJ_BLITAR_2026 = {
    "id_lpj": "LPJ-20260813-BLITAR",
    "id_rab_ref": "",
    "tgl_laporan": date(2026, 8, 13),
    "nama_kegiatan": "Pertemuan Rutin Paguyuban Dharma Wanita Persatuan DJBC di lingkungan Kanwil DJBC Jawa Timur II",
    "tema_kegiatan": "Kebersamaan dalam Keceriaan, Bersinergi dalam Kebahagiaan",
    "konteks_hut": "menyambut HUT ke-81 Kemerdekaan Republik Indonesia",
    "tgl_acara": date(2026, 8, 10),
    "tgl_acara_manual": "",
    "waktu_acara": "08.30 - selesai",
    "tempat_acara": "KPPBC Tipe Madya Pabean C Blitar",
    "rangkaian_kegiatan": NL.join([
        "Pembukaan acara dengan sambutan dari Pembina DWP BC Blitar Bapak Nurtjahjo, Ketua DWP BC Blitar Ibu Hasrifah Nurtjahjo, dan Ibu Ketua DWP Kanwil DJBC Jatim II Ibu Ana Lukman.",
        "Kegiatan lomba menghias tumpeng dan lomba sambung lirik lagu Nasional untuk menyambut HUT ke-81 Kemerdekaan Republik Indonesia.",
        "Kunjungan ke Museum Bung Karno.",
    ]),
    "sumber_dana": "Kas Dharma Wanita KPPBC TMC Malang dan donatur",
    "ket_tabel_dana": "Kas DWP BC Malang & Donatur Anggota",
    "merge_ket_vertikal": True,
    "tampilkan_ucapan_donatur": True,
    "teks_ucapan_donatur": DEFAULT_UCAPAN_DONATUR,
    "catatan_ig": "*Note : dokumentasi kegiatan dapat dilihat di Instagram @dwpbcmalang",
    "jabatan_baris_1": "Ketua Dharma Wanita Persatuan",
    "jabatan_baris_2": "KPPBC Tipe Madya Cukai Malang",
    "nama_penandatangan": "Ny. Rini Johan Pandores",
    "tampilkan_ttd": True,
    "items_dana": [
        {"NO": 1, "KEBUTUHAN KEGIATAN": "Iuran DWP Gabungan di Blitar", "JUMLAH": 10, "HARGA SATUAN": 0, "HARGA TOTAL": 675000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 2, "KEBUTUHAN KEGIATAN": "Doorprize", "JUMLAH": 10, "HARGA SATUAN": 0, "HARGA TOTAL": 0, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 3, "KEBUTUHAN KEGIATAN": "Mobil", "JUMLAH": 2, "HARGA SATUAN": 0, "HARGA TOTAL": 0, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 4, "KEBUTUHAN KEGIATAN": "Hadiah Peserta Lomba Tumpeng", "JUMLAH": 3, "HARGA SATUAN": 250000, "HARGA TOTAL": 750000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 5, "KEBUTUHAN KEGIATAN": "Kerudung Merah", "JUMLAH": 10, "HARGA SATUAN": 50000, "HARGA TOTAL": 500000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 6, "KEBUTUHAN KEGIATAN": "Baju Atasan Merah (Blouse)", "JUMLAH": 9, "HARGA SATUAN": 90000, "HARGA TOTAL": 810000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 7, "KEBUTUHAN KEGIATAN": "Sarapan Pagi (Roti)", "JUMLAH": 20, "HARGA SATUAN": 13000, "HARGA TOTAL": 260000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 8, "KEBUTUHAN KEGIATAN": "Snack (Keripik Usus)", "JUMLAH": 12, "HARGA SATUAN": 15000, "HARGA TOTAL": 180000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 9, "KEBUTUHAN KEGIATAN": "Makan Malam (Ayam Pemuda)", "JUMLAH": 13, "HARGA SATUAN": 100000, "HARGA TOTAL": 1300000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
        {"NO": 10, "KEBUTUHAN KEGIATAN": "Bahan Latihan Tumpeng, Support Makan Siang & Kopi Untuk Latihan", "JUMLAH": 3, "HARGA SATUAN": 200000, "HARGA TOTAL": 600000, "KETERANGAN": "Kas DWP BC Malang & Donatur Anggota"},
    ],
}

SEED_HISTORY_RAB_JEMBER = dict(PRESET_RAB_JEMBER)
SEED_HISTORY_LPJ_BLITAR_2026 = dict(PRESET_LPJ_BLITAR_2026)

# 5 Menu Navigasi Sidebar Jernih & Terfokus V9.0
MENU_RAB = "📋 1. Buat / Edit RAB"
MENU_LPJ = "📝 2. Buat / Edit LPJ"
MENU_PREVIEW = "🖨️ 3. Preview & Download"
MENU_CLOUD = "🗄️ 4. Riwayat & Hapus Data"
MENU_KOMPARASI = "📊 5. Komparasi RAB vs LPJ"
ALL_MENUS = [MENU_RAB, MENU_LPJ, MENU_PREVIEW, MENU_CLOUD, MENU_KOMPARASI]


# ==============================================================================
# BAGIAN 2: HELPER FUNGSI BAHASA, AKUNTANSI SEJAJAR & RUMUS IURAN DWP V9.0
# ==============================================================================

def parse_num(val) -> int:
    """Mengubah nilai string format Indonesia (Rp 250.000, 112rb, -) menjadi integer murni."""
    if val is None or isinstance(val, bool):
        return 0
    if isinstance(val, (int, float)):
        if pd.isna(val):
            return 0
        return int(round(float(val)))

    s = str(val).strip().lower()
    if not s or s in ("-", "rp -", "rp-", "nihil", "gratis", "free", "nan", "none"):
        return 0

    s = s.replace("rp", "").replace("idr", "").replace(TAB, "").strip()

    multiplier = 1
    if "juta" in s or "jt" in s:
        multiplier = 1000000
        s = s.replace("juta", "").replace("jt", "").strip()
    elif "ribu" in s or "rb" in s:
        multiplier = 1000
        s = s.replace("ribu", "").replace("rb", "").strip()

    if multiplier > 1:
        s_clean = s.replace(".", "").replace(",", ".")
        digits_dot = "".join(ch for ch in s_clean if ch.isdigit() or ch == ".")
        if not digits_dot:
            return 0
        try:
            return int(round(float(digits_dot) * multiplier))
        except Exception:
            return 0

    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.split(",")[0]
    elif "," in s and len(s.split(",")[-1]) <= 2 and len(s.split(",")[0]) > 3:
        s = s.split(",")[0]

    digits = "".join(ch for ch in s if ch.isdigit())
    return int(digits) if digits else 0


def fmt_ribuan(val) -> str:
    """Mengubah integer menjadi string titik ribuan Indonesia (2094000 -> 2.094.000)."""
    n = parse_num(val)
    if n < 0:
        return "-" + f"{abs(n):,}".replace(",", ".")
    return f"{n:,}".replace(",", ".")


def fmt_nominal_only(val, is_dash_if_zero: bool = True) -> str:
    """
    Mengembalikan hanya angka titik ribuan ("450.000", "2.094.000")
    atau tanda strip ("-" jika val <= 0 dan is_dash_if_zero=True, atau "0" jika False).
    """
    n = parse_num(val)
    if n <= 0:
        return "-" if is_dash_if_zero else "0"
    return fmt_ribuan(n)


def fmt_rp_cell(val, is_dash_if_zero: bool = True) -> str:
    """Format teks gabungan Rp untuk ringkasan metrik UI."""
    return "Rp " + fmt_nominal_only(val, is_dash_if_zero=is_dash_if_zero)


def fmt_rp_html_split(val, is_dash_if_zero: bool = True, bold: bool = False) -> str:
    """
    Merender flex container di Live Preview HTML A4 dengan Rp rata kiri
    dan nominal/strip rata kanan agar 100% identik dengan cetakan Word/PDF.
    """
    fw = "700" if bold else "400"
    nom = fmt_nominal_only(val, is_dash_if_zero)
    return (
        f'<div style="display:flex;justify-content:space-between;align-items:center;'
        f'width:100%;white-space:nowrap;font-weight:{fw};">'
        f"<span>Rp</span><span>{nom}</span></div>"
    )


def calc_tiered_iuran_dwp(jml: int) -> int:
    """
    Rumus baku Iuran DWP Gabungan Kanwil DJBC Jatim II:
    - Paket dasar 5 orang pertama = Rp 300.000
    - Setiap tambahan orang di atas 5 orang = Rp 75.000 / orang
    - jml <= 0 -> 0; jml <= 5 -> 300000; jml > 5 -> 300000 + ((jml - 5) * 75000)
    """
    q = parse_num(jml)
    if q <= 0:
        return 0
    if q <= 5:
        return 300000
    return 300000 + ((q - 5) * 75000)


def _is_tiered_iuran_nominal(val: int) -> bool:
    """Memeriksa apakah sebuah nominal merupakan hasil rumus paket Iuran DWP (300rb + k*75rb)."""
    n = parse_num(val)
    if n == 0:
        return True
    if n >= 300000 and (n - 300000) % 75000 == 0:
        return True
    return False


def _terbilang_dasar(n: int) -> str:
    satuan = [
        "", "satu", "dua", "tiga", "empat", "lima",
        "enam", "tujuh", "delapan", "sembilan", "sepuluh", "sebelas"
    ]
    if n < 12:
        return satuan[n]
    if n < 20:
        return _terbilang_dasar(n - 10) + " belas"
    if n < 100:
        hasil = _terbilang_dasar(n // 10) + " puluh"
        if n % 10 > 0:
            hasil = hasil + " " + _terbilang_dasar(n % 10)
        return hasil
    if n < 200:
        sisa = _terbilang_dasar(n - 100)
        return ("seratus " + sisa).strip()
    if n < 1000:
        hasil = _terbilang_dasar(n // 100) + " ratus"
        if n % 100 > 0:
            hasil = hasil + " " + _terbilang_dasar(n % 100)
        return hasil
    if n < 2000:
        sisa = _terbilang_dasar(n - 1000)
        return ("seribu " + sisa).strip()
    if n < 1000000:
        hasil = _terbilang_dasar(n // 1000) + " ribu"
        if n % 1000 > 0:
            hasil = hasil + " " + _terbilang_dasar(n % 1000)
        return hasil
    if n < 1000000000:
        hasil = _terbilang_dasar(n // 1000000) + " juta"
        if n % 1000000 > 0:
            hasil = hasil + " " + _terbilang_dasar(n % 1000000)
        return hasil
    if n < 1000000000000:
        hasil = _terbilang_dasar(n // 1000000000) + " miliar"
        if n % 1000000000 > 0:
            hasil = hasil + " " + _terbilang_dasar(n % 1000000000)
        return hasil
    return str(n)


def terbilang(n) -> str:
    """Mengubah nominal angka menjadi teks Terbilang Rupiah Title Case."""
    angka = parse_num(n)
    if angka <= 0:
        return "Nol Rupiah"
    teks = " ".join(_terbilang_dasar(angka).split())
    return teks.title() + " Rupiah"


def format_hari_tanggal_indo(dt) -> str:
    """Contoh: date(2026, 8, 13) -> Kamis, 13 Agustus 2026."""
    if isinstance(dt, str):
        s = dt.strip()
        if "," in s and any(b in s for b in ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]):
            return s
        try:
            dt = datetime.strptime(s[:10], "%Y-%m-%d").date()
        except Exception:
            return s
    if isinstance(dt, datetime):
        dt = dt.date()
    if not isinstance(dt, date):
        dt = date.today()

    nama_hari = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
    nama_bulan = [
        "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember"
    ]
    return f"{nama_hari[dt.weekday()]}, {dt.day} {nama_bulan[dt.month]} {dt.year}"


def normalize_items_df(records, default_ket: str = "", prev_records=None, include_hapus_col: bool = False) -> pd.DataFrame:
    """
    Menerapkan Logika 4-Mode (termasuk Mode 2A Rumus Pintar Iuran DWP Gabungan).
    Secara default tidak menyertakan kolom HAPUS? agar output cetakan Word/PDF/Database bersih,
    dan hanya menyertakan kolom HAPUS? bila include_hapus_col=True untuk UI st.data_editor.
    """
    if isinstance(records, pd.DataFrame):
        raw_list = records.to_dict(orient="records")
    elif isinstance(records, list):
        raw_list = records
    else:
        raw_list = []

    if isinstance(prev_records, pd.DataFrame):
        prev_list = prev_records.to_dict(orient="records")
    elif isinstance(prev_records, list):
        prev_list = prev_records
    else:
        prev_list = []

    rows = []
    idx_counter = 1
    for idx_pos, item in enumerate(raw_list):
        kebutuhan = str(item.get("KEBUTUHAN KEGIATAN", item.get("kebutuhan", "")) or "").strip()
        if not kebutuhan and parse_num(item.get("HARGA TOTAL", 0)) == 0 and parse_num(item.get("HARGA SATUAN", 0)) == 0:
            continue

        jml = max(0, parse_num(item.get("JUMLAH", item.get("jumlah", 1))))
        h_sat = max(0, parse_num(item.get("HARGA SATUAN", item.get("harga_satuan", 0))))
        h_tot_in = max(0, parse_num(item.get("HARGA TOTAL", item.get("harga_total", 0))))
        ket_raw = item.get("KETERANGAN", item.get("keterangan", default_ket))
        ket = str(ket_raw if ket_raw is not None else default_ket).strip()
        hapus_flag = bool(item.get("HAPUS?", False))

        prev_item = prev_list[idx_pos] if idx_pos < len(prev_list) else None
        prev_jml = max(0, parse_num(prev_item.get("JUMLAH", prev_item.get("jumlah", jml)))) if prev_item else None
        prev_htot = max(0, parse_num(prev_item.get("HARGA TOTAL", prev_item.get("harga_total", h_tot_in)))) if prev_item else None

        keb_lower = kebutuhan.lower()
        is_iuran_dwp = "iuran dwp" in keb_lower

        if h_sat > 0:
            qty_calc = jml if jml > 0 else 1
            h_tot = qty_calc * h_sat
        elif is_iuran_dwp and h_sat == 0:
            jml_changed = (prev_jml is not None and jml != prev_jml)
            htot_manually_overridden = (
                prev_htot is not None
                and h_tot_in != prev_htot
                and not jml_changed
                and h_tot_in != calc_tiered_iuran_dwp(jml)
            )
            if htot_manually_overridden:
                h_tot = h_tot_in
            elif jml_changed or h_tot_in == 0 or _is_tiered_iuran_nominal(h_tot_in):
                h_tot = calc_tiered_iuran_dwp(jml)
            else:
                h_tot = h_tot_in
        elif h_tot_in > 0:
            h_tot = h_tot_in
        else:
            h_tot = 0

        row_dict = {
            "NO": idx_counter,
            "KEBUTUHAN KEGIATAN": kebutuhan,
            "JUMLAH": jml,
            "HARGA SATUAN": h_sat,
            "HARGA TOTAL": h_tot,
            "KETERANGAN": ket,
        }
        if include_hapus_col:
            row_dict["HAPUS?"] = hapus_flag
        rows.append(row_dict)
        idx_counter += 1

    if not rows:
        empty_row = {
            "NO": 1,
            "KEBUTUHAN KEGIATAN": "Kebutuhan Kegiatan",
            "JUMLAH": 1,
            "HARGA SATUAN": 0,
            "HARGA TOTAL": 0,
            "KETERANGAN": default_ket,
        }
        if include_hapus_col:
            empty_row["HAPUS?"] = False
        rows.append(empty_row)

    return pd.DataFrame(rows)


def sync_iuran_note_rab(catatan_text: str, df_items: pd.DataFrame) -> str:
    """
    Otomatis memperbarui kalimat rincian di *Note : RAB menjadi:
    'Iuran DWP Gabungan untuk 5 orang Rp 300.000 + ({jml - 5} orang x Rp 75.000)'
    setiap kali jumlah orang pada baris Iuran DWP berubah.
    """
    if df_items is None or df_items.empty:
        return catatan_text or ""

    iuran_jml = None
    for _, r in df_items.iterrows():
        keb = str(r.get("KEBUTUHAN KEGIATAN", "")).lower()
        if "iuran dwp" in keb:
            iuran_jml = parse_num(r.get("JUMLAH", 0))
            break

    if iuran_jml is None or iuran_jml <= 0:
        return catatan_text or ""

    if iuran_jml > 5:
        new_iuran_line = f"Iuran DWP Gabungan untuk 5 orang Rp 300.000 + ({iuran_jml - 5} orang x Rp 75.000)"
    else:
        new_iuran_line = f"Iuran DWP Gabungan untuk {iuran_jml} orang (Paket 5 orang) Rp 300.000"

    lines = [ln.strip() for ln in (catatan_text or "").split(NL) if ln.strip()]
    replaced = False
    out_lines = []
    for ln in lines:
        clean_ln = ln.lstrip("•-* ").strip()
        if "iuran dwp" in clean_ln.lower():
            out_lines.append(new_iuran_line)
            replaced = True
        else:
            out_lines.append(clean_ln)

    if not replaced:
        out_lines.insert(0, new_iuran_line)

    return NL.join(out_lines)


def sync_data_editor_live(edited_raw: pd.DataFrame, state_df_key: str, ver_key: str, default_ket: str = "", is_rab: bool = False) -> tuple:
    """
    Membandingkan dataframe hasil editan dengan state sebelumnya (prev_records).
    Mengembalikan tuple (df_calc_clean, checked_indices) di mana checked_indices adalah
    daftar baris yang dicentang pada kolom 'HAPUS?'.
    """
    prev_df = st.session_state.get(state_df_key, pd.DataFrame())
    df_with_hapus = normalize_items_df(edited_raw, default_ket=default_ket, prev_records=prev_df, include_hapus_col=True)
    df_calc = df_with_hapus.drop(columns=["HAPUS?"], errors="ignore").copy()

    checked_indices = []
    if "HAPUS?" in df_with_hapus.columns:
        checked_indices = [int(i) for i, val in enumerate(df_with_hapus["HAPUS?"].tolist()) if bool(val)]

    raw_records = edited_raw.to_dict(orient="records") if isinstance(edited_raw, pd.DataFrame) else []
    calc_records = df_calc.to_dict(orient="records")

    need_force_refresh = False
    if len(raw_records) != len(calc_records):
        need_force_refresh = True
    else:
        for r_raw, r_calc in zip(raw_records, calc_records):
            if parse_num(r_raw.get("HARGA TOTAL", 0)) != parse_num(r_calc.get("HARGA TOTAL", 0)):
                need_force_refresh = True
                break
            if parse_num(r_raw.get("NO", 0)) != parse_num(r_calc.get("NO", 0)):
                need_force_refresh = True
                break

    st.session_state[state_df_key] = df_calc

    if is_rab:
        cur_note = st.session_state.get("catatan_note_rab", "")
        updated_note = sync_iuran_note_rab(cur_note, df_calc)
        if updated_note != cur_note:
            st.session_state["catatan_note_rab"] = updated_note
            need_force_refresh = True

    if need_force_refresh:
        st.session_state[ver_key] = int(st.session_state.get(ver_key, 0)) + 1
        st.rerun()

    return df_calc, checked_indices


def delete_rows_from_items_df(df_items: pd.DataFrame, indices_to_drop: list, default_ket: str = "") -> pd.DataFrame:
    """Menghapus baris berdasarkan daftar indeks (0-based) dan mengurutkan ulang nomor NO."""
    if df_items is None or df_items.empty:
        return normalize_items_df([], default_ket=default_ket)
    drop_set = set(indices_to_drop)
    remaining = [row for idx, row in enumerate(df_items.to_dict(orient="records")) if idx not in drop_set]
    return normalize_items_df(remaining, default_ket=default_ket)


def append_blank_row_to_items_df(df_items: pd.DataFrame, default_ket: str = "") -> pd.DataFrame:
    """Menambahkan 1 baris baru di bawah tabel rincian dana."""
    cur_list = df_items.to_dict(orient="records") if isinstance(df_items, pd.DataFrame) and not df_items.empty else []
    cur_list.append({
        "NO": len(cur_list) + 1,
        "KEBUTUHAN KEGIATAN": f"Item Kegiatan {len(cur_list) + 1}",
        "JUMLAH": 1,
        "HARGA SATUAN": 0,
        "HARGA TOTAL": 0,
        "KETERANGAN": default_ket,
    })
    return normalize_items_df(cur_list, default_ket=default_ket)


def detect_location_mismatch(lokasi_judul: str, df_items: pd.DataFrame) -> list:
    """Mendeteksi jika di Judul RAB tertulis kota X namun ada baris item yang masih tertulis kota Y."""
    kota_list = ["jember", "blitar", "malang", "surabaya", "kediri", "madiun", "probolinggo", "pasuruan", "banyuwangi", "tulungagung", "bojonegoro", "gresik", "sidoarjo", "batu"]
    lok_lower = (lokasi_judul or "").lower()
    kota_judul = [k for k in kota_list if k in lok_lower]
    if not kota_judul or df_items is None or df_items.empty:
        return []

    target_kota = kota_judul[0].title()
    warnings = []
    for _, r in df_items.iterrows():
        keb = str(r.get("KEBUTUHAN KEGIATAN", ""))
        keb_low = keb.lower()
        for k in kota_list:
            if k in keb_low and k.title() != target_kota:
                warnings.append({
                    "no": int(r.get("NO", 0)),
                    "item": keb,
                    "wrong_city": k.title(),
                    "target_city": target_kota,
                })
    return warnings


def replace_city_across_payload(doc_mode: str, payload: dict, new_city: str) -> dict:
    """Mengganti nama kota lama dengan kota tujuan baru saat menyalin dokumen dari History."""
    cloned = json.loads(json.dumps(payload, default=str))
    kota_list = ["Jember", "Blitar", "Malang", "Surabaya", "Kediri", "Madiun", "Probolinggo", "Pasuruan", "Banyuwangi", "Tulungagung", "Bojonegoro", "Gresik", "Sidoarjo", "Batu"]
    target_city = (new_city or "").strip().title()
    if not target_city:
        return cloned

    def _swap_city(text: str, upper_mode: bool = False) -> str:
        res = str(text or "")
        for k in kota_list:
            if k.lower() == "malang" and "tmc malang" in res.lower():
                continue
            if k.lower() in res.lower():
                rep = target_city.upper() if upper_mode else target_city
                res = res.replace(k.upper(), rep).replace(k.title(), rep).replace(k.lower(), rep)
        return res

    if doc_mode == "RAB":
        cloned["id_rab"] = f"RAB-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{target_city.upper()}"
        cloned["lokasi_acara_rab"] = f"(DI KPPBC TMP C {target_city.upper()})"
        cloned["status_konversi"] = "BELUM DIKONVERSI"
        new_items = []
        for it in cloned.get("items_rab", []):
            it_copy = dict(it)
            it_copy["KEBUTUHAN KEGIATAN"] = _swap_city(it_copy.get("KEBUTUHAN KEGIATAN", ""), upper_mode=False)
            new_items.append(it_copy)
        cloned["items_rab"] = new_items
    else:
        cloned["id_lpj"] = f"LPJ-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{target_city.upper()}"
        cloned["tempat_acara"] = f"KPPBC Tipe Madya Pabean C {target_city}"
        if target_city.lower() != "blitar" and any(w in str(cloned.get("rangkaian_kegiatan", "")).lower() for w in ["nurtjahjo", "bung karno", "blitar"]):
            cloned["rangkaian_kegiatan"] = NL.join([
                f"Pembukaan acara dan sambutan pada kegiatan di KPPBC Tipe Madya Pabean C {target_city}.",
                "Pelaksanaan kegiatan inti pertemuan rutin dan ramah tamah anggota Dharma Wanita Persatuan.",
                "Penutupan acara dan dokumentasi bersama.",
            ])
            cloned["tema_kegiatan"] = ""
            cloned["konteks_hut"] = ""
        else:
            cloned["rangkaian_kegiatan"] = _swap_city(cloned.get("rangkaian_kegiatan", ""), upper_mode=False)
        new_items = []
        for it in cloned.get("items_dana", []):
            it_copy = dict(it)
            it_copy["KEBUTUHAN KEGIATAN"] = _swap_city(it_copy.get("KEBUTUHAN KEGIATAN", ""), upper_mode=False)
            new_items.append(it_copy)
        cloned["items_dana"] = new_items

    return cloned


def smart_paste_parser(raw_text: str, default_ket: str = "") -> list:
    """Mengekstrak daftar belanja dari teks WhatsApp/Catatan ke format 4-Mode secara cerdas."""
    parsed = []
    if not raw_text or not raw_text.strip():
        return parsed

    lines = [ln.strip() for ln in raw_text.split(NL) if ln.strip()]
    for ln in lines:
        clean_ln = ln
        if "." in clean_ln[:4]:
            left_num, rest = clean_ln.split(".", 1)
            if left_num.strip().isdigit():
                clean_ln = rest.strip()

        if "|" in clean_ln or ";" in clean_ln:
            sep = "|" if "|" in clean_ln else ";"
            parts = [p.strip() for p in clean_ln.split(sep)]
            nama = parts[0] if len(parts) > 0 else "Item Kegiatan"
            jml = parse_num(parts[1]) if len(parts) > 1 else 1
            val_3 = parse_num(parts[2]) if len(parts) > 2 else 0
            val_4 = parse_num(parts[3]) if len(parts) > 3 else 0
            ket = parts[4] if len(parts) > 4 and parts[4] else default_ket

            if "iuran dwp" in nama.lower() and val_3 == 0 and val_4 == 0:
                h_sat, h_tot = 0, calc_tiered_iuran_dwp(jml)
            elif len(parts) >= 4 and val_3 == 0 and val_4 > 0:
                h_sat, h_tot = 0, val_4
            elif "total" in (parts[2].lower() if len(parts) > 2 else "") or "borongan" in clean_ln.lower():
                h_sat, h_tot = 0, val_3
            else:
                h_sat = val_3
                h_tot = (jml if jml > 0 else 1) * h_sat if h_sat > 0 else 0

            parsed.append({
                "KEBUTUHAN KEGIATAN": nama,
                "JUMLAH": jml if jml > 0 else 1,
                "HARGA SATUAN": h_sat,
                "HARGA TOTAL": h_tot,
                "KETERANGAN": ket,
            })
            continue

        lower_ln = clean_ln.lower()
        found_mult = False
        for sym in [" x ", " @ ", " * "]:
            if sym in lower_ln:
                idx_sym = lower_ln.rfind(sym)
                left_part = clean_ln[:idx_sym].strip()
                right_part = clean_ln[idx_sym + len(sym):].strip()
                h_sat = parse_num(right_part)
                tokens = left_part.split()
                if tokens and tokens[-1].isdigit():
                    jml = int(tokens[-1])
                    nama = " ".join(tokens[:-1]).strip()
                else:
                    jml = 1
                    nama = left_part
                parsed.append({
                    "KEBUTUHAN KEGIATAN": nama or "Item Kegiatan",
                    "JUMLAH": jml,
                    "HARGA SATUAN": h_sat,
                    "HARGA TOTAL": jml * h_sat,
                    "KETERANGAN": default_ket,
                })
                found_mult = True
                break
        if found_mult:
            continue

        if " total " in lower_ln:
            idx_tot = lower_ln.rfind(" total ")
            left_part = clean_ln[:idx_tot].strip()
            right_part = clean_ln[idx_tot + 7:].strip()
            h_tot = parse_num(right_part)
            tokens = left_part.split()
            if tokens and tokens[-1].isdigit():
                jml = int(tokens[-1])
                nama = " ".join(tokens[:-1]).strip()
            else:
                jml = 1
                nama = left_part
            parsed.append({
                "KEBUTUHAN KEGIATAN": nama or "Item Borongan",
                "JUMLAH": jml,
                "HARGA SATUAN": 0,
                "HARGA TOTAL": h_tot,
                "KETERANGAN": default_ket,
            })
            continue

        tokens = clean_ln.split()
        if len(tokens) >= 2:
            last_tok = tokens[-1]
            second_last = tokens[-2] if len(tokens) >= 2 else ""
            if last_tok in ("-", "0", "nihil") and second_last.isdigit():
                jml = int(second_last)
                nama = " ".join(tokens[:-2]).strip()
                h_tot_auto = calc_tiered_iuran_dwp(jml) if "iuran dwp" in nama.lower() else 0
                parsed.append({
                    "KEBUTUHAN KEGIATAN": nama or "Fasilitas",
                    "JUMLAH": jml,
                    "HARGA SATUAN": 0,
                    "HARGA TOTAL": h_tot_auto,
                    "KETERANGAN": default_ket,
                })
                continue
            elif last_tok.isdigit() and int(last_tok) < 500:
                jml = int(last_tok)
                nama = " ".join(tokens[:-1]).strip()
                h_tot_auto = calc_tiered_iuran_dwp(jml) if "iuran dwp" in nama.lower() else 0
                parsed.append({
                    "KEBUTUHAN KEGIATAN": nama or "Fasilitas",
                    "JUMLAH": jml,
                    "HARGA SATUAN": 0,
                    "HARGA TOTAL": h_tot_auto,
                    "KETERANGAN": default_ket,
                })
                continue
            else:
                nominal = parse_num(last_tok)
                if second_last.isdigit() and int(second_last) < 1000 and nominal > 0:
                    jml = int(second_last)
                    nama = " ".join(tokens[:-2]).strip()
                    parsed.append({
                        "KEBUTUHAN KEGIATAN": nama or "Item Kegiatan",
                        "JUMLAH": jml,
                        "HARGA SATUAN": nominal,
                        "HARGA TOTAL": jml * nominal,
                        "KETERANGAN": default_ket,
                    })
                    continue

        parsed.append({
            "KEBUTUHAN KEGIATAN": clean_ln,
            "JUMLAH": 1,
            "HARGA SATUAN": 0,
            "HARGA TOTAL": calc_tiered_iuran_dwp(1) if "iuran dwp" in clean_ln.lower() else 0,
            "KETERANGAN": default_ket,
        })

    return parsed


def build_paragraf_pembuka(nama_kegiatan: str, tema_kegiatan: str, konteks_hut: str) -> str:
    """Menyusun paragraf pembuka LPJ adaptif sesuai ketersediaan Tema dan Konteks HUT."""
    nama = (nama_kegiatan or "").strip() or "Kegiatan Dharma Wanita Persatuan KPPBC TMC Malang"
    tema = (tema_kegiatan or "").strip().strip('"')
    hut = (konteks_hut or "").strip()

    kalimat_1 = (
        "Puji syukur kehadirat Allah SWT yang telah melimpahkan rahmat dan rizki-Nya "
        f"kepada kita semua, yang mana telah terlaksana {nama}."
    )
    if tema and hut:
        kalimat_2 = f' Kegiatan ini mengusung tema "{tema}" serta {hut}, yang dilaksanakan pada :'
    elif tema:
        kalimat_2 = f' Kegiatan ini mengusung tema "{tema}", yang dilaksanakan pada :'
    elif hut:
        kalimat_2 = f" Kegiatan ini dilaksanakan dalam rangka {hut}, pada :"
    else:
        kalimat_2 = " Kegiatan tersebut dilaksanakan pada :"
    return kalimat_1 + kalimat_2


def build_mail_merge_dictionary(doc_mode: str, payload: dict) -> list:
    """Memetakan seluruh input form ke dalam tag Mail Merge (<<TAG>>) standar tata naskah DWP KPPBC TMC Malang."""
    if doc_mode == "RAB":
        df_items = normalize_items_df(payload.get("items_rab", []), payload.get("ket_tabel_rab", ""))
        total_rab = int(df_items["HARGA TOTAL"].sum()) if not df_items.empty else 0
        b3 = (payload.get("lokasi_acara_rab") or "").strip().upper()
        if b3 and not b3.startswith("("):
            b3 = f"({b3})"
        return [
            {"TAG MAIL MERGE": "<<ID_RAB>>", "NILAI / ISI DOKUMEN": str(payload.get("id_rab", "RAB-2026-JEMBER"))},
            {"TAG MAIL MERGE": "<<JUDUL_BARIS_1>>", "NILAI / ISI DOKUMEN": str(payload.get("baris_1_judul", "RANCANGAN ANGGARAN BIAYA")).upper()},
            {"TAG MAIL MERGE": "<<NAMA_KEGIATAN_KAPITAL>>", "NILAI / ISI DOKUMEN": str(payload.get("nama_kegiatan_rab", "")).upper()},
            {"TAG MAIL MERGE": "<<LOKASI_ACARA_KAPITAL>>", "NILAI / ISI DOKUMEN": b3},
            {"TAG MAIL MERGE": "<<SUMBER_DANA_RAB>>", "NILAI / ISI DOKUMEN": str(payload.get("sumber_dana_rab", ""))},
            {"TAG MAIL MERGE": "<<TABEL_6KOLOM_RAB>>", "NILAI / ISI DOKUMEN": f"{len(df_items)} Baris Rincian Anggaran (Rp Rata Kiri Sejajar TOTAL)"},
            {"TAG MAIL MERGE": "<<TOTAL_RAB_RP>>", "NILAI / ISI DOKUMEN": fmt_rp_cell(total_rab, is_dash_if_zero=False)},
            {"TAG MAIL MERGE": "<<TERBILANG_RAB>>", "NILAI / ISI DOKUMEN": terbilang(total_rab)},
            {"TAG MAIL MERGE": "<<CATATAN_NOTE_BULLET>>", "NILAI / ISI DOKUMEN": str(payload.get("catatan_note_rab", "")).replace(NL, " | ")},
        ]

    df_items = normalize_items_df(payload.get("items_dana", []), payload.get("ket_tabel_dana", ""))
    total_lpj = int(df_items["HARGA TOTAL"].sum()) if not df_items.empty else 0
    tgl_surat = format_hari_tanggal_indo(payload.get("tgl_laporan", date.today()))
    tgl_acara = (payload.get("tgl_acara_manual") or "").strip() or format_hari_tanggal_indo(payload.get("tgl_acara", date.today()))
    pembuka = build_paragraf_pembuka(
        payload.get("nama_kegiatan", ""),
        payload.get("tema_kegiatan", ""),
        payload.get("konteks_hut", ""),
    )
    rangkaian_lines = [ln.strip() for ln in str(payload.get("rangkaian_kegiatan", "")).split(NL) if ln.strip()]
    ucapan = str(payload.get("teks_ucapan_donatur", DEFAULT_UCAPAN_DONATUR)) if payload.get("tampilkan_ucapan_donatur", True) else "(Tidak Ditampilkan)"

    return [
        {"TAG MAIL MERGE": "<<ID_LPJ>>", "NILAI / ISI DOKUMEN": str(payload.get("id_lpj", "LPJ-20260813-BLITAR"))},
        {"TAG MAIL MERGE": "<<ID_RAB_REFERENSI>>", "NILAI / ISI DOKUMEN": str(payload.get("id_rab_ref", "-") or "-")},
        {"TAG MAIL MERGE": "<<TGL_SURAT_LPJ>>", "NILAI / ISI DOKUMEN": tgl_surat},
        {"TAG MAIL MERGE": "<<PARAGRAF_PEMBUKA_ADAPTIF>>", "NILAI / ISI DOKUMEN": pembuka},
        {"TAG MAIL MERGE": "<<HARI_TGL_ACARA>>", "NILAI / ISI DOKUMEN": tgl_acara},
        {"TAG MAIL MERGE": "<<WAKTU_ACARA>>", "NILAI / ISI DOKUMEN": str(payload.get("waktu_acara", "08.30 - selesai"))},
        {"TAG MAIL MERGE": "<<TEMPAT_ACARA>>", "NILAI / ISI DOKUMEN": str(payload.get("tempat_acara", ""))},
        {"TAG MAIL MERGE": "<<RANGKAIAN_ACARA_NOMOR>>", "NILAI / ISI DOKUMEN": f"{len(rangkaian_lines)} Butir Rangkaian Acara"},
        {"TAG MAIL MERGE": "<<SUMBER_DANA_LPJ>>", "NILAI / ISI DOKUMEN": str(payload.get("sumber_dana", ""))},
        {"TAG MAIL MERGE": "<<KET_TABEL_MERGED>>", "NILAI / ISI DOKUMEN": str(payload.get("ket_tabel_dana", ""))},
        {"TAG MAIL MERGE": "<<TABEL_6KOLOM_LPJ>>", "NILAI / ISI DOKUMEN": f"{len(df_items)} Baris Realisasi Dana (Rp Rata Kiri Sejajar TOTAL)"},
        {"TAG MAIL MERGE": "<<TOTAL_LPJ_RP>>", "NILAI / ISI DOKUMEN": fmt_rp_cell(total_lpj, is_dash_if_zero=False)},
        {"TAG MAIL MERGE": "<<TERBILANG_LPJ>>", "NILAI / ISI DOKUMEN": terbilang(total_lpj)},
        {"TAG MAIL MERGE": "<<PARAGRAF_UCAPAN_DONATUR>>", "NILAI / ISI DOKUMEN": ucapan[:120] + ("..." if len(ucapan) > 120 else "")},
        {"TAG MAIL MERGE": "<<CATATAN_IG>>", "NILAI / ISI DOKUMEN": str(payload.get("catatan_ig", ""))},
        {"TAG MAIL MERGE": "<<JABATAN_BARIS_1>>", "NILAI / ISI DOKUMEN": str(payload.get("jabatan_baris_1", "Ketua Dharma Wanita Persatuan"))},
        {"TAG MAIL MERGE": "<<JABATAN_BARIS_2>>", "NILAI / ISI DOKUMEN": str(payload.get("jabatan_baris_2", "KPPBC Tipe Madya Cukai Malang"))},
        {"TAG MAIL MERGE": "<<NAMA_PENANDATANGAN>>", "NILAI / ISI DOKUMEN": str(payload.get("nama_penandatangan", "Ny. Rini Johan Pandores"))},
    ]


# ==============================================================================
# BAGIAN 3: TRIPLE-LAYER AUTO-HISTORY & FULL CRUD (SIMPAN, HAPUS & ANTI-DOBEL)
# ==============================================================================

def _make_rab_row_and_rekap(rab_payload: dict, waktu_simpan: str = None) -> tuple:
    """Membentuk dictionary baris DATABASE_RAB dan list REKAP_ITEM_DANA dari payload RAB."""
    id_rab = rab_payload.get("id_rab") or ("RAB-" + datetime.now().strftime("%Y%m%d-%H%M%S"))
    rab_payload["id_rab"] = id_rab
    w_simpan = waktu_simpan or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    df_items = normalize_items_df(rab_payload.get("items_rab", []), rab_payload.get("ket_tabel_rab", ""))
    total_rab = int(df_items["HARGA TOTAL"].sum()) if not df_items.empty else 0

    serializable = dict(rab_payload)
    serializable["items_rab"] = df_items.to_dict(orient="records")

    rab_row_dict = {
        "ID_RAB": id_rab,
        "WAKTU_SIMPAN": w_simpan,
        "NAMA_KEGIATAN": rab_payload.get("nama_kegiatan_rab", ""),
        "LOKASI_ACARA": rab_payload.get("lokasi_acara_rab", ""),
        "SUMBER_DANA": rab_payload.get("sumber_dana_rab", ""),
        "KET_TABEL_DANA": rab_payload.get("ket_tabel_rab", ""),
        "TOTAL_RAB": total_rab,
        "CATATAN_NOTE": rab_payload.get("catatan_note_rab", ""),
        "STATUS_KONVERSI": rab_payload.get("status_konversi", "BELUM DIKONVERSI"),
        "PAYLOAD_JSON": json.dumps(serializable, ensure_ascii=False),
    }

    rekap_dicts = []
    for _, r in df_items.iterrows():
        rekap_dicts.append({
            "ID_DOKUMEN": id_rab,
            "JENIS_DOK": "RAB",
            "TGL_ACARA": rab_payload.get("lokasi_acara_rab", ""),
            "NAMA_KEGIATAN": rab_payload.get("nama_kegiatan_rab", ""),
            "NO_ITEM": int(r["NO"]),
            "KEBUTUHAN_KEGIATAN": str(r["KEBUTUHAN KEGIATAN"]),
            "JUMLAH": int(r["JUMLAH"]),
            "HARGA_SATUAN": int(r["HARGA SATUAN"]),
            "HARGA_TOTAL": int(r["HARGA TOTAL"]),
            "KETERANGAN": str(r["KETERANGAN"]),
        })
    return rab_row_dict, rekap_dicts


def _make_lpj_row_and_rekap(lpj_payload: dict, waktu_simpan: str = None) -> tuple:
    """Membentuk dictionary baris DATABASE_LPJ dan list REKAP_ITEM_DANA dari payload LPJ."""
    id_lpj = lpj_payload.get("id_lpj") or ("LPJ-" + datetime.now().strftime("%Y%m%d-%H%M%S"))
    lpj_payload["id_lpj"] = id_lpj
    w_simpan = waktu_simpan or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    tgl_lap_val = lpj_payload.get("tgl_laporan", date.today())
    if isinstance(tgl_lap_val, (date, datetime)):
        tgl_lap_str = tgl_lap_val.strftime("%Y-%m-%d")
    else:
        tgl_lap_str = str(tgl_lap_val)

    tgl_acara_str = (lpj_payload.get("tgl_acara_manual") or "").strip() or format_hari_tanggal_indo(lpj_payload.get("tgl_acara", date.today()))

    df_items = normalize_items_df(lpj_payload.get("items_dana", []), lpj_payload.get("ket_tabel_dana", ""))
    total_dana = int(df_items["HARGA TOTAL"].sum()) if not df_items.empty else 0

    serializable = dict(lpj_payload)
    if isinstance(serializable.get("tgl_laporan"), (date, datetime)):
        serializable["tgl_laporan"] = serializable["tgl_laporan"].strftime("%Y-%m-%d")
    if isinstance(serializable.get("tgl_acara"), (date, datetime)):
        serializable["tgl_acara"] = serializable["tgl_acara"].strftime("%Y-%m-%d")
    serializable["items_dana"] = df_items.to_dict(orient="records")

    lpj_row_dict = {
        "ID_LPJ": id_lpj,
        "ID_RAB_REF": lpj_payload.get("id_rab_ref", ""),
        "WAKTU_SIMPAN": w_simpan,
        "TGL_LAPORAN": tgl_lap_str,
        "NAMA_KEGIATAN": lpj_payload.get("nama_kegiatan", ""),
        "TEMA_KEGIATAN": lpj_payload.get("tema_kegiatan", ""),
        "KONTEKS_HUT": lpj_payload.get("konteks_hut", ""),
        "TGL_ACARA": tgl_acara_str,
        "WAKTU_ACARA": lpj_payload.get("waktu_acara", ""),
        "TEMPAT_ACARA": lpj_payload.get("tempat_acara", ""),
        "SUMBER_DANA": lpj_payload.get("sumber_dana", ""),
        "KET_TABEL_DANA": lpj_payload.get("ket_tabel_dana", ""),
        "TOTAL_DANA": total_dana,
        "PENANDATANGAN": lpj_payload.get("nama_penandatangan", ""),
        "PAYLOAD_JSON": json.dumps(serializable, ensure_ascii=False),
    }

    rekap_dicts = []
    for _, r in df_items.iterrows():
        rekap_dicts.append({
            "ID_DOKUMEN": id_lpj,
            "JENIS_DOK": "LPJ",
            "TGL_ACARA": tgl_acara_str,
            "NAMA_KEGIATAN": lpj_payload.get("nama_kegiatan", ""),
            "NO_ITEM": int(r["NO"]),
            "KEBUTUHAN_KEGIATAN": str(r["KEBUTUHAN KEGIATAN"]),
            "JUMLAH": int(r["JUMLAH"]),
            "HARGA_SATUAN": int(r["HARGA SATUAN"]),
            "HARGA_TOTAL": int(r["HARGA TOTAL"]),
            "KETERANGAN": str(r["KETERANGAN"]),
        })
    return lpj_row_dict, rekap_dicts


def load_persistent_local_history() -> dict:
    """
    Membaca file JSON lokal permanen `database_riwayat_dwp_malang.json`
    dan memastikan SEED_HISTORY_RAB_JEMBER (Rp 2.094.000) serta SEED_HISTORY_LPJ_BLITAR_2026
    (Rp 5.075.000, id_rab_ref="") tercatat sebagai Data Riwayat Bawaan.
    """
    seed_rab_row, seed_rab_rk = _make_rab_row_and_rekap(dict(SEED_HISTORY_RAB_JEMBER), waktu_simpan="2026-08-01 08:00:00")
    seed_lpj_row, seed_lpj_rk = _make_lpj_row_and_rekap(dict(SEED_HISTORY_LPJ_BLITAR_2026), waktu_simpan="2026-08-13 09:00:00")

    data = {"rab_rows": [], "lpj_rows": [], "rekap_rows": []}
    if os.path.isfile(LOCAL_HISTORY_JSON_FILE):
        try:
            with open(LOCAL_HISTORY_JSON_FILE, "r", encoding="utf-8") as f:
                loaded = json.load(f)
                if isinstance(loaded, dict):
                    data["rab_rows"] = loaded.get("rab_rows", [])
                    data["lpj_rows"] = loaded.get("lpj_rows", [])
                    data["rekap_rows"] = loaded.get("rekap_rows", [])
        except Exception:
            pass

    # Pastikan LPJ Blitar bawaan tidak pernah punya id_rab_ref ke RAB Jember
    for r_lpj in data["lpj_rows"]:
        if r_lpj.get("ID_LPJ") == "LPJ-20260813-BLITAR" and r_lpj.get("ID_RAB_REF") == "RAB-2026-JEMBER":
            r_lpj["ID_RAB_REF"] = ""
            try:
                p_obj = json.loads(r_lpj.get("PAYLOAD_JSON", "{}"))
                p_obj["id_rab_ref"] = ""
                r_lpj["PAYLOAD_JSON"] = json.dumps(p_obj, ensure_ascii=False)
            except Exception:
                pass

    rab_ids = {r.get("ID_RAB") for r in data["rab_rows"]}
    if seed_rab_row["ID_RAB"] not in rab_ids:
        data["rab_rows"].append(seed_rab_row)
        data["rekap_rows"].extend(seed_rab_rk)

    lpj_ids = {r.get("ID_LPJ") for r in data["lpj_rows"]}
    if seed_lpj_row["ID_LPJ"] not in lpj_ids:
        data["lpj_rows"].append(seed_lpj_row)
        data["rekap_rows"].extend(seed_lpj_rk)

    try:
        with open(LOCAL_HISTORY_JSON_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

    return data


def save_persistent_local_history(rab_rows: list = None, lpj_rows: list = None, rekap_rows: list = None) -> bool:
    """Menyimpan daftar riwayat RAB, LPJ, dan Rekap Item Dana ke file JSON lokal `database_riwayat_dwp_malang.json`."""
    r_rows = rab_rows if rab_rows is not None else st.session_state.get("local_db_rab", [])
    l_rows = lpj_rows if lpj_rows is not None else st.session_state.get("local_db_lpj", [])
    rk_rows = rekap_rows if rekap_rows is not None else st.session_state.get("local_rekap_dana", [])

    payload = {
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "rab_rows": r_rows,
        "lpj_rows": l_rows,
        "rekap_rows": rk_rows,
    }
    try:
        with open(LOCAL_HISTORY_JSON_FILE, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


def get_active_spreadsheet_id() -> str:
    """Otomatis membaca st.secrets['database']['spreadsheet_id'] jika tersedia, atau fallback ke DEFAULT_SPREADSHEET_ID."""
    custom_id = st.session_state.get("cfg_sheet_id", "").strip()
    if custom_id:
        return custom_id
    try:
        if "database" in st.secrets and "spreadsheet_id" in st.secrets["database"]:
            sec_id = str(st.secrets["database"]["spreadsheet_id"]).strip()
            if sec_id:
                return sec_id
    except Exception:
        pass
    return DEFAULT_SPREADSHEET_ID


@st.cache_resource(show_spinner=False)
def get_gspread_client(creds_json_str: str = ""):
    """Membuka koneksi gspread dengan dukungan otomatis dari upload, st.secrets, atau file lokal `.gcp_service_account_dwp.json`."""
    if not HAS_GSPREAD:
        return None, "Library gspread belum terpasang di environment ini."

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    creds_dict = None

    if creds_json_str and creds_json_str.strip():
        try:
            creds_dict = json.loads(creds_json_str)
        except Exception as e:
            return None, f"Format JSON Service Account tidak valid: {e}"

    if not creds_dict:
        try:
            if "gcp_service_account" in st.secrets:
                creds_dict = dict(st.secrets["gcp_service_account"])
        except Exception:
            creds_dict = None

    if not creds_dict and os.path.isfile(LOCAL_GCP_SA_FILE):
        try:
            with open(LOCAL_GCP_SA_FILE, "r", encoding="utf-8") as f:
                creds_dict = json.load(f)
        except Exception:
            creds_dict = None

    if not creds_dict:
        return None, "Riwayat Lokal Permanen Aktif (JSON)"

    try:
        creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
        client = gspread.authorize(creds)
        return client, "Terhubung ke Google Sheets Cloud"
    except Exception as e:
        return None, f"Gagal otentikasi Google Sheets: {e}"


def ensure_worksheets_ready(sh):
    """Memastikan 3 worksheet (DATABASE_RAB, DATABASE_LPJ, REKAP_ITEM_DANA) siap digunakan."""
    existing_titles = [ws.title for ws in sh.worksheets()]

    if SHEET_DB_RAB not in existing_titles:
        ws_rab = sh.add_worksheet(title=SHEET_DB_RAB, rows=500, cols=len(COLS_DB_RAB))
        ws_rab.batch_update([{"range": "A1:J1", "values": [COLS_DB_RAB]}], value_input_option="USER_ENTERED")
    else:
        ws_rab = sh.worksheet(SHEET_DB_RAB)

    if SHEET_DB_LPJ not in existing_titles:
        ws_lpj = sh.add_worksheet(title=SHEET_DB_LPJ, rows=500, cols=len(COLS_DB_LPJ))
        ws_lpj.batch_update([{"range": "A1:O1", "values": [COLS_DB_LPJ]}], value_input_option="USER_ENTERED")
    else:
        ws_lpj = sh.worksheet(SHEET_DB_LPJ)

    if SHEET_REKAP_DANA not in existing_titles:
        ws_rk = sh.add_worksheet(title=SHEET_REKAP_DANA, rows=2500, cols=len(COLS_REKAP_DANA))
        ws_rk.batch_update([{"range": "A1:J1", "values": [COLS_REKAP_DANA]}], value_input_option="USER_ENTERED")
    else:
        ws_rk = sh.worksheet(SHEET_REKAP_DANA)

    return ws_rab, ws_lpj, ws_rk


@st.cache_data(ttl=180, show_spinner=False)
def fetch_sheet_bundle_cached(spreadsheet_id: str, creds_json_str: str = "") -> dict:
    """Membaca DATABASE_RAB, DATABASE_LPJ, dan REKAP_ITEM_DANA sekaligus dengan 1 panggilan API."""
    client, status_msg = get_gspread_client(creds_json_str)
    if client is None:
        return {"ok": False, "msg": status_msg, "rab_rows": [], "lpj_rows": [], "rekap_rows": []}

    try:
        sh = client.open_by_key(spreadsheet_id)
        ensure_worksheets_ready(sh)
        batch_res = sh.values_batch_get([
            f"{SHEET_DB_RAB}!A1:J500",
            f"{SHEET_DB_LPJ}!A1:O500",
            f"{SHEET_REKAP_DANA}!A1:J2500",
        ])
        value_ranges = batch_res.get("valueRanges", [])

        rab_vals = value_ranges[0].get("values", []) if len(value_ranges) > 0 else []
        lpj_vals = value_ranges[1].get("values", []) if len(value_ranges) > 1 else []
        rk_vals = value_ranges[2].get("values", []) if len(value_ranges) > 2 else []

        def _to_records(vals, expected_cols):
            if not vals or len(vals) < 2:
                return []
            headers = vals[0]
            out = []
            for r in vals[1:]:
                if not any(str(c).strip() for c in r):
                    continue
                row_dict = {}
                for i, col_name in enumerate(expected_cols):
                    col_idx = headers.index(col_name) if col_name in headers else i
                    row_dict[col_name] = r[col_idx] if col_idx < len(r) else ""
                out.append(row_dict)
            return out

        return {
            "ok": True,
            "msg": "Cloud Google Sheets Aktif",
            "rab_rows": _to_records(rab_vals, COLS_DB_RAB),
            "lpj_rows": _to_records(lpj_vals, COLS_DB_LPJ),
            "rekap_rows": _to_records(rk_vals, COLS_REKAP_DANA),
        }
    except Exception as e:
        return {"ok": False, "msg": f"Gagal membaca GSheet: {e}", "rab_rows": [], "lpj_rows": [], "rekap_rows": []}


def get_sheet_bundle(spreadsheet_id: str, creds_json_str: str = "") -> dict:
    """
    Menggabungkan (merge) History dari Lapis 1 (st.session_state),
    Lapis 2 (File JSON Lokal `database_riwayat_dwp_malang.json`), dan Lapis 3 (Google Sheets Cloud)
    serta menghormati daftar ID yang baru saja dihapus (`deleted_doc_ids`).
    """
    deleted_ids = set(st.session_state.get("deleted_doc_ids", []))
    local_file_data = load_persistent_local_history()
    sess_rab = st.session_state.get("local_db_rab", [])
    sess_lpj = st.session_state.get("local_db_lpj", [])
    sess_rk = st.session_state.get("local_rekap_dana", [])

    res = fetch_sheet_bundle_cached(spreadsheet_id, creds_json_str)

    merged_rab = []
    seen_rab = set()
    for source_list in [sess_rab, local_file_data.get("rab_rows", []), res.get("rab_rows", [])]:
        for r in source_list:
            rid = r.get("ID_RAB")
            if rid and rid not in seen_rab and rid not in deleted_ids:
                seen_rab.add(rid)
                merged_rab.append(r)

    merged_lpj = []
    seen_lpj = set()
    for source_list in [sess_lpj, local_file_data.get("lpj_rows", []), res.get("lpj_rows", [])]:
        for r in source_list:
            lid = r.get("ID_LPJ")
            if lid and lid not in seen_lpj and lid not in deleted_ids:
                if lid == "LPJ-20260813-BLITAR" and r.get("ID_RAB_REF") == "RAB-2026-JEMBER":
                    r["ID_RAB_REF"] = ""
                seen_lpj.add(lid)
                merged_lpj.append(r)

    merged_rk = []
    seen_rk = set()
    for source_list in [sess_rk, local_file_data.get("rekap_rows", []), res.get("rekap_rows", [])]:
        for r in source_list:
            doc_id = r.get("ID_DOKUMEN")
            key_rk = (doc_id, str(r.get("NO_ITEM")))
            if doc_id and doc_id not in deleted_ids and key_rk not in seen_rk:
                seen_rk.add(key_rk)
                merged_rk.append(r)

    st.session_state["local_db_rab"] = merged_rab
    st.session_state["local_db_lpj"] = merged_lpj
    st.session_state["local_rekap_dana"] = merged_rk

    return {
        "ok": bool(res.get("ok")),
        "msg": res.get("msg", "Riwayat Lokal Permanen Aktif"),
        "rab_rows": merged_rab,
        "lpj_rows": merged_lpj,
        "rekap_rows": merged_rk,
    }


def save_rab_to_storage(rab_payload: dict, spreadsheet_id: str = "", creds_json_str: str = "") -> tuple:
    """Menyimpan dokumen RAB ke Lapis 1 (session_state), Lapis 2 (JSON Lokal Permanen), dan Lapis 3 (Google Sheets)."""
    rab_row_dict, rekap_dicts = _make_rab_row_and_rekap(rab_payload)
    id_rab = rab_row_dict["ID_RAB"]

    deleted_ids = set(st.session_state.get("deleted_doc_ids", []))
    if id_rab in deleted_ids:
        deleted_ids.remove(id_rab)
        st.session_state["deleted_doc_ids"] = list(deleted_ids)

    local_rab = [row for row in st.session_state.get("local_db_rab", []) if row.get("ID_RAB") != id_rab]
    local_rab.insert(0, rab_row_dict)
    st.session_state["local_db_rab"] = local_rab

    local_rk = [row for row in st.session_state.get("local_rekap_dana", []) if row.get("ID_DOKUMEN") != id_rab]
    local_rk.extend(rekap_dicts)
    st.session_state["local_rekap_dana"] = local_rk

    save_persistent_local_history(
        rab_rows=local_rab,
        lpj_rows=st.session_state.get("local_db_lpj", []),
        rekap_rows=local_rk,
    )

    target_sheet_id = spreadsheet_id or get_active_spreadsheet_id()
    client, status_msg = get_gspread_client(creds_json_str)
    if client is None:
        return True, f"RAB ({id_rab}) berhasil disimpan ke Riwayat Permanen. ({status_msg})"

    try:
        sh = client.open_by_key(target_sheet_id)
        ws_rab, _, ws_rk = ensure_worksheets_ready(sh)

        existing_vals = ws_rab.get_all_values()
        new_matrix = [COLS_DB_RAB]
        updated = False
        for r in existing_vals[1:]:
            if not any(str(c).strip() for c in r):
                continue
            if len(r) > 0 and r[0] == id_rab:
                new_matrix.append([rab_row_dict[c] for c in COLS_DB_RAB])
                updated = True
            else:
                padded = list(r) + [""] * max(0, len(COLS_DB_RAB) - len(r))
                new_matrix.append(padded[:len(COLS_DB_RAB)])
        if not updated:
            new_matrix.append([rab_row_dict[c] for c in COLS_DB_RAB])

        ws_rab.clear()
        ws_rab.batch_update(
            [{"range": f"A1:J{len(new_matrix)}", "values": new_matrix}],
            value_input_option="USER_ENTERED",
        )

        existing_rk = ws_rk.get_all_values()
        new_rk = [COLS_REKAP_DANA]
        for r in existing_rk[1:]:
            if not any(str(c).strip() for c in r):
                continue
            if len(r) > 0 and r[0] != id_rab:
                padded = list(r) + [""] * max(0, len(COLS_REKAP_DANA) - len(r))
                new_rk.append(padded[:len(COLS_REKAP_DANA)])
        for d in rekap_dicts:
            new_rk.append([d[c] for c in COLS_REKAP_DANA])

        ws_rk.clear()
        ws_rk.batch_update(
            [{"range": f"A1:J{len(new_rk)}", "values": new_rk}],
            value_input_option="USER_ENTERED",
        )

        fetch_sheet_bundle_cached.clear()
        return True, f"RAB ({id_rab}) berhasil disimpan ke Riwayat Lokal & Google Sheets!"
    except Exception as e:
        return True, f"RAB ({id_rab}) tersimpan di Riwayat Lokal, namun gagal sinkron ke Cloud: {e}"


def save_lpj_to_storage(lpj_payload: dict, spreadsheet_id: str = "", creds_json_str: str = "") -> tuple:
    """Menyimpan dokumen LPJ ke Lapis 1 (session_state), Lapis 2 (JSON Lokal Permanen), dan Lapis 3 (Google Sheets)."""
    lpj_row_dict, rekap_dicts = _make_lpj_row_and_rekap(lpj_payload)
    id_lpj = lpj_row_dict["ID_LPJ"]

    deleted_ids = set(st.session_state.get("deleted_doc_ids", []))
    if id_lpj in deleted_ids:
        deleted_ids.remove(id_lpj)
        st.session_state["deleted_doc_ids"] = list(deleted_ids)

    local_lpj = [row for row in st.session_state.get("local_db_lpj", []) if row.get("ID_LPJ") != id_lpj]
    local_lpj.insert(0, lpj_row_dict)
    st.session_state["local_db_lpj"] = local_lpj

    local_rk = [row for row in st.session_state.get("local_rekap_dana", []) if row.get("ID_DOKUMEN") != id_lpj]
    local_rk.extend(rekap_dicts)
    st.session_state["local_rekap_dana"] = local_rk

    save_persistent_local_history(
        rab_rows=st.session_state.get("local_db_rab", []),
        lpj_rows=local_lpj,
        rekap_rows=local_rk,
    )

    target_sheet_id = spreadsheet_id or get_active_spreadsheet_id()
    client, status_msg = get_gspread_client(creds_json_str)
    if client is None:
        return True, f"LPJ ({id_lpj}) berhasil disimpan ke Riwayat Permanen. ({status_msg})"

    try:
        sh = client.open_by_key(target_sheet_id)
        _, ws_lpj, ws_rk = ensure_worksheets_ready(sh)

        existing_vals = ws_lpj.get_all_values()
        new_matrix = [COLS_DB_LPJ]
        updated = False
        for r in existing_vals[1:]:
            if not any(str(c).strip() for c in r):
                continue
            if len(r) > 0 and r[0] == id_lpj:
                new_matrix.append([lpj_row_dict[c] for c in COLS_DB_LPJ])
                updated = True
            else:
                padded = list(r) + [""] * max(0, len(COLS_DB_LPJ) - len(r))
                new_matrix.append(padded[:len(COLS_DB_LPJ)])
        if not updated:
            new_matrix.append([lpj_row_dict[c] for c in COLS_DB_LPJ])

        ws_lpj.clear()
        ws_lpj.batch_update(
            [{"range": f"A1:O{len(new_matrix)}", "values": new_matrix}],
            value_input_option="USER_ENTERED",
        )

        existing_rk = ws_rk.get_all_values()
        new_rk = [COLS_REKAP_DANA]
        for r in existing_rk[1:]:
            if not any(str(c).strip() for c in r):
                continue
            if len(r) > 0 and r[0] != id_lpj:
                padded = list(r) + [""] * max(0, len(COLS_REKAP_DANA) - len(r))
                new_rk.append(padded[:len(COLS_REKAP_DANA)])
        for d in rekap_dicts:
            new_rk.append([d[c] for c in COLS_REKAP_DANA])

        ws_rk.clear()
        ws_rk.batch_update(
            [{"range": f"A1:J{len(new_rk)}", "values": new_rk}],
            value_input_option="USER_ENTERED",
        )

        fetch_sheet_bundle_cached.clear()
        return True, f"LPJ ({id_lpj}) berhasil disimpan ke Riwayat Lokal & Google Sheets!"
    except Exception as e:
        return True, f"LPJ ({id_lpj}) tersimpan di Riwayat Lokal, namun gagal sinkron ke Cloud: {e}"


def delete_document_from_storage(doc_mode: str, doc_id: str, spreadsheet_id: str = "", creds_json_str: str = "") -> tuple:
    """
    PERBAIKAN MUTLAK V9.0 (Full CRUD): Menghapus dokumen RAB atau LPJ dari:
    1) st.session_state (local_db_rab / local_db_lpj & local_rekap_dana)
    2) File JSON lokal `database_riwayat_dwp_malang.json`
    3) Google Sheets (`DATABASE_RAB` / `DATABASE_LPJ` & `REKAP_ITEM_DANA`)
    """
    clean_id = (doc_id or "").strip()
    if not clean_id:
        return False, "ID Dokumen yang ingin dihapus kosong."

    deleted_ids = set(st.session_state.get("deleted_doc_ids", []))
    deleted_ids.add(clean_id)
    st.session_state["deleted_doc_ids"] = list(deleted_ids)

    if doc_mode == "RAB":
        local_rab = [r for r in st.session_state.get("local_db_rab", []) if r.get("ID_RAB") != clean_id]
        st.session_state["local_db_rab"] = local_rab
    else:
        local_lpj = [r for r in st.session_state.get("local_db_lpj", []) if r.get("ID_LPJ") != clean_id]
        st.session_state["local_db_lpj"] = local_lpj

    local_rk = [r for r in st.session_state.get("local_rekap_dana", []) if r.get("ID_DOKUMEN") != clean_id]
    st.session_state["local_rekap_dana"] = local_rk

    save_persistent_local_history(
        rab_rows=st.session_state.get("local_db_rab", []),
        lpj_rows=st.session_state.get("local_db_lpj", []),
        rekap_rows=local_rk,
    )

    target_sheet_id = spreadsheet_id or get_active_spreadsheet_id()
    client, _ = get_gspread_client(creds_json_str)
    if client is None:
        return True, f"Dokumen {doc_mode} ({clean_id}) berhasil dihapus dari Riwayat Lokal."

    try:
        sh = client.open_by_key(target_sheet_id)
        ws_rab, ws_lpj, ws_rk = ensure_worksheets_ready(sh)

        if doc_mode == "RAB":
            existing_vals = ws_rab.get_all_values()
            new_matrix = [COLS_DB_RAB]
            for r in existing_vals[1:]:
                if any(str(c).strip() for c in r) and (len(r) == 0 or r[0] != clean_id):
                    padded = list(r) + [""] * max(0, len(COLS_DB_RAB) - len(r))
                    new_matrix.append(padded[:len(COLS_DB_RAB)])
            ws_rab.clear()
            ws_rab.batch_update([{"range": f"A1:J{len(new_matrix)}", "values": new_matrix}], value_input_option="USER_ENTERED")
        else:
            existing_vals = ws_lpj.get_all_values()
            new_matrix = [COLS_DB_LPJ]
            for r in existing_vals[1:]:
                if any(str(c).strip() for c in r) and (len(r) == 0 or r[0] != clean_id):
                    padded = list(r) + [""] * max(0, len(COLS_DB_LPJ) - len(r))
                    new_matrix.append(padded[:len(COLS_DB_LPJ)])
            ws_lpj.clear()
            ws_lpj.batch_update([{"range": f"A1:O{len(new_matrix)}", "values": new_matrix}], value_input_option="USER_ENTERED")

        existing_rk = ws_rk.get_all_values()
        new_rk = [COLS_REKAP_DANA]
        for r in existing_rk[1:]:
            if any(str(c).strip() for c in r) and (len(r) == 0 or r[0] != clean_id):
                padded = list(r) + [""] * max(0, len(COLS_REKAP_DANA) - len(r))
                new_rk.append(padded[:len(COLS_REKAP_DANA)])
        ws_rk.clear()
        ws_rk.batch_update([{"range": f"A1:J{len(new_rk)}", "values": new_rk}], value_input_option="USER_ENTERED")

        fetch_sheet_bundle_cached.clear()
        return True, f"Dokumen {doc_mode} ({clean_id}) berhasil dihapus dari Riwayat Lokal & Google Sheets!"
    except Exception as e:
        return True, f"Dokumen {doc_mode} ({clean_id}) dihapus dari Lokal (Cloud warning: {e})"


def delete_rab_from_storage(id_rab: str, spreadsheet_id: str = "", creds_json_str: str = "") -> tuple:
    """Wrapper khusus untuk menghapus dokumen RAB berdasarkan ID_RAB."""
    return delete_document_from_storage("RAB", id_rab, spreadsheet_id, creds_json_str)


def delete_lpj_from_storage(id_lpj: str, spreadsheet_id: str = "", creds_json_str: str = "") -> tuple:
    """Wrapper khusus untuk menghapus dokumen LPJ berdasarkan ID_LPJ."""
    return delete_document_from_storage("LPJ", id_lpj, spreadsheet_id, creds_json_str)


def deduplicate_history_storage(spreadsheet_id: str = "", creds_json_str: str = "") -> tuple:
    """
    PERBAIKAN MUTLAK V9.0: Menghapus dokumen dobel (yang memiliki Nama Kegiatan, Lokasi/Tempat,
    dan Total Dana yang sama persis akibat klik simpan berulang kali), dengan menyisakan 1 dokumen terbaru.
    """
    rab_rows = st.session_state.get("local_db_rab", [])
    lpj_rows = st.session_state.get("local_db_lpj", [])

    removed_rab_ids = []
    seen_rab_sig = set()
    clean_rab = []
    for r in rab_rows:
        sig = (
            str(r.get("NAMA_KEGIATAN", "")).strip().upper(),
            str(r.get("LOKASI_ACARA", "")).strip().upper(),
            parse_num(r.get("TOTAL_RAB", 0)),
        )
        if sig in seen_rab_sig:
            removed_rab_ids.append(r.get("ID_RAB"))
        else:
            seen_rab_sig.add(sig)
            clean_rab.append(r)

    removed_lpj_ids = []
    seen_lpj_sig = set()
    clean_lpj = []
    for r in lpj_rows:
        sig = (
            str(r.get("NAMA_KEGIATAN", "")).strip().upper(),
            str(r.get("TEMPAT_ACARA", "")).strip().upper(),
            parse_num(r.get("TOTAL_DANA", 0)),
        )
        if sig in seen_lpj_sig:
            removed_lpj_ids.append(r.get("ID_LPJ"))
        else:
            seen_lpj_sig.add(sig)
            clean_lpj.append(r)

    kept_rab_ids = {r.get("ID_RAB") for r in clean_rab}
    kept_lpj_ids = {r.get("ID_LPJ") for r in clean_lpj}
    for rid in removed_rab_ids:
        if rid not in kept_rab_ids:
            delete_document_from_storage("RAB", rid, spreadsheet_id, creds_json_str)
    for lid in removed_lpj_ids:
        if lid not in kept_lpj_ids:
            delete_document_from_storage("LPJ", lid, spreadsheet_id, creds_json_str)
    st.session_state["local_db_rab"] = clean_rab
    st.session_state["local_db_lpj"] = clean_lpj
    save_persistent_local_history(
        rab_rows=clean_rab,
        lpj_rows=clean_lpj,
        rekap_rows=st.session_state.get("local_rekap_dana", []),
    )

    total_removed = len(removed_rab_ids) + len(removed_lpj_ids)
    if total_removed == 0:
        return True, "Tidak ditemukan data dobel. Seluruh dokumen di Riwayat sudah unik!"
    return True, f"Berhasil membersihkan {total_removed} dokumen dobel ({len(removed_rab_ids)} RAB & {len(removed_lpj_ids)} LPJ)!"


# ==============================================================================
# BAGIAN 4: ENGINE DOKUMEN WORD V9.0 (RP RATA KIRI SEJAJAR TOTAL & 2 PARAGRAF JABATAN)
# ==============================================================================

def find_template_file(filename: str) -> str:
    """Mencari file template .docx di direktori kerja atau subdirektori."""
    candidates = [
        filename,
        os.path.join(os.getcwd(), filename),
    ]
    for root_dir, _, files in os.walk(os.getcwd()):
        if filename in files:
            candidates.append(os.path.join(root_dir, filename))
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    return ""


def get_default_assets_bytes() -> dict:
    """Mengembalikan byte gambar Logo Kiri, Logo Kanan, dan Tanda Tangan Asli Ny. Rini Johan Pandores."""
    logo1_bytes = base64.b64decode(EMBEDDED_LOGO1_B64)
    logo2_bytes = base64.b64decode(EMBEDDED_LOGO2_B64)
    sig_bytes = base64.b64decode(EMBEDDED_SIG_B64)

    tpl_bytes = st.session_state.get("uploaded_template_bytes")
    tpl_path = find_template_file(TEMPLATE_LPJ_FILENAME)

    try:
        if tpl_bytes:
            doc_tmp = docx.Document(io.BytesIO(tpl_bytes))
        elif tpl_path:
            doc_tmp = docx.Document(tpl_path)
        else:
            doc_tmp = None

        if doc_tmp is not None:
            for _, part in doc_tmp.part.related_parts.items():
                if "image" in part.content_type:
                    sig_bytes = part.blob
    except Exception:
        pass

    return {
        "logo1": logo1_bytes,
        "logo2": logo2_bytes,
        "sig": sig_bytes,
    }


def _set_cell_margins(cell, top=42, bottom=42, left=65, right=65):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m_name, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m_name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_cell_shading(cell, hex_color: str):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _set_table_borders(table, color="505050", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f"</w:tblBorders>"
    )
    tblPr.append(borders)


def _build_fallback_header(section, logo1_bytes: bytes, logo2_bytes: bytes):
    """Membangun Kop Surat Resmi DWP KPPBC TMC Malang secara terprogram jika file template tidak ditemukan."""
    header = section.header
    header.is_linked_to_previous = False
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)

    htbl = header.add_table(rows=1, cols=3, width=Cm(19.5))
    htbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(2.5), Cm(14.2), Cm(2.8)]
    for i, w in enumerate(widths):
        htbl.columns[i].width = w
        htbl.rows[0].cells[i].width = w
        htbl.rows[0].cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p_left = htbl.rows[0].cells[0].paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left.add_run().add_picture(io.BytesIO(logo1_bytes), width=Cm(1.95))

    c_mid = htbl.rows[0].cells[1]
    p1 = c_mid.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(1)
    r1 = p1.add_run("Dharma Wanita Persatuan")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.italic = True

    p2 = c_mid.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(1)
    r2 = p2.add_run(
        "DIREKTORAT JENDERAL BEA DAN CUKAI KANTOR WILAYAH DJBC JAWA TIMUR II" + NL +
        "KANTOR PENGAWASAN DAN PELAYANAN BEA DAN CUKAI TIPE MADYA CUKAI MALANG"
    )
    r2.font.name = "Arial"
    r2.font.size = Pt(9)
    r2.font.bold = True

    p3 = c_mid.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(0)
    r3 = p3.add_run(
        "JALAN SURABAYA NOMOR 2, MALANG 65145" + NL +
        "TELEPON (0341) 551628; FAKSIMILE (0341)554193; LAMAN WWW.BEACUKAIMALANG.COM" + NL +
        "PUSAT KONTAK LAYANAN 1500225; SUREL INFO@CUSTOMS.GO.ID"
    )
    r3.font.name = "Calibri"
    r3.font.size = Pt(6.5)

    p_right = htbl.rows[0].cells[2].paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right.add_run().add_picture(io.BytesIO(logo2_bytes), width=Cm(2.25))

    p_line = header.add_paragraph()
    p_line.paragraph_format.space_before = Pt(2)
    p_line.paragraph_format.space_after = Pt(4)
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        '<w:bottom w:val="single" w:sz="12" w:space="1" w:color="000000"/>'
        "</w:pBdr>"
    )
    p_line._p.get_or_add_pPr().append(pBdr)


def _prepare_cloned_document():
    """Membuka dokumen dengan Kop Surat Resmi DWP KPPBC TMC Malang yang di-clone 100% utuh."""
    assets = get_default_assets_bytes()
    tpl_bytes = st.session_state.get("uploaded_template_bytes")
    tpl_path = find_template_file(TEMPLATE_LPJ_FILENAME) or find_template_file(TEMPLATE_RAB_FILENAME)

    used_cloned = False
    if tpl_bytes:
        doc = docx.Document(io.BytesIO(tpl_bytes))
        used_cloned = True
    elif tpl_path:
        doc = docx.Document(tpl_path)
        used_cloned = True
    else:
        doc = docx.Document()

    if used_cloned:
        hdr_ref = None
        for p in doc.paragraphs:
            pPr = p._p.find(qn("w:pPr"))
            if pPr is not None:
                sPr = pPr.find(qn("w:sectPr"))
                if sPr is not None:
                    href = sPr.find(qn("w:headerReference"))
                    if href is not None:
                        hdr_ref = href

        body_el = doc._body._body
        final_sectPr = body_el.find(qn("w:sectPr"))
        if hdr_ref is not None and final_sectPr is not None:
            if final_sectPr.find(qn("w:headerReference")) is None:
                final_sectPr.insert(0, hdr_ref)

        for child in list(body_el):
            if child.tag != qn("w:sectPr"):
                body_el.remove(child)
    else:
        sec = doc.sections[0]
        sec.page_width = Cm(21.59)
        sec.page_height = Cm(27.94)
        sec.top_margin = Cm(3.85)
        sec.bottom_margin = Cm(1.20)
        sec.left_margin = Cm(0.88)
        sec.right_margin = Cm(0.50)
        sec.header_distance = Cm(1.28)
        _build_fallback_header(sec, assets["logo1"], assets["logo2"])

    for sec in doc.sections:
        if sec.bottom_margin < Cm(1.0):
            sec.bottom_margin = Cm(1.0)

    return doc, assets


def _format_currency_cell_split(cell, val, col_width_cm: float, bold: bool = False, size_pt: float = 10.5, is_dash_if_zero: bool = True):
    """
    Format Akuntansi Sejajar Lurus di Word (.docx):
    - Paragraf di dalam sel diatur Rata Kiri (WD_ALIGN_PARAGRAPH.LEFT) + <w:noWrap/>
    - Tambahkan 1 Tab-Stop Rata Kanan (WD_TAB_ALIGNMENT.RIGHT) tepat di lebar bersih sel: Cm(max(1.8, col_width_cm - 0.35))
    - Isi teks paragraf dengan "Rp" + TAB + fmt_nominal_only(val, is_dash_if_zero)
    """
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_margins(cell, top=42, bottom=42, left=65, right=65)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(OxmlElement("w:noWrap"))

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.05
    pf.tab_stops.clear_all()
    pf.tab_stops.add_tab_stop(Cm(max(1.8, col_width_cm - 0.35)), WD_TAB_ALIGNMENT.RIGHT)

    for r_old in list(p.runs):
        p._p.remove(r_old._r)

    r = p.add_run("Rp" + TAB + fmt_nominal_only(val, is_dash_if_zero=is_dash_if_zero))
    r.font.name = "Arial"
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def _build_6col_financial_table(doc, df_items: pd.DataFrame, merge_ket: bool, ket_global: str):
    """Membangun Tabel 6 Kolom Presisi dengan Format Akuntansi Split-Alignment (Rp Rata Kiri, Angka Rata Kanan)."""
    n_items = len(df_items)
    total_rows = n_items + 3
    col_widths_cm = [1.05, 5.80, 1.85, 2.95, 3.10, 4.45]

    table = doc.add_table(rows=total_rows, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table, color="505050", sz="4")

    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    tblPr.append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="400" w:type="dxa"/>'))

    for col_idx, w_cm in enumerate(col_widths_cm):
        table.columns[col_idx].width = Cm(w_cm)

    def _format_cell_text(cell, text: str, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size_pt=10.5, no_wrap=False):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(cell, top=42, bottom=42, left=65, right=65)
        if no_wrap:
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(OxmlElement("w:noWrap"))
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        for r_old in list(p.runs):
            p._p.remove(r_old._r)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    r0 = table.rows[0]
    r1 = table.rows[1]
    for h_row in (r0, r1):
        trPr = h_row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
        trPr.append(OxmlElement("w:cantSplit"))

    r0.cells[3].merge(r0.cells[4])
    for c_idx in (0, 1, 2, 5):
        r0.cells[c_idx].merge(r1.cells[c_idx])

    _format_cell_text(r0.cells[0], "NO", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.5, no_wrap=True)
    _format_cell_text(r0.cells[1], "KEBUTUHAN KEGIATAN", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.5)
    _format_cell_text(r0.cells[2], "JUMLAH", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.5, no_wrap=True)
    _format_cell_text(r0.cells[3], "DANA", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.5)
    _format_cell_text(r0.cells[5], "KETERANGAN", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.5)

    _format_cell_text(r1.cells[3], "HARGA SATUAN", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.0, no_wrap=True)
    _format_cell_text(r1.cells[4], "HARGA TOTAL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.0, no_wrap=True)

    for h_r in (r0, r1):
        for c_idx, w_cm in enumerate(col_widths_cm):
            h_r.cells[c_idx].width = Cm(w_cm)
            _set_cell_shading(h_r.cells[c_idx], "F8FAFC")

    for idx_df, row_data in df_items.iterrows():
        row_pos = 2 + idx_df
        tr_row = table.rows[row_pos]
        tr_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

        for c_idx, w_cm in enumerate(col_widths_cm):
            tr_row.cells[c_idx].width = Cm(w_cm)

        no_str = str(int(row_data["NO"]))
        keb_str = str(row_data["KEBUTUHAN KEGIATAN"])
        jml_str = str(int(row_data["JUMLAH"])) if int(row_data["JUMLAH"]) > 0 else "-"
        ket_item = str(row_data.get("KETERANGAN", ket_global) or "")

        _format_cell_text(tr_row.cells[0], no_str, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size_pt=10.5)
        _format_cell_text(tr_row.cells[1], keb_str, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size_pt=10.5)
        _format_cell_text(tr_row.cells[2], jml_str, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size_pt=10.5)

        _format_currency_cell_split(tr_row.cells[3], row_data["HARGA SATUAN"], col_widths_cm[3], bold=False, size_pt=10.5, is_dash_if_zero=True)
        _format_currency_cell_split(tr_row.cells[4], row_data["HARGA TOTAL"], col_widths_cm[4], bold=False, size_pt=10.5, is_dash_if_zero=True)

        if not merge_ket:
            _format_cell_text(tr_row.cells[5], ket_item, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size_pt=10.0)

    if merge_ket and n_items > 0:
        first_ket_cell = table.rows[2].cells[5]
        if n_items > 1:
            last_ket_cell = table.rows[2 + n_items - 1].cells[5]
            first_ket_cell.merge(last_ket_cell)
        _format_cell_text(first_ket_cell, ket_global, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size_pt=10.5)

    r_tot = table.rows[total_rows - 1]
    r_tot._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for c_idx, w_cm in enumerate(col_widths_cm):
        r_tot.cells[c_idx].width = Cm(w_cm)

    r_tot.cells[0].merge(r_tot.cells[3])
    total_keseluruhan = int(df_items["HARGA TOTAL"].sum()) if not df_items.empty else 0

    _format_cell_text(r_tot.cells[0], "TOTAL", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=10.5)
    _format_currency_cell_split(r_tot.cells[4], total_keseluruhan, col_widths_cm[4], bold=True, size_pt=10.5, is_dash_if_zero=False)
    _format_cell_text(r_tot.cells[5], "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size_pt=10.5)


def generate_rab_docx_bytes(rab_payload: dict) -> bytes:
    """Mesin pembuat dokumen Word (.docx) khusus RAB sesuai pakem resmi RAB DWP GABUNGAN DI JEMBER.docx."""
    doc, _ = _prepare_cloned_document()

    def _add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, left_cm=0.73, right_cm=0.80, first_line_cm=0.0, before_pt=0, after_pt=3, bold=False, italic=False, size_pt=11.0):
        p = doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.left_indent = Cm(left_cm)
        pf.right_indent = Cm(right_cm)
        pf.first_line_indent = Cm(first_line_cm)
        pf.space_before = Pt(before_pt)
        pf.space_after = Pt(after_pt)
        pf.line_spacing = 1.15
        pf.keep_together = True
        if text:
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(size_pt)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        return p

    b1 = (rab_payload.get("baris_1_judul") or "RANCANGAN ANGGARAN BIAYA").strip().upper()
    b2 = (rab_payload.get("nama_kegiatan_rab") or "PELAKSANAAN PERTEMUAN RUTIN PAGUYUBAN DWP").strip().upper()
    b3 = (rab_payload.get("lokasi_acara_rab") or "").strip().upper()
    if b3 and not b3.startswith("("):
        b3 = f"({b3})"

    p_t1 = _add_p(b1, align=WD_ALIGN_PARAGRAPH.CENTER, left_cm=0.73, right_cm=0.80, before_pt=10, after_pt=1, bold=True, size_pt=11.0)
    p_t1.paragraph_format.keep_with_next = True

    p_t2 = _add_p(b2, align=WD_ALIGN_PARAGRAPH.CENTER, left_cm=0.73, right_cm=0.80, before_pt=0, after_pt=1, bold=True, size_pt=11.0)
    p_t2.paragraph_format.keep_with_next = True

    if b3:
        p_t3 = _add_p(b3, align=WD_ALIGN_PARAGRAPH.CENTER, left_cm=0.73, right_cm=0.80, before_pt=0, after_pt=12, bold=True, size_pt=11.0)
        p_t3.paragraph_format.keep_with_next = True

    sumber = (rab_payload.get("sumber_dana_rab") or "Kas Dharma Wanita KPPBC TMC Malang").strip()
    p_pengantar = _add_p(
        f"Dana untuk kegiatan ini berasal dari {sumber}. Dengan rincian dana sebagai berikut :",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left_cm=0.73,
        right_cm=0.80,
        first_line_cm=1.25,
        before_pt=4,
        after_pt=8,
        size_pt=11.0,
    )
    p_pengantar.paragraph_format.keep_with_next = True

    df_items = normalize_items_df(rab_payload.get("items_rab", []), rab_payload.get("ket_tabel_rab", ""))
    _build_6col_financial_table(
        doc,
        df_items=df_items,
        merge_ket=bool(rab_payload.get("merge_ket_rab", False)),
        ket_global=str(rab_payload.get("ket_tabel_rab", "") or ""),
    )

    note_raw = (rab_payload.get("catatan_note_rab") or "").strip()
    note_lines = [ln.strip() for ln in note_raw.split(NL) if ln.strip()]
    if note_lines:
        p_note_hdr = _add_p(
            "*Note :",
            align=WD_ALIGN_PARAGRAPH.LEFT,
            left_cm=0.73,
            right_cm=0.80,
            before_pt=10,
            after_pt=3,
            italic=True,
            size_pt=10.0,
        )
        p_note_hdr.paragraph_format.keep_with_next = True

        for ln in note_lines:
            clean_ln = ln.lstrip("•-* ").strip()
            p_b = doc.add_paragraph()
            p_b.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf_b = p_b.paragraph_format
            pf_b.left_indent = Cm(1.55)
            pf_b.right_indent = Cm(0.80)
            pf_b.first_line_indent = Cm(-0.50)
            pf_b.space_before = Pt(1)
            pf_b.space_after = Pt(2)
            pf_b.line_spacing = 1.12
            pf_b.tab_stops.add_tab_stop(Cm(1.55), WD_TAB_ALIGNMENT.LEFT)
            r_b = p_b.add_run(BULLET_CHAR + TAB + clean_ln)
            r_b.font.name = "Arial"
            r_b.font.size = Pt(10.0)

    out_buf = io.BytesIO()
    doc.save(out_buf)
    return out_buf.getvalue()


def generate_lpj_docx_bytes(payload: dict, custom_sig_bytes: bytes = None, attachment_images: list = None) -> bytes:
    """Mesin utama pembuat dokumen Word (.docx) LPJ sesuai pakem resmi LPJ_DWP GABUNGAN DI BLITAR_2026-1.docx."""
    doc, assets = _prepare_cloned_document()
    sig_to_use = custom_sig_bytes if custom_sig_bytes else assets["sig"]

    def _add_styled_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, left_cm=0.73, right_cm=0.80, first_line_cm=0.0, space_before_pt=0, space_after_pt=3, italic=False, bold=False, size_pt=11.0):
        p = doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.left_indent = Cm(left_cm)
        pf.right_indent = Cm(right_cm)
        pf.first_line_indent = Cm(first_line_cm)
        pf.space_before = Pt(space_before_pt)
        pf.space_after = Pt(space_after_pt)
        pf.line_spacing = 1.15
        pf.keep_together = True
        if text:
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(size_pt)
            r.font.italic = italic
            r.font.bold = bold
            r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        return p

    tgl_laporan_indo = format_hari_tanggal_indo(payload.get("tgl_laporan", date.today()))
    _add_styled_p(
        tgl_laporan_indo,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        left_cm=0.73,
        right_cm=1.32,
        space_before_pt=6,
        space_after_pt=10,
        size_pt=11.0,
    )

    _add_styled_p(
        "Assalamu’alaikum Wr. Wb.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        left_cm=0.73,
        right_cm=0.80,
        space_before_pt=2,
        space_after_pt=8,
        italic=True,
        size_pt=11.0,
    )

    teks_pembuka = build_paragraf_pembuka(
        payload.get("nama_kegiatan", ""),
        payload.get("tema_kegiatan", ""),
        payload.get("konteks_hut", ""),
    )
    _add_styled_p(
        teks_pembuka,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left_cm=0.73,
        right_cm=0.80,
        first_line_cm=1.27,
        space_before_pt=0,
        space_after_pt=6,
        size_pt=11.0,
    )

    tgl_acara_indo = (payload.get("tgl_acara_manual") or "").strip()
    if not tgl_acara_indo:
        tgl_acara_indo = format_hari_tanggal_indo(payload.get("tgl_acara", date.today()))

    waktu_acara = (payload.get("waktu_acara") or "08.30 - selesai").strip()
    tempat_acara = (payload.get("tempat_acara") or "-").strip()

    for label_kiri, nilai_kanan in [
        ("Hari, Tanggal", tgl_acara_indo),
        ("Waktu", waktu_acara),
        ("Tempat", tempat_acara),
    ]:
        p_info = _add_styled_p(
            "",
            align=WD_ALIGN_PARAGRAPH.LEFT,
            left_cm=1.98,
            right_cm=0.80,
            space_before_pt=1,
            space_after_pt=2,
            size_pt=11.0,
        )
        p_info.paragraph_format.tab_stops.add_tab_stop(Cm(4.85), WD_TAB_ALIGNMENT.LEFT)
        r_lbl = p_info.add_run(label_kiri + TAB + ": " + nilai_kanan)
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(11.0)

    _add_styled_p(
        "Dengan rangkaian kegiatan sebagai berikut :",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        left_cm=1.98,
        right_cm=0.80,
        space_before_pt=6,
        space_after_pt=3,
        size_pt=11.0,
    )

    rangkaian_raw = (payload.get("rangkaian_kegiatan") or "").strip()
    rangkaian_list = [ln.strip() for ln in rangkaian_raw.split(NL) if ln.strip()]
    for idx_r, item_r in enumerate(rangkaian_list, start=1):
        clean_r = item_r
        if "." in clean_r[:4]:
            l_num, r_rest = clean_r.split(".", 1)
            if l_num.strip().isdigit():
                clean_r = r_rest.strip()

        p_r = doc.add_paragraph()
        p_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf_r = p_r.paragraph_format
        pf_r.left_indent = Cm(2.61)
        pf_r.right_indent = Cm(0.80)
        pf_r.first_line_indent = Cm(-0.63)
        pf_r.space_before = Pt(1)
        pf_r.space_after = Pt(2)
        pf_r.line_spacing = 1.12
        pf_r.tab_stops.add_tab_stop(Cm(2.61), WD_TAB_ALIGNMENT.LEFT)
        run_r = p_r.add_run(f"{idx_r}." + TAB + clean_r)
        run_r.font.name = "Arial"
        run_r.font.size = Pt(11.0)

    sumber_dana = (payload.get("sumber_dana") or "Kas Dharma Wanita KPPBC TMC Malang dan donatur").strip()
    _add_styled_p(
        f"Dana untuk kegiatan ini berasal dari {sumber_dana}. Dengan rincian dana sebagai berikut :",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left_cm=0.73,
        right_cm=0.80,
        first_line_cm=1.27,
        space_before_pt=6,
        space_after_pt=6,
        size_pt=11.0,
    )

    df_items = normalize_items_df(payload.get("items_dana", []), payload.get("ket_tabel_dana", ""))
    _build_6col_financial_table(
        doc,
        df_items=df_items,
        merge_ket=bool(payload.get("merge_ket_vertikal", True)),
        ket_global=str(payload.get("ket_tabel_dana", "Kas DWP BC Malang & Donatur Anggota") or ""),
    )

    if payload.get("tampilkan_ucapan_donatur", True):
        teks_donatur = (payload.get("teks_ucapan_donatur") or DEFAULT_UCAPAN_DONATUR).strip()
        if teks_donatur:
            p_don = _add_styled_p(
                teks_donatur,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                left_cm=0.73,
                right_cm=0.80,
                first_line_cm=1.25,
                space_before_pt=10,
                space_after_pt=6,
                size_pt=11.0,
            )
            p_don.paragraph_format.keep_with_next = True

    nama_keg = (payload.get("nama_kegiatan") or "Kegiatan Dharma Wanita Persatuan").strip()
    teks_penutup = (
        f"Demikian laporan pertanggung jawaban {nama_keg} ini kami sampaikan, "
        "atas perhatian dan bantuan yang telah diberikan kami ucapkan terima kasih."
    )
    _add_styled_p(
        teks_penutup,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        left_cm=0.73,
        right_cm=0.80,
        first_line_cm=1.25,
        space_before_pt=4,
        space_after_pt=6,
        size_pt=11.0,
    )

    _add_styled_p(
        "Wassalamu'alaikum Wr. Wb.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        left_cm=0.73,
        right_cm=0.80,
        space_before_pt=2,
        space_after_pt=6,
        italic=True,
        size_pt=11.0,
    )

    catatan_ig = (payload.get("catatan_ig") or "").strip()
    if catatan_ig:
        p_ig = _add_styled_p(
            "",
            align=WD_ALIGN_PARAGRAPH.LEFT,
            left_cm=0.73,
            right_cm=0.80,
            space_before_pt=4,
            space_after_pt=8,
            size_pt=10.0,
        )
        if "@dwpbcmalang" in catatan_ig:
            parts_ig = catatan_ig.split("@dwpbcmalang", 1)
            r_ig1 = p_ig.add_run(parts_ig[0])
            r_ig1.font.name = "Arial"
            r_ig1.font.size = Pt(10.0)
            r_ig1.font.italic = True
            r_ig2 = p_ig.add_run("@dwpbcmalang")
            r_ig2.font.name = "Arial"
            r_ig2.font.size = Pt(10.0)
            r_ig2.font.bold = True
            if parts_ig[1]:
                r_ig3 = p_ig.add_run(parts_ig[1])
                r_ig3.font.name = "Arial"
                r_ig3.font.size = Pt(10.0)
        else:
            r_ig = p_ig.add_run(catatan_ig)
            r_ig.font.name = "Arial"
            r_ig.font.size = Pt(10.0)
            r_ig.font.italic = True

    jab_1 = (payload.get("jabatan_baris_1") or "Ketua Dharma Wanita Persatuan").strip()
    jab_2 = (payload.get("jabatan_baris_2") or "KPPBC Tipe Madya Cukai Malang").strip()
    nama_ttd = (payload.get("nama_penandatangan") or "Ny. Rini Johan Pandores").strip()

    p_jab_1 = _add_styled_p(
        jab_1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        left_cm=12.14,
        right_cm=0.80,
        space_before_pt=6,
        space_after_pt=0,
        size_pt=11.0,
    )
    p_jab_1.paragraph_format.keep_with_next = True

    p_jab_2 = _add_styled_p(
        jab_2,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        left_cm=12.14,
        right_cm=0.80,
        space_before_pt=0,
        space_after_pt=2,
        size_pt=11.0,
    )
    p_jab_2.paragraph_format.keep_with_next = True

    p_sig_img = doc.add_paragraph()
    p_sig_img.paragraph_format.left_indent = Cm(12.14)
    p_sig_img.paragraph_format.space_before = Pt(2)
    p_sig_img.paragraph_format.space_after = Pt(2)
    p_sig_img.paragraph_format.keep_with_next = True

    if payload.get("tampilkan_ttd", True) and sig_to_use:
        try:
            r_sig = p_sig_img.add_run()
            r_sig.add_picture(io.BytesIO(sig_to_use), width=Cm(3.35))
        except Exception:
            p_sig_img.paragraph_format.space_before = Pt(18)
            p_sig_img.paragraph_format.space_after = Pt(18)
    else:
        p_sig_img.paragraph_format.space_before = Pt(20)
        p_sig_img.paragraph_format.space_after = Pt(20)

    _add_styled_p(
        nama_ttd,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        left_cm=12.14,
        right_cm=0.80,
        space_before_pt=2,
        space_after_pt=4,
        size_pt=11.0,
    )

    if attachment_images:
        doc.add_page_break()
        p_lamp = _add_styled_p(
            "LAMPIRAN DOKUMENTASI & BUKTI KWITANSI KEGIATAN",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            left_cm=0.73,
            right_cm=0.80,
            space_before_pt=6,
            space_after_pt=10,
            bold=True,
            size_pt=12.0,
        )
        p_lamp.paragraph_format.keep_with_next = True

        n_imgs = len(attachment_images)
        grid_rows = (n_imgs + 1) // 2
        gtbl = doc.add_table(rows=grid_rows, cols=2)
        gtbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(gtbl, color="CBD5E1", sz="4")

        for idx_img, img_item in enumerate(attachment_images):
            r_i = idx_img // 2
            c_i = idx_img % 2
            cell = gtbl.rows[r_i].cells[c_i]
            cell.width = Cm(9.2)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
            p_c = cell.paragraphs[0]
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p_c.add_run().add_picture(io.BytesIO(img_item["bytes"]), width=Cm(8.2))
                p_cap = cell.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cap = p_cap.add_run(img_item.get("caption", f"Dokumentasi {idx_img + 1}"))
                r_cap.font.name = "Arial"
                r_cap.font.size = Pt(9.5)
                r_cap.font.italic = True
            except Exception:
                p_c.add_run("[Gambar tidak dapat dimuat]")

    out_buf = io.BytesIO()
    doc.save(out_buf)
    return out_buf.getvalue()


@st.cache_data(show_spinner=False)
def convert_docx_bytes_to_pdf_bytes(docx_bytes: bytes) -> tuple:
    """Mengonversi byte .docx menjadi byte .pdf menggunakan LibreOffice Headless (Cached Anti-Lag)."""
    lo_bin = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo_bin:
        return None, "LibreOffice Headless tidak ditemukan di server. Gunakan tombol Cetak PDF Browser di bawah."

    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            in_docx = os.path.join(tmpdir, "DOKUMEN_DWP.docx")
            out_pdf = os.path.join(tmpdir, "DOKUMEN_DWP.pdf")
            with open(in_docx, "wb") as f:
                f.write(docx_bytes)

            lo_profile = os.path.join(tmpdir, "lo_profile")
            cmd = [
                lo_bin,
                f"-env:UserInstallation=file://{lo_profile}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmpdir,
                in_docx,
            ]
            proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=45)
            if os.path.isfile(out_pdf) and os.path.getsize(out_pdf) > 0:
                with open(out_pdf, "rb") as f_pdf:
                    return f_pdf.read(), "PDF berhasil dibuat dengan LibreOffice Headless."
            err_txt = proc.stderr.decode("utf-8", errors="ignore")
            return None, f"Konversi PDF gagal: {err_txt}"
    except Exception as e:
        return None, f"Kendala konversi PDF: {e}"


# ==============================================================================
# BAGIAN 5: DIRECT DOWNLOAD BASE64 & LIVE PREVIEW KERTAS A4 V9.0 (SPLIT RP)
# ==============================================================================

def make_direct_download_html(file_bytes: bytes, filename: str, mime_type: str, label: str, bg_gradient: str) -> str:
    """Membuat tombol unduh langsung berbasis Base64 Data URI (Anti-Error 404 Streamlit)."""
    b64_data = base64.b64encode(file_bytes).decode("ascii")
    return (
        f'<a href="data:{mime_type};base64,{b64_data}" download="{filename}" '
        f'style="display:flex;align-items:center;justify-content:center;gap:8px;'
        f'width:100%;min-height:48px;padding:10px 16px;border-radius:10px;'
        f'background:{bg_gradient};color:#FFFFFF !important;'
        f'-webkit-text-fill-color:#FFFFFF !important;font-weight:700;font-size:14.5px;'
        f'text-decoration:none;box-shadow:0 4px 12px rgba(15,23,42,0.14);'
        f'border:1px solid rgba(255,255,255,0.25);text-align:center;">'
        f"{label}</a>"
    )


def _build_kop_html(logo1_b64: str, logo2_b64: str) -> str:
    return f"""
    <table class="kop-table">
        <tr>
            <td style="width:78px;text-align:center;">
                <img src="data:image/jpeg;base64,{logo1_b64}" style="width:68px;height:auto;"/>
            </td>
            <td style="text-align:center;">
                <div style="font-family:'Times New Roman',serif;font-size:20px;font-weight:700;font-style:italic;letter-spacing:0.2px;">
                    Dharma Wanita Persatuan
                </div>
                <div style="font-size:11px;font-weight:700;line-height:1.25;margin-top:2px;">
                    DIREKTORAT JENDERAL BEA DAN CUKAI KANTOR WILAYAH DJBC JAWA TIMUR II<br/>
                    KANTOR PENGAWASAN DAN PELAYANAN BEA DAN CUKAI TIPE MADYA CUKAI MALANG
                </div>
                <div style="font-size:8.5px;margin-top:3px;line-height:1.25;">
                    JALAN SURABAYA NOMOR 2, MALANG 65145<br/>
                    TELEPON (0341) 551628; FAKSIMILE (0341)554193; LAMAN <span style="color:#1D4ED8;text-decoration:underline;">WWW.BEACUKAIMALANG.COM</span><br/>
                    PUSAT KONTAK LAYANAN 1500225; SUREL <span style="color:#1D4ED8;text-decoration:underline;">INFO@CUSTOMS.GO.ID</span>
                </div>
            </td>
            <td style="width:88px;text-align:center;">
                <img src="data:image/png;base64,{logo2_b64}" style="width:78px;height:auto;"/>
            </td>
        </tr>
    </table>
    """


def build_a4_live_preview_html(doc_mode: str, payload: dict, custom_sig_bytes: bytes = None) -> str:
    """Membangun pratinjau visual HTML Kertas A4 untuk RAB maupun LPJ dengan `fmt_rp_html_split`."""
    assets = get_default_assets_bytes()
    logo1_b64 = base64.b64encode(assets["logo1"]).decode("ascii")
    logo2_b64 = base64.b64encode(assets["logo2"]).decode("ascii")
    kop_html = _build_kop_html(logo1_b64, logo2_b64)

    if doc_mode == "RAB":
        b1 = (payload.get("baris_1_judul") or "RANCANGAN ANGGARAN BIAYA").strip().upper()
        b2 = (payload.get("nama_kegiatan_rab") or "").strip().upper()
        b3 = (payload.get("lokasi_acara_rab") or "").strip().upper()
        if b3 and not b3.startswith("("):
            b3 = f"({b3})"

        sumber = (payload.get("sumber_dana_rab") or "Kas Dharma Wanita KPPBC TMC Malang").strip()
        df_items = normalize_items_df(payload.get("items_rab", []), payload.get("ket_tabel_rab", ""))
        n_items = len(df_items)
        merge_ket = bool(payload.get("merge_ket_rab", False))
        ket_global = (payload.get("ket_tabel_rab") or "").strip()
        total_dana = int(df_items["HARGA TOTAL"].sum()) if not df_items.empty else 0

        rows_html = ""
        for idx_df, r in df_items.iterrows():
            no_str = str(int(r["NO"]))
            keb_str = str(r["KEBUTUHAN KEGIATAN"])
            jml_str = str(int(r["JUMLAH"])) if int(r["JUMLAH"]) > 0 else "-"
            hsat_html = fmt_rp_html_split(r["HARGA SATUAN"], is_dash_if_zero=True, bold=False)
            htot_html = fmt_rp_html_split(r["HARGA TOTAL"], is_dash_if_zero=True, bold=False)
            ket_item = str(r.get("KETERANGAN", ket_global) or "")

            if merge_ket:
                if idx_df == 0:
                    ket_td = f"<td rowspan='{n_items}' style='text-align:center;vertical-align:middle;padding:6px;'>{ket_global}</td>"
                else:
                    ket_td = ""
            else:
                ket_td = f"<td style='text-align:center;vertical-align:middle;padding:4px 6px;'>{ket_item}</td>"

            rows_html += (
                "<tr>"
                f"<td style='text-align:center;'>{no_str}</td>"
                f"<td style='text-align:left;'>{keb_str}</td>"
                f"<td style='text-align:center;'>{jml_str}</td>"
                f"<td>{hsat_html}</td>"
                f"<td>{htot_html}</td>"
                f"{ket_td}"
                "</tr>"
            )

        note_raw = (payload.get("catatan_note_rab") or "").strip()
        note_lines = [ln.strip() for ln in note_raw.split(NL) if ln.strip()]
        note_html = ""
        if note_lines:
            li_items = "".join(f"<li style='margin-bottom:4px;'>{ln.lstrip('•-* ').strip()}</li>" for ln in note_lines)
            note_html = (
                "<div style='margin:14px 24px 4px 24px;font-style:italic;font-size:12.5px;'>*Note :</div>"
                f"<ul style='margin:2px 24px 12px 44px;font-size:12.5px;'>{li_items}</ul>"
            )

        body_content = f"""
            <div style="text-align:center;font-weight:700;font-size:13.8px;line-height:1.4;margin:14px 24px 14px 24px;">
                <div>{b1}</div>
                <div>{b2}</div>
                <div>{b3}</div>
            </div>
            <p class="lpj-p-indent">
                Dana untuk kegiatan ini berasal dari {sumber}. Dengan rincian dana sebagai berikut :
            </p>
            <table class="dana-table">
                <thead>
                    <tr>
                        <th rowspan="2" style="width:6%;">NO</th>
                        <th rowspan="2" style="width:31%;">KEBUTUHAN KEGIATAN</th>
                        <th rowspan="2" style="width:9%;">JUMLAH</th>
                        <th colspan="2" style="width:32%;">DANA</th>
                        <th rowspan="2" style="width:22%;">KETERANGAN</th>
                    </tr>
                    <tr>
                        <th style="width:16%;">HARGA SATUAN</th>
                        <th style="width:16%;">HARGA TOTAL</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                    <tr>
                        <td colspan="4" style="text-align:center;font-weight:700;background:#F8FAFC;">TOTAL</td>
                        <td style="background:#F8FAFC;">{fmt_rp_html_split(total_dana, is_dash_if_zero=False, bold=True)}</td>
                        <td style="background:#F8FAFC;"></td>
                    </tr>
                </tbody>
            </table>
            {note_html}
        """
    else:
        sig_bytes = custom_sig_bytes if custom_sig_bytes else assets["sig"]
        sig_b64 = base64.b64encode(sig_bytes).decode("ascii") if sig_bytes else ""

        tgl_laporan_indo = format_hari_tanggal_indo(payload.get("tgl_laporan", date.today()))
        tgl_acara_indo = (payload.get("tgl_acara_manual") or "").strip() or format_hari_tanggal_indo(payload.get("tgl_acara", date.today()))
        waktu_acara = (payload.get("waktu_acara") or "08.30 - selesai").strip()
        tempat_acara = (payload.get("tempat_acara") or "-").strip()
        teks_pembuka = build_paragraf_pembuka(
            payload.get("nama_kegiatan", ""),
            payload.get("tema_kegiatan", ""),
            payload.get("konteks_hut", ""),
        )

        rangkaian_raw = (payload.get("rangkaian_kegiatan") or "").strip()
        rangkaian_items = [ln.strip() for ln in rangkaian_raw.split(NL) if ln.strip()]
        ol_html = ""
        for item_r in rangkaian_items:
            clean_r = item_r
            if "." in clean_r[:4]:
                l_num, r_rest = clean_r.split(".", 1)
                if l_num.strip().isdigit():
                    clean_r = r_rest.strip()
            ol_html += f"<li style='margin-bottom:3px;'>{clean_r}</li>"

        df_items = normalize_items_df(payload.get("items_dana", []), payload.get("ket_tabel_dana", ""))
        n_items = len(df_items)
        merge_ket = bool(payload.get("merge_ket_vertikal", True))
        ket_global = (payload.get("ket_tabel_dana") or "Kas DWP BC Malang & Donatur Anggota").strip()
        total_dana = int(df_items["HARGA TOTAL"].sum()) if not df_items.empty else 0

        rows_html = ""
        for idx_df, r in df_items.iterrows():
            no_str = str(int(r["NO"]))
            keb_str = str(r["KEBUTUHAN KEGIATAN"])
            jml_str = str(int(r["JUMLAH"])) if int(r["JUMLAH"]) > 0 else "-"
            hsat_html = fmt_rp_html_split(r["HARGA SATUAN"], is_dash_if_zero=True, bold=False)
            htot_html = fmt_rp_html_split(r["HARGA TOTAL"], is_dash_if_zero=True, bold=False)
            ket_item = str(r.get("KETERANGAN", ket_global) or ket_global)

            if merge_ket:
                if idx_df == 0:
                    ket_td = f"<td rowspan='{n_items}' style='text-align:center;vertical-align:middle;padding:6px;'>{ket_global}</td>"
                else:
                    ket_td = ""
            else:
                ket_td = f"<td style='text-align:center;vertical-align:middle;padding:4px 6px;'>{ket_item}</td>"

            rows_html += (
                "<tr>"
                f"<td style='text-align:center;'>{no_str}</td>"
                f"<td style='text-align:left;'>{keb_str}</td>"
                f"<td style='text-align:center;'>{jml_str}</td>"
                f"<td>{hsat_html}</td>"
                f"<td>{htot_html}</td>"
                f"{ket_td}"
                "</tr>"
            )

        ucapan_html = ""
        if payload.get("tampilkan_ucapan_donatur", True):
            teks_don = (payload.get("teks_ucapan_donatur") or DEFAULT_UCAPAN_DONATUR).strip()
            if teks_don:
                ucapan_html = f"<p class='lpj-p-indent'>{teks_don}</p>"

        nama_keg = (payload.get("nama_kegiatan") or "Kegiatan Dharma Wanita Persatuan").strip()
        penutup_html = (
            f"<p class='lpj-p-indent'>Demikian laporan pertanggung jawaban {nama_keg} ini kami sampaikan, "
            "atas perhatian dan bantuan yang telah diberikan kami ucapkan terima kasih.</p>"
        )

        catatan_ig = (payload.get("catatan_ig") or "").strip()
        ig_html = ""
        if catatan_ig:
            ig_formatted = catatan_ig.replace("@dwpbcmalang", "<b>@dwpbcmalang</b>")
            ig_html = f"<p style='margin:10px 0 12px 24px;font-size:12.5px;font-style:italic;'>{ig_formatted}</p>"

        jab_1 = (payload.get("jabatan_baris_1") or "Ketua Dharma Wanita Persatuan").strip()
        jab_2 = (payload.get("jabatan_baris_2") or "KPPBC Tipe Madya Cukai Malang").strip()
        nama_ttd = (payload.get("nama_penandatangan") or "Ny. Rini Johan Pandores").strip()

        if payload.get("tampilkan_ttd", True) and sig_b64:
            sig_img_html = f"<img src='data:image/jpeg;base64,{sig_b64}' style='height:58px;margin:4px 0;display:block;'/>"
        else:
            sig_img_html = "<div style='height:54px;'></div>"

        body_content = f"""
            <div style="text-align:right;margin:8px 28px 14px 24px;">{tgl_laporan_indo}</div>
            <div style="margin:6px 24px 10px 24px;font-style:italic;">Assalamu’alaikum Wr. Wb.</div>
            <p class="lpj-p-indent">{teks_pembuka}</p>
            <table class="info-table">
                <tr><td style="width:110px;">Hari, Tanggal</td><td>: {tgl_acara_indo}</td></tr>
                <tr><td>Waktu</td><td>: {waktu_acara}</td></tr>
                <tr><td>Tempat</td><td>: {tempat_acara}</td></tr>
            </table>
            <div style="margin:8px 24px 4px 64px;">Dengan rangkaian kegiatan sebagai berikut :</div>
            <ol style="margin:2px 28px 10px 82px;padding-left:14px;">{ol_html}</ol>
            <p class="lpj-p-indent">
                Dana untuk kegiatan ini berasal dari {payload.get("sumber_dana", "")}. Dengan rincian dana sebagai berikut :
            </p>
            <table class="dana-table">
                <thead>
                    <tr>
                        <th rowspan="2" style="width:6%;">NO</th>
                        <th rowspan="2" style="width:31%;">KEBUTUHAN KEGIATAN</th>
                        <th rowspan="2" style="width:9%;">JUMLAH</th>
                        <th colspan="2" style="width:32%;">DANA</th>
                        <th rowspan="2" style="width:22%;">KETERANGAN</th>
                    </tr>
                    <tr>
                        <th style="width:16%;">HARGA SATUAN</th>
                        <th style="width:16%;">HARGA TOTAL</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                    <tr>
                        <td colspan="4" style="text-align:center;font-weight:700;background:#F8FAFC;">TOTAL</td>
                        <td style="background:#F8FAFC;">{fmt_rp_html_split(total_dana, is_dash_if_zero=False, bold=True)}</td>
                        <td style="background:#F8FAFC;"></td>
                    </tr>
                </tbody>
            </table>
            {ucapan_html}
            {penutup_html}
            <div style="margin:8px 24px 6px 24px;font-style:italic;">Wassalamu'alaikum Wr. Wb.</div>
            {ig_html}
            <div style="margin-left:58%;margin-top:14px;line-height:1.35;">
                <div>{jab_1}</div>
                <div>{jab_2}</div>
                {sig_img_html}
                <div style="font-weight:600;">{nama_ttd}</div>
            </div>
        """

    return f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8"/>
    <style>
        :root {{ color-scheme: light !important; }}
        body {{
            background: #E2E8F0;
            margin: 0;
            padding: 16px 8px;
            font-family: Arial, Helvetica, sans-serif;
            color: #000000 !important;
        }}
        .toolbar-print {{
            max-width: 215mm;
            margin: 0 auto 12px auto;
            display: flex;
            justify-content: space-between;
            align-items: center;
            background: #FFFFFF;
            padding: 10px 16px;
            border-radius: 10px;
            border: 1px solid #CBD5E1;
            box-shadow: 0 2px 6px rgba(15,23,42,0.06);
        }}
        .btn-print {{
            background: linear-gradient(135deg, #2B98F0, #1D7FE0);
            color: #FFFFFF !important;
            border: none;
            padding: 9px 16px;
            border-radius: 8px;
            font-weight: 700;
            font-size: 13.5px;
            cursor: pointer;
        }}
        .a4-sheet {{
            max-width: 215mm;
            min-height: 279mm;
            margin: 0 auto;
            background: #FFFFFF !important;
            color: #000000 !important;
            padding: 12mm 14mm 15mm 16mm;
            box-shadow: 0 8px 28px rgba(15,23,42,0.15);
            box-sizing: border-box;
            font-size: 13.5px;
            line-height: 1.42;
        }}
        .kop-table {{
            width: 100%;
            border-collapse: collapse;
            border-bottom: 2.2px solid #000000;
            margin-bottom: 12px;
        }}
        .kop-table td {{
            border: none !important;
            vertical-align: middle;
            padding-bottom: 6px;
        }}
        .lpj-p-indent {{
            text-align: justify;
            text-indent: 36px;
            margin: 6px 24px;
        }}
        .info-table {{
            margin: 6px 0 6px 64px;
            border-collapse: collapse;
        }}
        .info-table td {{
            padding: 1.5px 6px 1.5px 0;
            vertical-align: top;
            font-size: 13.5px;
        }}
        .dana-table {{
            width: calc(100% - 24px);
            margin: 8px 12px 12px 12px;
            border-collapse: collapse;
            font-size: 12.8px;
        }}
        .dana-table th, .dana-table td {{
            border: 1px solid #475569;
            padding: 4px 6px;
            color: #000000 !important;
        }}
        .dana-table th {{
            background: #F8FAFC;
            font-weight: 700;
            text-align: center;
        }}
        @media print {{
            body {{ background: #FFFFFF; padding: 0; }}
            .toolbar-print {{ display: none !important; }}
            .a4-sheet {{ box-shadow: none; margin: 0; max-width: 100%; padding: 8mm 10mm; }}
        }}
    </style>
    </head>
    <body>
        <div class="toolbar-print">
            <div style="font-size:13px;font-weight:700;color:#1E293B;">
                📄 Pratinjau Kertas A4 Resmi — Mode Dokumen: {doc_mode} DWP KPPBC TMC Malang
            </div>
            <button class="btn-print" onclick="window.print()">🖨️ Cetak / Simpan PDF via Browser</button>
        </div>
        <div class="a4-sheet">
            {kop_html}
            {body_content}
        </div>
    </body>
    </html>
    """


# ==============================================================================
# BAGIAN 6: SPESIFIKASI CSS NATIVE CONTAINER V9.0 (ZERO EMPTY WHITE BARS)
# ==============================================================================

def inject_fily_admin_css():
    st.markdown(
        """
        <style>
        :root {
            color-scheme: light !important;
        }
        html, body, [data-testid="stAppViewContainer"], [data-testid="stMain"] {
            background-color: #F1F5F9 !important;
            color: #0F172A !important;
            font-family: 'Inter', 'Segoe UI', Roboto, sans-serif;
        }
        [data-testid="stHeader"] {
            background-color: rgba(241, 245, 249, 0.90) !important;
        }
        .block-container {
            padding-top: 1.1rem !important;
            padding-bottom: 2.5rem !important;
            max-width: 1440px !important;
        }
        /* Native st.container(border=True) Card Styling — 0% Empty White Bars */
        [data-testid="stMain"] [data-testid="stVerticalBlockBorderWrapper"] {
            background-color: #FFFFFF !important;
            border-radius: 14px !important;
            border: 1px solid #E2E8F0 !important;
            padding: 18px 20px !important;
            margin-bottom: 16px !important;
            box-shadow: 0 2px 10px rgba(15, 23, 42, 0.04) !important;
        }
        /* Dark Slate Navy Sidebar (#1E2235) */
        [data-testid="stSidebar"], [data-testid="stSidebar"] > div:first-child {
            background-color: #1E2235 !important;
            border-right: 1px solid #2D334D !important;
        }
        [data-testid="stSidebar"] h1,
        [data-testid="stSidebar"] h2,
        [data-testid="stSidebar"] h3,
        [data-testid="stSidebar"] h4,
        [data-testid="stSidebar"] p,
        [data-testid="stSidebar"] span,
        [data-testid="stSidebar"] label,
        [data-testid="stSidebar"] caption {
            color: #F8FAFC !important;
            -webkit-text-fill-color: #F8FAFC !important;
        }
        /* Navigasi Radio di Sidebar tampil seperti Menu Web Modern */
        [data-testid="stSidebar"] div[role="radiogroup"] > label {
            background-color: rgba(255, 255, 255, 0.04) !important;
            border: 1px solid rgba(255, 255, 255, 0.08) !important;
            border-left: 4px solid transparent !important;
            border-radius: 10px !important;
            padding: 10px 14px !important;
            margin-bottom: 6px !important;
            min-height: 46px !important;
            transition: all 0.18s ease !important;
            cursor: pointer !important;
        }
        [data-testid="stSidebar"] div[role="radiogroup"] > label:hover {
            background-color: #282E47 !important;
            border-left: 4px solid #2B98F0 !important;
        }
        [data-testid="stSidebar"] div[role="radiogroup"] > label[data-checked="true"],
        [data-testid="stSidebar"] div[role="radiogroup"] > label:has(input:checked) {
            background: linear-gradient(90deg, rgba(43,152,240,0.25) 0%, rgba(40,46,71,0.95) 100%) !important;
            border: 1px solid rgba(43, 152, 240, 0.45) !important;
            border-left: 4px solid #00D09C !important;
            font-weight: 700 !important;
        }
        /* Input di dalam Expander Sidebar tetap berlatar Putih & Teks Gelap */
        [data-testid="stSidebar"] input,
        [data-testid="stSidebar"] textarea,
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {
            background-color: #FFFFFF !important;
            color: #0F172A !important;
            -webkit-text-fill-color: #0F172A !important;
        }
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] * {
            color: #0F172A !important;
            -webkit-text-fill-color: #0F172A !important;
        }
        /* Area Utama: Paksa semua teks gelap pekat #0F172A */
        [data-testid="stMain"] p,
        [data-testid="stMain"] span,
        [data-testid="stMain"] label,
        [data-testid="stMain"] h1,
        [data-testid="stMain"] h2,
        [data-testid="stMain"] h3,
        [data-testid="stMain"] h4,
        [data-testid="stMain"] h5,
        [data-testid="stMain"] li,
        [data-testid="stMain"] td,
        [data-testid="stMain"] th {
            color: #0F172A !important;
            -webkit-text-fill-color: #0F172A !important;
        }
        /* Input Ergonomis Mobile-First (Anti Auto-Zoom HP: 16px, min-height 46px) */
        [data-testid="stMain"] input,
        [data-testid="stMain"] textarea,
        [data-testid="stMain"] select,
        [data-testid="stMain"] div[data-baseweb="input"] input,
        [data-testid="stMain"] div[data-baseweb="textarea"] textarea,
        [data-testid="stMain"] div[data-baseweb="select"] > div {
            background-color: #FFFFFF !important;
            color: #0F172A !important;
            -webkit-text-fill-color: #0F172A !important;
            font-size: 16px !important;
            min-height: 46px !important;
            border-radius: 8px !important;
            border: 1.5px solid #CBD5E1 !important;
        }
        div[data-baseweb="popover"], ul[data-baseweb="menu"], li[role="option"] {
            background-color: #FFFFFF !important;
            color: #0F172A !important;
            -webkit-text-fill-color: #0F172A !important;
            font-size: 16px !important;
        }
        /* Tombol Utama Modern */
        [data-testid="stMain"] .stButton > button {
            min-height: 45px !important;
            border-radius: 10px !important;
            font-size: 14px !important;
            font-weight: 700 !important;
            background: linear-gradient(135deg, #2B98F0 0%, #1D7FE0 100%) !important;
            color: #FFFFFF !important;
            -webkit-text-fill-color: #FFFFFF !important;
            border: none !important;
            box-shadow: 0 4px 12px rgba(43, 152, 240, 0.20) !important;
            transition: all 0.18s ease !important;
        }
        [data-testid="stMain"] .stButton > button p,
        [data-testid="stMain"] .stButton > button span {
            color: #FFFFFF !important;
            -webkit-text-fill-color: #FFFFFF !important;
        }
        [data-testid="stMain"] .stButton > button:hover {
            transform: translateY(-1px);
            box-shadow: 0 6px 16px rgba(43, 152, 240, 0.30) !important;
        }
        [data-testid="stSidebar"] .stButton > button {
            min-height: 42px !important;
            border-radius: 9px !important;
            font-size: 13.5px !important;
            font-weight: 700 !important;
            background: linear-gradient(135deg, #2B98F0 0%, #1D7FE0 100%) !important;
            color: #FFFFFF !important;
            -webkit-text-fill-color: #FFFFFF !important;
            border: 1px solid rgba(255,255,255,0.18) !important;
        }
        /* Top Welcome Bar & 4 Pill KPI Cards ("Fily" Style) */
        .fily-topbar {
            background: #FFFFFF !important;
            padding: 14px 20px;
            border-radius: 14px;
            border: 1px solid #E2E8F0;
            display: flex;
            justify-content: space-between;
            align-items: center;
            flex-wrap: wrap;
            gap: 12px;
            margin-bottom: 14px;
            box-shadow: 0 2px 8px rgba(15, 23, 42, 0.04);
        }
        .fily-pill {
            border-radius: 999px;
            padding: 14px 20px;
            color: #FFFFFF !important;
            display: flex;
            align-items: center;
            justify-content: space-between;
            box-shadow: 0 6px 18px rgba(15, 23, 42, 0.10);
            margin-bottom: 14px;
        }
        .fily-pill * {
            color: #FFFFFF !important;
            -webkit-text-fill-color: #FFFFFF !important;
        }
        .fily-pill-icon {
            width: 46px;
            height: 46px;
            border-radius: 50%;
            background: rgba(255, 255, 255, 0.22);
            display: flex;
            align-items: center;
            justify-content: center;
            font-size: 21px;
            flex-shrink: 0;
        }
        .pill-emerald { background: linear-gradient(135deg, #00D09C 0%, #00B386 100%); }
        .pill-blue    { background: linear-gradient(135deg, #2B98F0 0%, #1D7FE0 100%); }
        .pill-rose    { background: linear-gradient(135deg, #F85370 0%, #E63956 100%); }
        .pill-amber   { background: linear-gradient(135deg, #FBA01D 0%, #F58800 100%); }
        </style>
        """,
        unsafe_allow_html=True,
    )


# ==============================================================================
# BAGIAN 7: MANAJEMEN STATE, KONVERSI CERDAS RAB->LPJ & BAR RIWAYAT V9.0
# ==============================================================================

def load_rab_into_state(data: dict):
    """Memuat seluruh field RAB ke st.session_state."""
    st.session_state["id_rab"] = data.get("id_rab", "RAB-" + datetime.now().strftime("%Y%m%d-%H%M%S"))
    st.session_state["baris_1_judul"] = data.get("baris_1_judul", "RANCANGAN ANGGARAN BIAYA")
    st.session_state["nama_kegiatan_rab"] = data.get("nama_kegiatan_rab", "")
    st.session_state["lokasi_acara_rab"] = data.get("lokasi_acara_rab", "")
    st.session_state["sumber_dana_rab"] = data.get("sumber_dana_rab", "Kas Dharma Wanita KPPBC TMC Malang")
    st.session_state["ket_tabel_rab"] = data.get("ket_tabel_rab", "")
    st.session_state["merge_ket_rab"] = bool(data.get("merge_ket_rab", False))
    df_norm = normalize_items_df(data.get("items_rab", []), st.session_state["ket_tabel_rab"])
    st.session_state["df_items_rab"] = df_norm
    st.session_state["catatan_note_rab"] = sync_iuran_note_rab(data.get("catatan_note_rab", ""), df_norm)
    st.session_state["status_konversi"] = data.get("status_konversi", "BELUM DIKONVERSI")
    st.session_state["rab_editor_ver"] = int(st.session_state.get("rab_editor_ver", 0)) + 1


def load_lpj_into_state(data: dict):
    """Memuat seluruh field LPJ ke st.session_state."""
    st.session_state["id_lpj"] = data.get("id_lpj", "LPJ-" + datetime.now().strftime("%Y%m%d-%H%M%S"))
    st.session_state["id_rab_ref"] = data.get("id_rab_ref", "")

    tgl_lap = data.get("tgl_laporan", date(2026, 8, 13))
    if isinstance(tgl_lap, str):
        try:
            tgl_lap = datetime.strptime(tgl_lap[:10], "%Y-%m-%d").date()
        except Exception:
            tgl_lap = date(2026, 8, 13)
    st.session_state["tgl_laporan"] = tgl_lap

    st.session_state["nama_kegiatan"] = data.get("nama_kegiatan", "")
    st.session_state["tema_kegiatan"] = data.get("tema_kegiatan", "")
    st.session_state["konteks_hut"] = data.get("konteks_hut", "")

    tgl_acr = data.get("tgl_acara", date(2026, 8, 10))
    if isinstance(tgl_acr, str):
        try:
            tgl_acr = datetime.strptime(tgl_acr[:10], "%Y-%m-%d").date()
        except Exception:
            tgl_acr = date(2026, 8, 10)
    st.session_state["tgl_acara"] = tgl_acr

    st.session_state["tgl_acara_manual"] = data.get("tgl_acara_manual", "")
    st.session_state["waktu_acara"] = data.get("waktu_acara", "08.30 - selesai")
    st.session_state["tempat_acara"] = data.get("tempat_acara", "")
    st.session_state["rangkaian_kegiatan"] = data.get("rangkaian_kegiatan", "")
    st.session_state["sumber_dana"] = data.get("sumber_dana", "Kas Dharma Wanita KPPBC TMC Malang dan donatur")
    st.session_state["ket_tabel_dana"] = data.get("ket_tabel_dana", "Kas DWP BC Malang & Donatur Anggota")
    st.session_state["merge_ket_vertikal"] = bool(data.get("merge_ket_vertikal", True))
    st.session_state["tampilkan_ucapan_donatur"] = bool(data.get("tampilkan_ucapan_donatur", True))
    st.session_state["teks_ucapan_donatur"] = data.get("teks_ucapan_donatur", DEFAULT_UCAPAN_DONATUR)
    st.session_state["catatan_ig"] = data.get("catatan_ig", "*Note : dokumentasi kegiatan dapat dilihat di Instagram @dwpbcmalang")
    st.session_state["jabatan_baris_1"] = data.get("jabatan_baris_1", "Ketua Dharma Wanita Persatuan")
    st.session_state["jabatan_baris_2"] = data.get("jabatan_baris_2", "KPPBC Tipe Madya Cukai Malang")
    st.session_state["nama_penandatangan"] = data.get("nama_penandatangan", "Ny. Rini Johan Pandores")
    st.session_state["tampilkan_ttd"] = bool(data.get("tampilkan_ttd", True))
    st.session_state["df_items_dana"] = normalize_items_df(data.get("items_dana", []), st.session_state["ket_tabel_dana"])
    st.session_state["lpj_editor_ver"] = int(st.session_state.get("lpj_editor_ver", 0)) + 1


def collect_rab_payload(df_rab_override: pd.DataFrame = None) -> dict:
    df_use = df_rab_override if df_rab_override is not None else st.session_state.get("df_items_rab", pd.DataFrame())
    df_clean = normalize_items_df(df_use, st.session_state.get("ket_tabel_rab", ""), include_hapus_col=False)
    return {
        "id_rab": st.session_state.get("id_rab", "RAB-2026-JEMBER"),
        "baris_1_judul": st.session_state.get("baris_1_judul", "RANCANGAN ANGGARAN BIAYA"),
        "nama_kegiatan_rab": st.session_state.get("nama_kegiatan_rab", ""),
        "lokasi_acara_rab": st.session_state.get("lokasi_acara_rab", ""),
        "sumber_dana_rab": st.session_state.get("sumber_dana_rab", ""),
        "ket_tabel_rab": st.session_state.get("ket_tabel_rab", ""),
        "merge_ket_rab": st.session_state.get("merge_ket_rab", False),
        "catatan_note_rab": st.session_state.get("catatan_note_rab", ""),
        "status_konversi": st.session_state.get("status_konversi", "BELUM DIKONVERSI"),
        "items_rab": df_clean.to_dict(orient="records"),
    }


def collect_lpj_payload(df_lpj_override: pd.DataFrame = None) -> dict:
    df_use = df_lpj_override if df_lpj_override is not None else st.session_state.get("df_items_dana", pd.DataFrame())
    df_clean = normalize_items_df(df_use, st.session_state.get("ket_tabel_dana", ""), include_hapus_col=False)
    return {
        "id_lpj": st.session_state.get("id_lpj", "LPJ-20260813-BLITAR"),
        "id_rab_ref": st.session_state.get("id_rab_ref", ""),
        "tgl_laporan": st.session_state.get("tgl_laporan", date(2026, 8, 13)),
        "nama_kegiatan": st.session_state.get("nama_kegiatan", ""),
        "tema_kegiatan": st.session_state.get("tema_kegiatan", ""),
        "konteks_hut": st.session_state.get("konteks_hut", ""),
        "tgl_acara": st.session_state.get("tgl_acara", date(2026, 8, 10)),
        "tgl_acara_manual": st.session_state.get("tgl_acara_manual", ""),
        "waktu_acara": st.session_state.get("waktu_acara", "08.30 - selesai"),
        "tempat_acara": st.session_state.get("tempat_acara", ""),
        "rangkaian_kegiatan": st.session_state.get("rangkaian_kegiatan", ""),
        "sumber_dana": st.session_state.get("sumber_dana", ""),
        "ket_tabel_dana": st.session_state.get("ket_tabel_dana", ""),
        "merge_ket_vertikal": st.session_state.get("merge_ket_vertikal", True),
        "tampilkan_ucapan_donatur": st.session_state.get("tampilkan_ucapan_donatur", True),
        "teks_ucapan_donatur": st.session_state.get("teks_ucapan_donatur", DEFAULT_UCAPAN_DONATUR),
        "catatan_ig": st.session_state.get("catatan_ig", ""),
        "jabatan_baris_1": st.session_state.get("jabatan_baris_1", "Ketua Dharma Wanita Persatuan"),
        "jabatan_baris_2": st.session_state.get("jabatan_baris_2", "KPPBC Tipe Madya Cukai Malang"),
        "nama_penandatangan": st.session_state.get("nama_penandatangan", "Ny. Rini Johan Pandores"),
        "tampilkan_ttd": st.session_state.get("tampilkan_ttd", True),
        "items_dana": df_clean.to_dict(orient="records"),
    }


def convert_active_rab_to_lpj(df_rab_current: pd.DataFrame = None, sheet_id: str = "", creds_str: str = ""):
    """
    PERBAIKAN MUTLAK V9.0: Konversi RAB ke LPJ secara cerdas dan kontekstual:
    - Menarik Nama Kegiatan, Lokasi Acara, Sumber Dana, dan seluruh baris anggaran RAB
    - Menetapkan `id_rab_ref = id_rab`
    - Menyesuaikan `rangkaian_kegiatan`, `tema_kegiatan`, dan `konteks_hut` agar sesuai dengan lokasi RAB
      (TIDAK menyisakan teks rangkaian kegiatan Blitar / Bapak Nurtjahjo / Museum Bung Karno jika lokasi RAB bukan di Blitar!)
    """
    raw_nama_rab = (st.session_state.get("nama_kegiatan_rab") or "").strip()
    if raw_nama_rab.upper().startswith("PELAKSANAAN "):
        raw_nama_rab = raw_nama_rab[12:].strip()

    nama_lpj = raw_nama_rab.title()
    for acr in ["DWP", "DJBC", "KPPBC", "TMC", "TMP", "HUT", "RI", "II", "III"]:
        nama_lpj = nama_lpj.replace(acr.title(), acr)
    nama_lpj = nama_lpj.replace("Di Lingkungan", "di lingkungan")

    lok_rab = (st.session_state.get("lokasi_acara_rab") or "").strip().strip("()")
    if lok_rab.upper().startswith("DI "):
        lok_rab = lok_rab[3:].strip()
    tempat_lpj = lok_rab.title()
    for acr in ["KPPBC", "TMC", "TMP", "DJBC", "C", "A", "B"]:
        tempat_lpj = tempat_lpj.replace(f" {acr.title()} ", f" {acr} ").replace(f"{acr.title()} ", f"{acr} ")
    tempat_lpj = tempat_lpj.replace("Tmp C", "Tipe Madya Pabean C").replace("TMP C", "Tipe Madya Pabean C")

    sumber_rab = (st.session_state.get("sumber_dana_rab") or "Kas Dharma Wanita KPPBC TMC Malang").strip()
    if "donatur" not in sumber_rab.lower():
        sumber_lpj = sumber_rab + " dan donatur"
    else:
        sumber_lpj = sumber_rab

    ket_lpj = st.session_state.get("ket_tabel_dana") or "Kas DWP BC Malang & Donatur Anggota"
    df_src = df_rab_current if df_rab_current is not None else st.session_state.get("df_items_rab", pd.DataFrame())
    df_lpj_new = normalize_items_df(df_src, ket_lpj, include_hapus_col=False)
    df_lpj_new["KETERANGAN"] = ket_lpj

    # Cek apakah lokasi RAB bukan di Blitar; jika bukan Blitar namun rangkaian kegiatan masih berisi teks Blitar, ganti!
    cur_rangkaian = str(st.session_state.get("rangkaian_kegiatan") or "")
    is_rab_blitar = "blitar" in tempat_lpj.lower() if False else ("blitar" in tempat_lpj.lower())
    has_blitar_leftover = any(k in cur_rangkaian.lower() for k in ["blitar", "nurtjahjo", "bung karno", "tumpeng"])

    if not is_rab_blitar and (has_blitar_leftover or not cur_rangkaian.strip()):
        st.session_state["rangkaian_kegiatan"] = NL.join([
            f"Pembukaan acara dan sambutan pada kegiatan di {tempat_lpj or 'lokasi acara'}.",
            "Pelaksanaan kegiatan inti pertemuan rutin dan ramah tamah anggota Dharma Wanita Persatuan.",
            "Penutupan acara dan dokumentasi bersama.",
        ])
        st.session_state["tema_kegiatan"] = ""
        st.session_state["konteks_hut"] = ""

    kota_tag = "".join(ch for ch in tempat_lpj.split()[-1].upper() if ch.isalnum()) if tempat_lpj.strip() else "DWP"
    new_lpj_id = f"LPJ-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{kota_tag}"
    st.session_state["id_lpj"] = new_lpj_id
    st.session_state["id_rab_ref"] = st.session_state.get("id_rab", "RAB-AKTIF")
    st.session_state["nama_kegiatan"] = nama_lpj
    st.session_state["tempat_acara"] = tempat_lpj
    st.session_state["sumber_dana"] = sumber_lpj
    st.session_state["df_items_dana"] = df_lpj_new
    st.session_state["status_konversi"] = f"SUDAH DIKONVERSI KE {new_lpj_id}"
    st.session_state["preview_doc_mode"] = "LPJ"
    st.session_state["lpj_editor_ver"] = int(st.session_state.get("lpj_editor_ver", 0)) + 1

    save_rab_to_storage(collect_rab_payload(df_src), sheet_id, creds_str)
    save_lpj_to_storage(collect_lpj_payload(df_lpj_new), sheet_id, creds_str)


def init_app_state():
    if "app_v9_initialized" not in st.session_state:
        st.session_state["rab_editor_ver"] = 0
        st.session_state["lpj_editor_ver"] = 0
        st.session_state["deleted_doc_ids"] = []
        persisted = load_persistent_local_history()
        st.session_state["local_db_rab"] = persisted.get("rab_rows", [])
        st.session_state["local_db_lpj"] = persisted.get("lpj_rows", [])
        st.session_state["local_rekap_dana"] = persisted.get("rekap_rows", [])
        load_rab_into_state(SEED_HISTORY_RAB_JEMBER)
        load_lpj_into_state(SEED_HISTORY_LPJ_BLITAR_2026)
        st.session_state["active_menu"] = MENU_RAB
        st.session_state["preview_doc_mode"] = "RAB"
        st.session_state["app_v9_initialized"] = True


def switch_menu(target_menu: str, preview_mode: str = None):
    st.session_state["_nav_target"] = target_menu
    if preview_mode:
        st.session_state["preview_doc_mode"] = preview_mode
    st.rerun()


def render_history_quick_bar(context_key: str, default_doc_filter: str, bundle: dict, sheet_id_active: str, creds_active: str):
    """
    PERBAIKAN MUTLAK V9.0: Panel Pilih dari Riwayat / Salin / Cetak Ulang / Hapus Data (4 Tombol Lengkap)
    Menggunakan `with st.container(border=True):` asli Streamlit sehingga 0% pita putih kosong!
    """
    rab_rows = bundle.get("rab_rows", [])
    lpj_rows = bundle.get("lpj_rows", [])

    options_map = {}
    ordered_keys = []

    if default_doc_filter in ("ALL", "RAB"):
        for r in rab_rows:
            rid = r.get("ID_RAB", "RAB")
            lok = r.get("LOKASI_ACARA", "")
            tot = fmt_rp_cell(r.get("TOTAL_RAB", 0), is_dash_if_zero=False)
            wkt = str(r.get("WAKTU_SIMPAN", ""))[:16]
            label = f"[RAB] {rid} | {lok} ({tot}) — {wkt}"
            options_map[label] = ("RAB", r)
            ordered_keys.append(label)

    if default_doc_filter in ("ALL", "LPJ"):
        for r in lpj_rows:
            lid = r.get("ID_LPJ", "LPJ")
            tmp = r.get("TEMPAT_ACARA", "")
            tot = fmt_rp_cell(r.get("TOTAL_DANA", 0), is_dash_if_zero=False)
            wkt = str(r.get("WAKTU_SIMPAN", ""))[:16]
            label = f"[LPJ] {lid} | {tmp} ({tot}) — {wkt}"
            options_map[label] = ("LPJ", r)
            ordered_keys.append(label)

    if default_doc_filter == "RAB":
        for r in lpj_rows:
            lid = r.get("ID_LPJ", "LPJ")
            tmp = r.get("TEMPAT_ACARA", "")
            tot = fmt_rp_cell(r.get("TOTAL_DANA", 0), is_dash_if_zero=False)
            label = f"[LPJ] {lid} | {tmp} ({tot})"
            options_map[label] = ("LPJ", r)
            ordered_keys.append(label)
    elif default_doc_filter == "LPJ":
        for r in rab_rows:
            rid = r.get("ID_RAB", "RAB")
            lok = r.get("LOKASI_ACARA", "")
            tot = fmt_rp_cell(r.get("TOTAL_RAB", 0), is_dash_if_zero=False)
            label = f"[RAB] {rid} | {lok} ({tot})"
            options_map[label] = ("RAB", r)
            ordered_keys.append(label)

    with st.container(border=True):
        st.markdown("##### 🗂️ Pilih dari Riwayat / Salin ke Kota Baru / Cetak Ulang / Hapus Data")
        hc1, hc2 = st.columns([2.2, 1.1])
        with hc1:
            sel_hist_label = st.selectbox(
                "Pilih Dokumen dari Riwayat (RAB / LPJ):",
                options=ordered_keys if ordered_keys else ["-- Belum Ada Dokumen Riwayat --"],
                key=f"hist_sel_{context_key}",
            )
        with hc2:
            new_city_in = st.text_input(
                "Kota Tujuan Baru (Jika Menyalin Dokumen):",
                placeholder="Misal: Kediri / Pasuruan / Malang",
                key=f"hist_city_{context_key}",
            )

        hb1, hb2, hb3, hb4 = st.columns(4)
        with hb1:
            if st.button("📂 Buka & Edit", key=f"btn_open_hist_{context_key}", use_container_width=True):
                if sel_hist_label in options_map:
                    dtype, row_obj = options_map[sel_hist_label]
                    raw_j = row_obj.get("PAYLOAD_JSON", "")
                    if raw_j:
                        parsed_p = json.loads(raw_j)
                        if dtype == "RAB":
                            load_rab_into_state(parsed_p)
                            st.toast(f"Riwayat RAB ({row_obj.get('ID_RAB')}) berhasil dimuat!", icon="📂")
                            switch_menu(MENU_RAB)
                        else:
                            load_lpj_into_state(parsed_p)
                            st.toast(f"Riwayat LPJ ({row_obj.get('ID_LPJ')}) berhasil dimuat!", icon="📂")
                            switch_menu(MENU_LPJ)

        with hb2:
            if st.button("📑 Salin ke Kota Baru", key=f"btn_copy_hist_{context_key}", use_container_width=True):
                if sel_hist_label in options_map:
                    dtype, row_obj = options_map[sel_hist_label]
                    raw_j = row_obj.get("PAYLOAD_JSON", "")
                    if raw_j:
                        parsed_p = json.loads(raw_j)
                        target_city = (new_city_in or "").strip() or "Kediri"
                        cloned_p = replace_city_across_payload(dtype, parsed_p, target_city)
                        if dtype == "RAB":
                            load_rab_into_state(cloned_p)
                            save_rab_to_storage(collect_rab_payload(), sheet_id_active, creds_active)
                            st.toast(f"Salinan RAB baru untuk kota {target_city.title()} berhasil dibuat!", icon="📑")
                            switch_menu(MENU_RAB)
                        else:
                            load_lpj_into_state(cloned_p)
                            save_lpj_to_storage(collect_lpj_payload(), sheet_id_active, creds_active)
                            st.toast(f"Salinan LPJ baru untuk kota {target_city.title()} berhasil dibuat!", icon="📑")
                            switch_menu(MENU_LPJ)

        with hb3:
            if st.button("🖨️ Cetak Ulang (Word/PDF)", key=f"btn_reprint_hist_{context_key}", use_container_width=True):
                if sel_hist_label in options_map:
                    dtype, row_obj = options_map[sel_hist_label]
                    raw_j = row_obj.get("PAYLOAD_JSON", "")
                    if raw_j:
                        parsed_p = json.loads(raw_j)
                        if dtype == "RAB":
                            load_rab_into_state(parsed_p)
                            switch_menu(MENU_PREVIEW, preview_mode="RAB")
                        else:
                            load_lpj_into_state(parsed_p)
                            switch_menu(MENU_PREVIEW, preview_mode="LPJ")

        with hb4:
            if st.button("🗑️ Hapus Dokumen Terpilih", key=f"btn_del_hist_{context_key}", use_container_width=True):
                if sel_hist_label in options_map:
                    dtype, row_obj = options_map[sel_hist_label]
                    doc_id_del = row_obj.get("ID_RAB") if dtype == "RAB" else row_obj.get("ID_LPJ")
                    ok_del, msg_del = delete_document_from_storage(dtype, doc_id_del, sheet_id_active, creds_active)
                    if ok_del:
                        st.toast(msg_del, icon="🗑️")
                        st.rerun()
                    else:
                        st.error(msg_del)


def render_mail_merge_expander(doc_mode: str, payload: dict):
    """Expander Tabel Pemetaan Variabel Mail Merge (<<TAG>> -> Isi Dokumen)."""
    with st.expander(f"🏷️ Lihat Pemetaan Variabel Mail Merge ({doc_mode})", expanded=False):
        mm_list = build_mail_merge_dictionary(doc_mode, payload)
        st.dataframe(pd.DataFrame(mm_list), use_container_width=True, hide_index=True)


# ==============================================================================
# BAGIAN 8: MAIN APPLICATION V9.0 (CLEAN NATIVE CONTAINERS & FULL CRUD)
# ==============================================================================

def main():
    inject_fily_admin_css()
    init_app_state()

    if "_nav_target" in st.session_state:
        st.session_state["active_menu"] = st.session_state.pop("_nav_target")

    default_sheet_id = get_active_spreadsheet_id()

    # --- SIDEBAR: DARK SLATE NAVY ("FILY" ADMIN NAVIGATION V9.0) ---
    with st.sidebar:
        st.markdown(
            """
            <div style="display:flex;align-items:center;gap:12px;padding:6px 4px 16px 4px;border-bottom:1px solid #2D334D;margin-bottom:14px;">
                <div style="width:42px;height:42px;border-radius:10px;background:linear-gradient(135deg,#2B98F0,#00D09C);display:flex;align-items:center;justify-content:center;font-size:22px;box-shadow:0 4px 12px rgba(43,152,240,0.3);">
                    🏛️
                </div>
                <div>
                    <div style="font-weight:800;font-size:16px;letter-spacing:0.4px;color:#FFFFFF;">DWP BC MALANG</div>
                    <div style="font-size:11.5px;color:#94A3B8;">RAB & LPJ Suite V9.0</div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.caption("MENU UTAMA")
        active_menu = st.radio(
            "Menu Utama",
            options=ALL_MENUS,
            key="active_menu",
            label_visibility="collapsed",
        )

        df_rab_side = st.session_state.get("df_items_rab", pd.DataFrame())
        df_lpj_side = st.session_state.get("df_items_dana", pd.DataFrame())
        tot_rab_side = int(df_rab_side["HARGA TOTAL"].sum()) if not df_rab_side.empty else 0
        tot_lpj_side = int(df_lpj_side["HARGA TOTAL"].sum()) if not df_lpj_side.empty else 0

        st.markdown('<div style="border-top:1px solid #2D334D;margin:12px 0 10px 0;"></div>', unsafe_allow_html=True)
        st.caption("DOKUMEN SEDANG DIEDIT")
        st.markdown(
            f"""
            <div style="background:rgba(255,255,255,0.05);border:1px solid rgba(255,255,255,0.1);border-radius:9px;padding:9px 12px;font-size:12.5px;margin-bottom:10px;">
                <div>• <b>RAB:</b> {st.session_state.get('id_rab', '-')}</div>
                <div style="color:#00D09C !important;font-weight:700;margin-left:10px;margin-bottom:6px;">{fmt_rp_cell(tot_rab_side, False)}</div>
                <div>• <b>LPJ:</b> {st.session_state.get('id_lpj', '-')}</div>
                <div style="color:#2B98F0 !important;font-weight:700;margin-left:10px;">{fmt_rp_cell(tot_lpj_side, False)}</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.caption("BUAT FORMULIR BARU & CONTOH")
        if st.button("➕ Mulai RAB Baru", use_container_width=True):
            new_rab = dict(PRESET_RAB_JEMBER)
            new_rab["id_rab"] = "RAB-" + datetime.now().strftime("%Y%m%d-%H%M%S")
            load_rab_into_state(new_rab)
            st.toast(f"Formulir RAB baru ({new_rab['id_rab']}) siap diisi!", icon="➕")
            switch_menu(MENU_RAB)

        if st.button("➕ Mulai LPJ Baru", use_container_width=True):
            new_lpj = dict(PRESET_LPJ_BLITAR_2026)
            new_lpj["id_lpj"] = "LPJ-" + datetime.now().strftime("%Y%m%d-%H%M%S")
            new_lpj["id_rab_ref"] = ""
            load_lpj_into_state(new_lpj)
            st.toast(f"Formulir LPJ baru ({new_lpj['id_lpj']}) siap diisi!", icon="➕")
            switch_menu(MENU_LPJ)

        if st.button("📋 Muat Contoh RAB Jember", use_container_width=True):
            load_rab_into_state(SEED_HISTORY_RAB_JEMBER)
            st.toast("Contoh RAB DWP Gabungan di Jember (Rp 2.094.000) dimuat!", icon="📋")
            switch_menu(MENU_RAB)

        if st.button("📝 Muat Contoh LPJ Blitar 2026", use_container_width=True):
            load_lpj_into_state(SEED_HISTORY_LPJ_BLITAR_2026)
            st.toast("Contoh LPJ DWP Gabungan di Blitar (Rp 5.075.000) dimuat!", icon="📝")
            switch_menu(MENU_LPJ)

        st.markdown('<div style="border-top:1px solid #2D334D;margin:12px 0 10px 0;"></div>', unsafe_allow_html=True)
        with st.expander("⚙️ KONEKSI CLOUD & TEMPLATE", expanded=False):
            st.text_input("ID Google Sheets", value=default_sheet_id, key="cfg_sheet_id")
            up_sa = st.file_uploader("Upload JSON Service Account", type=["json"], key="up_sa_json")
            if up_sa is not None:
                raw_sa_txt = up_sa.getvalue().decode("utf-8", errors="ignore")
                st.session_state["creds_json_str"] = raw_sa_txt
                try:
                    with open(LOCAL_GCP_SA_FILE, "w", encoding="utf-8") as f_sa:
                        f_sa.write(raw_sa_txt)
                    st.success("Service Account disimpan permanen!")
                except Exception:
                    pass

            up_tpl = st.file_uploader("Upload Template .docx (Opsional)", type=["docx"], key="up_tpl_sidebar")
            if up_tpl is not None:
                st.session_state["uploaded_template_bytes"] = up_tpl.read()
                st.success("Template .docx aktif!")

            if st.button("🔄 Sinkronkan Cloud", use_container_width=True):
                fetch_sheet_bundle_cached.clear()
                st.rerun()

        sheet_id_active = get_active_spreadsheet_id()
        creds_active = st.session_state.get("creds_json_str", "")
        bundle = get_sheet_bundle(sheet_id_active, creds_active)

        status_dot = "🟢 Cloud & Riwayat Lokal Aktif" if bundle.get("ok") else "🟢 Riwayat Lokal Permanen Aktif"
        st.caption(f"Status: **{status_dot}**")

    df_rab_now = st.session_state.get("df_items_rab", pd.DataFrame())
    df_lpj_now = st.session_state.get("df_items_dana", pd.DataFrame())

    # --- TOP WELCOME BAR ---
    st.markdown(
        f"""
        <div class="fily-topbar">
            <div>
                <div style="font-size:17.5px;font-weight:800;color:#0F172A;">
                    👋 Sistem Generator RAB & LPJ DWP KPPBC Tipe Madya Cukai Malang (V9.0)
                </div>
                <div style="font-size:12.8px;color:#475569;margin-top:2px;">
                    Format Akuntansi <b>Rp Rata Kiri Sejajar TOTAL</b> • <b>Otomasi Iuran DWP (5 Orang 300rb + 75rb/org)</b> • <b>Fitur Hapus Baris & Riwayat</b>
                </div>
            </div>
            <div style="font-size:12.5px;font-weight:700;background:#EFF6FF;color:#1D4ED8;padding:6px 12px;border-radius:8px;border:1px solid #BFDBFE;">
                {status_dot}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    # ==========================================================================
    # VIEW 1: 📋 1. BUAT / EDIT RAB (PRA-KEGIATAN)
    # ==========================================================================
    if active_menu == MENU_RAB:
        render_history_quick_bar("rab", "RAB", bundle, sheet_id_active, creds_active)

        # LANGKAH 1: IDENTITAS KEGIATAN RAB (TANPA <<TAG>> DI LABEL)
        with st.container(border=True):
            st.markdown("#### 📌 Langkah 1: Identitas Kegiatan RAB")
            r_c1, r_c2 = st.columns([1.6, 1.0])
            with r_c1:
                st.text_area(
                    "Nama Kegiatan (Otomatis Dicetak Huruf Kapital di Dokumen):",
                    key="nama_kegiatan_rab",
                    height=76,
                )
            with r_c2:
                st.text_input(
                    "Lokasi Acara (Contoh: (DI KPPBC TMP C JEMBER)):",
                    key="lokasi_acara_rab",
                )
                st.text_input(
                    "Sumber Dana:",
                    key="sumber_dana_rab",
                )

        # LANGKAH 2: TABEL RINCIAN DANA RAB (DENGAN KOLOM HAPUS? & TOMBOL HAPUS BARIS)
        with st.container(border=True):
            st.markdown("#### 💰 Langkah 2: Tabel Rincian Dana RAB (Rp Rata Kiri & Bisa Hapus Baris)")
            st.caption(
                "💡 **Cara Pakai:** "
                "• **Baris Iuran DWP Gabungan**: Ketik `JUMLAH` orang (misal `7` atau `9`), maka `HARGA TOTAL` otomatis menghitung paket **Rp 300.000 + (Rp 75.000/orang di atas 5)** & memperbarui `*Note` di bawah! "
                "• **Barang Biasa**: Ketik `JUMLAH` & `HARGA SATUAN` (Total terhitung otomatis). "
                "• **Hapus Baris**: Centang kotak **🗑️ HAPUS?** di sisi kanan baris yang ingin dibuang, lalu klik tombol **🗑️ Hapus Baris yang Dicentang** di bawah tabel."
            )

            with st.expander("⚡ Copy-Paste Daftar Belanja dari WhatsApp / Tambah Baris Cepat", expanded=False):
                mp1, mp2 = st.columns([1.3, 1.0])
                with mp1:
                    wa_rab = st.text_area(
                        "Tempel rincian anggaran dari WhatsApp (Contoh: Iuran DWP Gabungan di Jember 9, Doorprize 7 x 50.000, Mobil 2 -):",
                        height=100,
                        key="wa_paste_rab",
                    )
                    b_p1, b_p2 = st.columns(2)
                    with b_p1:
                        if st.button("🔮 Tambah ke Tabel RAB", use_container_width=True):
                            items_new = smart_paste_parser(wa_rab, st.session_state.get("ket_tabel_rab", ""))
                            if items_new:
                                cur = st.session_state["df_items_rab"].to_dict(orient="records")
                                cur.extend(items_new)
                                df_new = normalize_items_df(cur, st.session_state.get("ket_tabel_rab", ""))
                                st.session_state["df_items_rab"] = df_new
                                st.session_state["catatan_note_rab"] = sync_iuran_note_rab(st.session_state.get("catatan_note_rab", ""), df_new)
                                st.session_state["rab_editor_ver"] = int(st.session_state.get("rab_editor_ver", 0)) + 1
                                st.rerun()
                    with b_p2:
                        if st.button("🔄 Ganti Seluruh Tabel RAB", use_container_width=True):
                            items_new = smart_paste_parser(wa_rab, st.session_state.get("ket_tabel_rab", ""))
                            if items_new:
                                df_new = normalize_items_df(items_new, st.session_state.get("ket_tabel_rab", ""))
                                st.session_state["df_items_rab"] = df_new
                                st.session_state["catatan_note_rab"] = sync_iuran_note_rab(st.session_state.get("catatan_note_rab", ""), df_new)
                                st.session_state["rab_editor_ver"] = int(st.session_state.get("rab_editor_ver", 0)) + 1
                                st.rerun()
                with mp2:
                    qr_nama = st.text_input("Nama Kebutuhan Kegiatan RAB", key="qr_nama")
                    qr_jml = st.number_input("Jumlah (Qty)", min_value=1, value=1, step=1, key="qr_jml")
                    qr_tipe = st.selectbox(
                        "Mode Perhitungan",
                        [
                            "Normal (Qty x Harga Satuan)",
                            "Paket Iuran DWP Gabungan (5 org 300rb + 75rb/org)",
                            "Borongan (Langsung Harga Total)",
                            "Fasilitas / Gratis (Rp -)",
                        ],
                        key="qr_tipe",
                    )
                    qr_nom = st.text_input("Nominal Harga Satuan / Borongan (Rp)", value="0", key="qr_nom")
                    if st.button("➕ Tambahkan ke Tabel RAB", use_container_width=True):
                        if qr_nama.strip():
                            v_nom = parse_num(qr_nom)
                            if qr_tipe.startswith("Normal"):
                                hs, ht = v_nom, int(qr_jml) * v_nom
                            elif qr_tipe.startswith("Paket Iuran"):
                                hs, ht = 0, calc_tiered_iuran_dwp(int(qr_jml))
                            elif qr_tipe.startswith("Borongan"):
                                hs, ht = 0, v_nom
                            else:
                                hs, ht = 0, 0
                            cur = st.session_state["df_items_rab"].to_dict(orient="records")
                            cur.append({
                                "KEBUTUHAN KEGIATAN": qr_nama.strip(),
                                "JUMLAH": int(qr_jml),
                                "HARGA SATUAN": hs,
                                "HARGA TOTAL": ht,
                                "KETERANGAN": st.session_state.get("ket_tabel_rab", ""),
                            })
                            df_new = normalize_items_df(cur, st.session_state.get("ket_tabel_rab", ""))
                            st.session_state["df_items_rab"] = df_new
                            st.session_state["catatan_note_rab"] = sync_iuran_note_rab(st.session_state.get("catatan_note_rab", ""), df_new)
                            st.session_state["rab_editor_ver"] = int(st.session_state.get("rab_editor_ver", 0)) + 1
                            st.rerun()

            rab_ver = int(st.session_state.get("rab_editor_ver", 0))
            df_rab_for_editor = normalize_items_df(
                st.session_state["df_items_rab"],
                default_ket=st.session_state.get("ket_tabel_rab", ""),
                include_hapus_col=True,
            )
            edited_rab_raw = st.data_editor(
                df_rab_for_editor,
                num_rows="dynamic",
                use_container_width=True,
                hide_index=True,
                column_config={
                    "NO": st.column_config.NumberColumn("NO", width="small", disabled=True),
                    "KEBUTUHAN KEGIATAN": st.column_config.TextColumn("KEBUTUHAN KEGIATAN", width="large", required=True),
                    "JUMLAH": st.column_config.NumberColumn("JUMLAH", min_value=0, step=1, width="small"),
                    "HARGA SATUAN": st.column_config.NumberColumn("HARGA SATUAN (Rp)", min_value=0, step=1000, format="%d"),
                    "HARGA TOTAL": st.column_config.NumberColumn("HARGA TOTAL (Rp)", min_value=0, step=1000, format="%d"),
                    "KETERANGAN": st.column_config.TextColumn("KETERANGAN", width="medium"),
                    "HAPUS?": st.column_config.CheckboxColumn("🗑️ HAPUS?", width="small", default=False),
                },
                key=f"editor_table_rab_{rab_ver}",
            )

            df_rab_live, checked_rab_indices = sync_data_editor_live(
                edited_raw=edited_rab_raw,
                state_df_key="df_items_rab",
                ver_key="rab_editor_ver",
                default_ket=st.session_state.get("ket_tabel_rab", ""),
                is_rab=True,
            )
            total_rab_live = int(df_rab_live["HARGA TOTAL"].sum()) if not df_rab_live.empty else 0

            # Baris Aksi Tabel: Tambah Baris Kosong, Hapus Baris yang Dicentang, atau Hapus Nomor Baris Tertentu
            act_c1, act_c2, act_c3 = st.columns([1.0, 1.3, 1.1])
            with act_c1:
                if st.button("➕ Tambah 1 Baris Kosong di Bawah", key="btn_add_row_rab", use_container_width=True):
                    df_added = append_blank_row_to_items_df(df_rab_live, st.session_state.get("ket_tabel_rab", ""))
                    st.session_state["df_items_rab"] = df_added
                    st.session_state["rab_editor_ver"] = rab_ver + 1
                    st.rerun()
            with act_c2:
                n_checked_rab = len(checked_rab_indices)
                lbl_del_rab = f"🗑️ Hapus Baris yang Dicentang ({n_checked_rab} Baris)" if n_checked_rab > 0 else "🗑️ Hapus Baris yang Dicentang (Centang Kolom 🗑️ HAPUS?)"
                if st.button(lbl_del_rab, key="btn_del_checked_rab", use_container_width=True):
                    if n_checked_rab > 0:
                        df_after_del = delete_rows_from_items_df(df_rab_live, checked_rab_indices, st.session_state.get("ket_tabel_rab", ""))
                        st.session_state["df_items_rab"] = df_after_del
                        st.session_state["catatan_note_rab"] = sync_iuran_note_rab(st.session_state.get("catatan_note_rab", ""), df_after_del)
                        st.session_state["rab_editor_ver"] = rab_ver + 1
                        st.toast(f"{n_checked_rab} baris berhasil dihapus dari tabel RAB!", icon="🗑️")
                        st.rerun()
                    else:
                        st.warning("Silakan centang kotak pada kolom '🗑️ HAPUS?' di sebelah kanan baris yang ingin dihapus terlebih dahulu.")
            with act_c3:
                del_no_rab = st.selectbox(
                    "Atau Pilih Nomor Baris untuk Dihapus:",
                    options=["-- Pilih Baris --"] + [f"Baris #{int(r['NO'])}: {r['KEBUTUHAN KEGIATAN']}" for _, r in df_rab_live.iterrows()],
                    key="sel_del_no_rab",
                    label_visibility="collapsed",
                )
                if del_no_rab != "-- Pilih Baris --":
                    idx_to_del = int(del_no_rab.split("#")[1].split(":")[0]) - 1
                    df_after_del = delete_rows_from_items_df(df_rab_live, [idx_to_del], st.session_state.get("ket_tabel_rab", ""))
                    st.session_state["df_items_rab"] = df_after_del
                    st.session_state["catatan_note_rab"] = sync_iuran_note_rab(st.session_state.get("catatan_note_rab", ""), df_after_del)
                    st.session_state["rab_editor_ver"] = rab_ver + 1
                    st.toast(f"{del_no_rab} berhasil dihapus!", icon="🗑️")
                    st.rerun()

            mismatches = detect_location_mismatch(st.session_state.get("lokasi_acara_rab", ""), df_rab_live)
            if mismatches:
                for m in mismatches:
                    st.warning(
                        f"⚠️ Detektor Konsistensi Lokasi: Judul RAB tertulis **{m['target_city']}**, "
                        f"tetapi Baris #{m['no']} ({m['item']}) menyebut **{m['wrong_city']}**."
                    )
                if st.button(f"🛠️ Sesuaikan Otomatis Nama Kota Menjadi '{mismatches[0]['target_city']}'", use_container_width=False):
                    fixed_rows = df_rab_live.to_dict(orient="records")
                    for row_f in fixed_rows:
                        for m in mismatches:
                            row_f["KEBUTUHAN KEGIATAN"] = str(row_f["KEBUTUHAN KEGIATAN"]).replace(m["wrong_city"], m["target_city"]).replace(m["wrong_city"].lower(), m["target_city"])
                    st.session_state["df_items_rab"] = normalize_items_df(fixed_rows, st.session_state.get("ket_tabel_rab", ""))
                    st.session_state["rab_editor_ver"] = rab_ver + 1
                    st.rerun()

            st.info(
                f"📌 **Total Item:** {len(df_rab_live)} Baris  |  "
                f"💵 **TOTAL ANGGARAN RAB:** **{fmt_rp_cell(total_rab_live, False)}** "
                f"*({terbilang(total_rab_live)})*"
            )

        # LANGKAH 3: CATATAN (*Note) & TOMBOL SIMPAN / CETAK DOKUMEN
        with st.container(border=True):
            st.markdown("#### 📝 Langkah 3: Catatan (*Note) & Tombol Simpan / Cetak Dokumen")
            st.text_area(
                "Catatan Rincian di Bawah Tabel (*Note : setiap 1 baris otomatis menjadi 1 poin bullet •):",
                key="catatan_note_rab",
                height=90,
            )

            with st.expander("⚙️ Pengaturan Lanjutan RAB (ID Dokumen, Judul Baris 1 & Kolom Keterangan)", expanded=False):
                cfg_r1, cfg_r2, cfg_r3 = st.columns([1.2, 1.2, 1.0])
                with cfg_r1:
                    st.text_input("ID Dokumen RAB", key="id_rab")
                with cfg_r2:
                    st.text_input("Judul Baris 1", key="baris_1_judul")
                with cfg_r3:
                    st.text_input("Keterangan Tabel RAB (Opsional)", key="ket_tabel_rab")
                    st.checkbox("Gabungkan (Merge) Vertikal Kolom Keterangan RAB", key="merge_ket_rab")

            rab_payload_now = collect_rab_payload(df_rab_live)
            render_mail_merge_expander("RAB", rab_payload_now)

            ba1, ba2, ba3 = st.columns(3)
            with ba1:
                if st.button("💾 Simpan RAB ke Riwayat", key="btn_save_rab_main", use_container_width=True):
                    ok_s, msg_s = save_rab_to_storage(rab_payload_now, sheet_id_active, creds_active)
                    if ok_s:
                        st.success(msg_s)
                    else:
                        st.error(msg_s)
            with ba2:
                if st.button("🚀 Buat LPJ dari RAB Ini (1-Klik)", key="btn_conv_rab_main", use_container_width=True):
                    convert_active_rab_to_lpj(df_rab_live, sheet_id_active, creds_active)
                    st.toast("Seluruh data RAB berhasil ditarik menjadi LPJ & menyesuaikan kota kegiatan!", icon="🚀")
                    switch_menu(MENU_LPJ)
            with ba3:
                if st.button("🖨️ Preview Kertas A4 & Download Word/PDF", key="btn_prev_rab_main", use_container_width=True):
                    save_rab_to_storage(rab_payload_now, sheet_id_active, creds_active)
                    switch_menu(MENU_PREVIEW, preview_mode="RAB")

    # ==========================================================================
    # VIEW 2: 📝 2. BUAT / EDIT LPJ (PASCA-KEGIATAN)
    # ==========================================================================
    elif active_menu == MENU_LPJ:
        render_history_quick_bar("lpj", "LPJ", bundle, sheet_id_active, creds_active)

        # LANGKAH 1: IDENTITAS SURAT & PELAKSANAAN KEGIATAN LPJ (LABEL BERSIH TANPA <<TAG>>)
        with st.container(border=True):
            hdr_l1, hdr_l2 = st.columns([2.0, 1.1])
            with hdr_l1:
                st.markdown("#### 📌 Langkah 1: Identitas Surat & Pelaksanaan Kegiatan LPJ")
            with hdr_l2:
                if st.button("🚀 Tarik Data dari RAB yang Sedang Aktif", use_container_width=True):
                    convert_active_rab_to_lpj(df_rab_now, sheet_id_active, creds_active)
                    st.toast("Data RAB berhasil ditarik ke formulir LPJ!", icon="🚀")
                    st.rerun()

            d1, d2 = st.columns([1.0, 1.8])
            with d1:
                st.date_input("Tanggal Surat LPJ", key="tgl_laporan")
            with d2:
                st.text_input("Nama Kegiatan", key="nama_kegiatan")

            d3, d4 = st.columns(2)
            with d3:
                st.text_input("Tema Kegiatan (Opsional — Kosongkan jika tidak ada tema khusus)", key="tema_kegiatan")
            with d4:
                st.text_input("Konteks Acara / HUT (Opsional — Kosongkan jika pertemuan rutin biasa)", key="konteks_hut")

            d5, d6, d7 = st.columns([1.0, 0.9, 1.3])
            with d5:
                st.date_input("Tanggal Pelaksanaan", key="tgl_acara")
            with d6:
                st.text_input("Waktu", key="waktu_acara")
            with d7:
                st.text_input("Tempat Pelaksanaan", key="tempat_acara")

            st.text_area(
                "Rangkaian Kegiatan (1 Baris = 1 Nomor Urut Otomatis):",
                key="rangkaian_kegiatan",
                height=95,
            )

        # LANGKAH 2: TABEL RINCIAN REALISASI DANA LPJ (DENGAN KOLOM HAPUS? & TOMBOL HAPUS BARIS)
        with st.container(border=True):
            st.markdown("#### 💰 Langkah 2: Tabel Realisasi Dana LPJ (Rp Rata Kiri & Bisa Hapus Baris)")
            ks1, ks2, ks3 = st.columns([1.3, 1.2, 0.8])
            with ks1:
                st.text_input("Sumber Dana:", key="sumber_dana")
            with ks2:
                st.text_input("Keterangan Kolom Tabel:", key="ket_tabel_dana")
            with ks3:
                st.write("")
                st.checkbox("Merge Vertikal Keterangan", key="merge_ket_vertikal")

            with st.expander("⚡ Copy-Paste Daftar Belanja LPJ dari WhatsApp", expanded=False):
                wa_lpj = st.text_area(
                    "Tempel rincian belanja LPJ dari WhatsApp:",
                    height=90,
                    placeholder="Iuran DWP Gabungan di Blitar 10" + NL + "Sarapan Pagi (Roti) 20 x 13.000",
                    key="wa_paste_lpj",
                )
                bl1, bl2 = st.columns(2)
                with bl1:
                    if st.button("🔮 Tambah ke Tabel LPJ", use_container_width=True):
                        new_l = smart_paste_parser(wa_lpj, st.session_state.get("ket_tabel_dana", ""))
                        if new_l:
                            cur = st.session_state["df_items_dana"].to_dict(orient="records")
                            cur.extend(new_l)
                            st.session_state["df_items_dana"] = normalize_items_df(cur, st.session_state.get("ket_tabel_dana", ""))
                            st.session_state["lpj_editor_ver"] = int(st.session_state.get("lpj_editor_ver", 0)) + 1
                            st.rerun()
                with bl2:
                    if st.button("🔄 Ganti Seluruh Tabel LPJ", use_container_width=True):
                        new_l = smart_paste_parser(wa_lpj, st.session_state.get("ket_tabel_dana", ""))
                        if new_l:
                            st.session_state["df_items_dana"] = normalize_items_df(new_l, st.session_state.get("ket_tabel_dana", ""))
                            st.session_state["lpj_editor_ver"] = int(st.session_state.get("lpj_editor_ver", 0)) + 1
                            st.rerun()

            lpj_ver = int(st.session_state.get("lpj_editor_ver", 0))
            df_lpj_for_editor = normalize_items_df(
                st.session_state["df_items_dana"],
                default_ket=st.session_state.get("ket_tabel_dana", ""),
                include_hapus_col=True,
            )
            edited_lpj_raw = st.data_editor(
                df_lpj_for_editor,
                num_rows="dynamic",
                use_container_width=True,
                hide_index=True,
                column_config={
                    "NO": st.column_config.NumberColumn("NO", width="small", disabled=True),
                    "KEBUTUHAN KEGIATAN": st.column_config.TextColumn("KEBUTUHAN KEGIATAN", width="large", required=True),
                    "JUMLAH": st.column_config.NumberColumn("JUMLAH", min_value=0, step=1, width="small"),
                    "HARGA SATUAN": st.column_config.NumberColumn("HARGA SATUAN (Rp)", min_value=0, step=1000, format="%d"),
                    "HARGA TOTAL": st.column_config.NumberColumn("HARGA TOTAL (Rp)", min_value=0, step=1000, format="%d"),
                    "KETERANGAN": st.column_config.TextColumn("KETERANGAN", width="medium"),
                    "HAPUS?": st.column_config.CheckboxColumn("🗑️ HAPUS?", width="small", default=False),
                },
                key=f"editor_table_lpj_{lpj_ver}",
            )

            df_lpj_live, checked_lpj_indices = sync_data_editor_live(
                edited_raw=edited_lpj_raw,
                state_df_key="df_items_dana",
                ver_key="lpj_editor_ver",
                default_ket=st.session_state.get("ket_tabel_dana", ""),
                is_rab=False,
            )
            total_lpj_live = int(df_lpj_live["HARGA TOTAL"].sum()) if not df_lpj_live.empty else 0

            lact_1, lact_2, lact_3 = st.columns([1.0, 1.3, 1.1])
            with lact_1:
                if st.button("➕ Tambah 1 Baris Kosong di Bawah", key="btn_add_row_lpj", use_container_width=True):
                    df_added = append_blank_row_to_items_df(df_lpj_live, st.session_state.get("ket_tabel_dana", ""))
                    st.session_state["df_items_dana"] = df_added
                    st.session_state["lpj_editor_ver"] = lpj_ver + 1
                    st.rerun()
            with lact_2:
                n_checked_lpj = len(checked_lpj_indices)
                lbl_del_lpj = f"🗑️ Hapus Baris yang Dicentang ({n_checked_lpj} Baris)" if n_checked_lpj > 0 else "🗑️ Hapus Baris yang Dicentang (Centang Kolom 🗑️ HAPUS?)"
                if st.button(lbl_del_lpj, key="btn_del_checked_lpj", use_container_width=True):
                    if n_checked_lpj > 0:
                        df_after_del = delete_rows_from_items_df(df_lpj_live, checked_lpj_indices, st.session_state.get("ket_tabel_dana", ""))
                        st.session_state["df_items_dana"] = df_after_del
                        st.session_state["lpj_editor_ver"] = lpj_ver + 1
                        st.toast(f"{n_checked_lpj} baris berhasil dihapus dari tabel LPJ!", icon="🗑️")
                        st.rerun()
                    else:
                        st.warning("Silakan centang kotak pada kolom '🗑️ HAPUS?' di sebelah kanan baris yang ingin dihapus terlebih dahulu.")
            with lact_3:
                del_no_lpj = st.selectbox(
                    "Atau Pilih Nomor Baris LPJ untuk Dihapus:",
                    options=["-- Pilih Baris --"] + [f"Baris #{int(r['NO'])}: {r['KEBUTUHAN KEGIATAN']}" for _, r in df_lpj_live.iterrows()],
                    key="sel_del_no_lpj",
                    label_visibility="collapsed",
                )
                if del_no_lpj != "-- Pilih Baris --":
                    idx_to_del = int(del_no_lpj.split("#")[1].split(":")[0]) - 1
                    df_after_del = delete_rows_from_items_df(df_lpj_live, [idx_to_del], st.session_state.get("ket_tabel_dana", ""))
                    st.session_state["df_items_dana"] = df_after_del
                    st.session_state["lpj_editor_ver"] = lpj_ver + 1
                    st.toast(f"{del_no_lpj} berhasil dihapus!", icon="🗑️")
                    st.rerun()

            st.info(
                f"📌 **Total Item:** {len(df_lpj_live)} Baris  |  "
                f"💵 **TOTAL REALISASI LPJ:** **{fmt_rp_cell(total_lpj_live, False)}** "
                f"*({terbilang(total_lpj_live)})*"
            )

        # LANGKAH 3: PENUTUP, PENANDATANGAN & SIMPAN / CETAK LPJ
        with st.container(border=True):
            st.markdown("#### 🖋️ Langkah 3: Penutup, Penandatangan & Cetak LPJ")
            cp1, cp2 = st.columns([1.25, 1.05])
            with cp1:
                st.checkbox("Tampilkan Paragraf Ucapan Terima Kasih Donatur", key="tampilkan_ucapan_donatur")
                if st.session_state.get("tampilkan_ucapan_donatur", True):
                    st.text_area("Redaksi Ucapan Terima Kasih Donatur:", key="teks_ucapan_donatur", height=85)
                st.text_input("Catatan Dokumentasi Instagram:", key="catatan_ig")
            with cp2:
                st.text_input("Jabatan Baris 1 (Paragraf 1):", key="jabatan_baris_1")
                st.text_input("Jabatan Baris 2 (Paragraf 2):", key="jabatan_baris_2")
                st.text_input("Nama Penandatangan:", key="nama_penandatangan")
                st.checkbox("Sertakan Gambar Tanda Tangan Asli (Ny. Rini Johan Pandores)", key="tampilkan_ttd")

            lpj_payload_now = collect_lpj_payload(df_lpj_live)
            render_mail_merge_expander("LPJ", lpj_payload_now)

            al1, al2 = st.columns(2)
            with al1:
                if st.button("💾 Simpan LPJ ke Riwayat", key="btn_save_lpj_main", use_container_width=True):
                    ok_l, msg_l = save_lpj_to_storage(lpj_payload_now, sheet_id_active, creds_active)
                    if ok_l:
                        st.success(msg_l)
                    else:
                        st.error(msg_l)
            with al2:
                if st.button("🖨️ Preview Kertas A4 & Download Word/PDF", key="btn_prev_lpj_main", use_container_width=True):
                    save_lpj_to_storage(lpj_payload_now, sheet_id_active, creds_active)
                    switch_menu(MENU_PREVIEW, preview_mode="LPJ")

    # ==========================================================================
    # VIEW 3: 🖨️ 3. PREVIEW & DOWNLOAD (WORD & PDF)
    # ==========================================================================
    elif active_menu == MENU_PREVIEW:
        render_history_quick_bar("prev", "ALL", bundle, sheet_id_active, creds_active)

        with st.container(border=True):
            p_col1, p_col2 = st.columns([1.3, 1.7])
            with p_col1:
                doc_mode_sel = st.radio(
                    "Pilih Dokumen yang Ingin Di-Preview & Diunduh:",
                    options=["RAB", "LPJ"],
                    format_func=lambda x: "📋 Dokumen RAB Resmi (Pra-Kegiatan)" if x == "RAB" else "📝 Dokumen LPJ Resmi (Pasca-Kegiatan)",
                    horizontal=True,
                    key="preview_doc_mode",
                )
            with p_col2:
                up_sig_custom = st.file_uploader("Upload Tanda Tangan Baru PNG/JPG (Opsional, khusus LPJ)", type=["png", "jpg", "jpeg"], key="up_sig_preview")
                up_fotos = st.file_uploader("Upload Lampiran Foto/Kwitansi Grid 2x2 (Opsional, khusus LPJ)", type=["png", "jpg", "jpeg"], accept_multiple_files=True, key="up_fotos_preview")

            custom_sig_bytes = up_sig_custom.getvalue() if up_sig_custom is not None else None
            attachment_imgs = []
            if up_fotos:
                for idx_f, f_obj in enumerate(up_fotos, start=1):
                    attachment_imgs.append({"bytes": f_obj.getvalue(), "caption": f"Dokumentasi {idx_f}: {f_obj.name}"})

            if doc_mode_sel == "RAB":
                active_payload = collect_rab_payload(df_rab_now)
                docx_bytes = generate_rab_docx_bytes(active_payload)
                clean_name = "".join(ch if ch.isalnum() else "_" for ch in active_payload.get("lokasi_acara_rab", "JEMBER")[:25]).strip("_")
                file_base = f"RAB_DWP_{clean_name or 'MALANG'}_{date.today().strftime('%Y%m%d')}"
            else:
                active_payload = collect_lpj_payload(df_lpj_now)
                docx_bytes = generate_lpj_docx_bytes(active_payload, custom_sig_bytes=custom_sig_bytes, attachment_images=attachment_imgs)
                clean_name = "".join(ch if ch.isalnum() else "_" for ch in active_payload.get("tempat_acara", "BLITAR")[:25]).strip("_")
                file_base = f"LPJ_DWP_{clean_name or 'MALANG'}_{date.today().strftime('%Y%m%d')}"

            pdf_bytes, pdf_msg = convert_docx_bytes_to_pdf_bytes(docx_bytes)

            render_mail_merge_expander(doc_mode_sel, active_payload)

            dl1, dl2, dl3 = st.columns([1.1, 1.15, 1.15])
            with dl1:
                if st.button(f"💾 Simpan {doc_mode_sel} Ini ke Riwayat", use_container_width=True):
                    if doc_mode_sel == "RAB":
                        ok_s, msg_s = save_rab_to_storage(active_payload, sheet_id_active, creds_active)
                    else:
                        ok_s, msg_s = save_lpj_to_storage(active_payload, sheet_id_active, creds_active)
                    if ok_s:
                        st.success(msg_s)
                    else:
                        st.error(msg_s)

            with dl2:
                st.markdown(
                    make_direct_download_html(
                        file_bytes=docx_bytes,
                        filename=f"{file_base}.docx",
                        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        label=f"📄 Unduh Word {doc_mode_sel} Resmi (.docx)",
                        bg_gradient="linear-gradient(135deg, #2B98F0 0%, #1D7FE0 100%)",
                    ),
                    unsafe_allow_html=True,
                )

            with dl3:
                if pdf_bytes:
                    st.markdown(
                        make_direct_download_html(
                            file_bytes=pdf_bytes,
                            filename=f"{file_base}.pdf",
                            mime_type="application/pdf",
                            label=f"📕 Unduh PDF {doc_mode_sel} Siap Sebar (.pdf)",
                            bg_gradient="linear-gradient(135deg, #00D09C 0%, #00B386 100%)",
                        ),
                        unsafe_allow_html=True,
                    )
                else:
                    st.info(pdf_msg)

        a4_html = build_a4_live_preview_html(doc_mode_sel, active_payload, custom_sig_bytes=custom_sig_bytes)
        st.components.v1.html(a4_html, height=1080, scrolling=True)

    # ==========================================================================
    # VIEW 4: 🗄️ 4. RIWAYAT & HAPUS DATA (FULL CRUD & PEMBERSIH DOBEL OTOMATIS)
    # ==========================================================================
    elif active_menu == MENU_CLOUD:
        render_history_quick_bar("cloud", "ALL", bundle, sheet_id_active, creds_active)

        rab_rows = bundle.get("rab_rows", [])
        lpj_rows = bundle.get("lpj_rows", [])
        rk_rows = bundle.get("rekap_rows", [])

        with st.container(border=True):
            top_h1, top_h2 = st.columns([2.1, 1.1])
            with top_h1:
                st.markdown("#### 🗄️ Manajemen Riwayat Dokumen & Pembersih Data Dobel")
                st.caption("Kelola seluruh arsip RAB & LPJ yang tersimpan di penyimpanan lokal maupun Google Sheets.")
            with top_h2:
                if st.button("🧹 Bersihkan Data Dobel Otomatis", use_container_width=True):
                    ok_dedup, msg_dedup = deduplicate_history_storage(sheet_id_active, creds_active)
                    if ok_dedup:
                        st.toast(msg_dedup, icon="🧹")
                        st.success(msg_dedup)
                        st.rerun()

            db_c1, db_c2 = st.columns(2)
            with db_c1:
                st.markdown(f"##### 📋 Daftar Riwayat RAB ({len(rab_rows)} Dokumen)")
                if rab_rows:
                    rab_del_opts = [f"{r.get('ID_RAB')} | {r.get('LOKASI_ACARA')} ({fmt_rp_cell(r.get('TOTAL_RAB', 0), False)})" for r in rab_rows]
                    sel_del_rab = st.selectbox("Pilih RAB untuk Dibuka / Dihapus:", options=rab_del_opts, key="sel_manage_rab")
                    sel_rab_id = sel_del_rab.split(" | ")[0].strip()
                    mb_r1, mb_r2 = st.columns(2)
                    with mb_r1:
                        if st.button("📂 Buka RAB Ini", key="btn_m_open_rab", use_container_width=True):
                            for r in rab_rows:
                                if r.get("ID_RAB") == sel_rab_id and r.get("PAYLOAD_JSON"):
                                    load_rab_into_state(json.loads(r["PAYLOAD_JSON"]))
                                    switch_menu(MENU_RAB)
                    with mb_r2:
                        if st.button("🗑️ Hapus RAB Ini", key="btn_m_del_rab", use_container_width=True):
                            ok_d, msg_d = delete_rab_from_storage(sel_rab_id, sheet_id_active, creds_active)
                            st.toast(msg_d, icon="🗑️")
                            st.rerun()
                    st.dataframe(pd.DataFrame(rab_rows).drop(columns=["PAYLOAD_JSON"], errors="ignore"), use_container_width=True, hide_index=True)
                else:
                    st.info("Belum ada arsip RAB tersimpan.")

            with db_c2:
                st.markdown(f"##### 📝 Daftar Riwayat LPJ ({len(lpj_rows)} Dokumen)")
                if lpj_rows:
                    lpj_del_opts = [f"{r.get('ID_LPJ')} | {r.get('TEMPAT_ACARA')} ({fmt_rp_cell(r.get('TOTAL_DANA', 0), False)})" for r in lpj_rows]
                    sel_del_lpj = st.selectbox("Pilih LPJ untuk Dibuka / Dihapus:", options=lpj_del_opts, key="sel_manage_lpj")
                    sel_lpj_id = sel_del_lpj.split(" | ")[0].strip()
                    mb_l1, mb_l2 = st.columns(2)
                    with mb_l1:
                        if st.button("📂 Buka LPJ Ini", key="btn_m_open_lpj", use_container_width=True):
                            for r in lpj_rows:
                                if r.get("ID_LPJ") == sel_lpj_id and r.get("PAYLOAD_JSON"):
                                    load_lpj_into_state(json.loads(r["PAYLOAD_JSON"]))
                                    switch_menu(MENU_LPJ)
                    with mb_l2:
                        if st.button("🗑️ Hapus LPJ Ini", key="btn_m_del_lpj", use_container_width=True):
                            ok_d, msg_d = delete_lpj_from_storage(sel_lpj_id, sheet_id_active, creds_active)
                            st.toast(msg_d, icon="🗑️")
                            st.rerun()
                    st.dataframe(pd.DataFrame(lpj_rows).drop(columns=["PAYLOAD_JSON"], errors="ignore"), use_container_width=True, hide_index=True)
                else:
                    st.info("Belum ada arsip LPJ tersimpan.")

        with st.container(border=True):
            st.markdown(f"##### 📊 Tabel Gabungan Rekap Item Dana ({len(rk_rows)} Baris Rincian Item RAB & LPJ)")
            if rk_rows:
                st.dataframe(pd.DataFrame(rk_rows), use_container_width=True, hide_index=True)
            else:
                st.info("Belum ada rincian item tersimpan.")

    # ==========================================================================
    # VIEW 5: 📊 5. KOMPARASI RAB VS LPJ (KHUSUS KEGIATAN BERPASANGAN)
    # ==========================================================================
    elif active_menu == MENU_KOMPARASI:
        rab_rows = bundle.get("rab_rows", [])
        lpj_rows = bundle.get("lpj_rows", [])

        with st.container(border=True):
            st.markdown("#### 📊 Komparasi Anggaran Rencana (RAB) vs Realisasi (LPJ)")
            st.caption(
                "Pilih pasangan dokumen RAB dan LPJ dari kegiatan yang sama untuk melihat perbandingan anggaran dan selisih realisasi secara akurat."
            )

            # Cari pasangan otomatis yang memiliki LPJ.ID_RAB_REF == RAB.ID_RAB
            rab_dict_map = {r.get("ID_RAB"): r for r in rab_rows if r.get("ID_RAB")}
            lpj_dict_map = {r.get("ID_LPJ"): r for r in lpj_rows if r.get("ID_LPJ")}

            paired_lpj = [r for r in lpj_rows if r.get("ID_RAB_REF") and r.get("ID_RAB_REF") in rab_dict_map]

            if not paired_lpj:
                st.info(
                    "ℹ️ **Info Relasi Dokumen:** Saat ini contoh bawaan **RAB Jember** (`RAB-2026-JEMBER`) dan **LPJ Blitar** (`LPJ-20260813-BLITAR`) "
                    "adalah **2 kegiatan berbeda kota** sehingga tidak dibandingkan secara paksa. "
                    "Untuk membuat pasangan RAB & LPJ dari kegiatan yang sama, buka menu **📋 1. Buat / Edit RAB** lalu klik tombol **🚀 Buat LPJ dari RAB Ini (1-Klik)**, "
                    "atau pilih dokumen secara manual di bawah ini."
                )

            rab_opts = [f"{r.get('ID_RAB')} | {r.get('LOKASI_ACARA')} ({fmt_rp_cell(r.get('TOTAL_RAB', 0), False)})" for r in rab_rows]
            lpj_opts = [f"{r.get('ID_LPJ')} | {r.get('TEMPAT_ACARA')} [Ref: {r.get('ID_RAB_REF') or 'Mandiri'}] ({fmt_rp_cell(r.get('TOTAL_DANA', 0), False)})" for r in lpj_rows]

            default_rab_idx = 0
            default_lpj_idx = 0
            if paired_lpj:
                target_lpj_id = paired_lpj[0].get("ID_LPJ")
                target_rab_id = paired_lpj[0].get("ID_RAB_REF")
                for i, r in enumerate(rab_rows):
                    if r.get("ID_RAB") == target_rab_id:
                        default_rab_idx = i
                        break
                for j, r in enumerate(lpj_rows):
                    if r.get("ID_LPJ") == target_lpj_id:
                        default_lpj_idx = j
                        break

            kc1, kc2 = st.columns(2)
            with kc1:
                sel_comp_rab = st.selectbox("Pilih Dokumen Rencana (RAB):", options=rab_opts if rab_opts else ["-- Kosong --"], index=default_rab_idx)
            with kc2:
                sel_comp_lpj = st.selectbox("Pilih Dokumen Realisasi (LPJ):", options=lpj_opts if lpj_opts else ["-- Kosong --"], index=default_lpj_idx)

            if rab_opts and lpj_opts:
                chosen_rab_id = sel_comp_rab.split(" | ")[0].strip()
                chosen_lpj_id = sel_comp_lpj.split(" | ")[0].strip()
                rab_obj = rab_dict_map.get(chosen_rab_id, {})
                lpj_obj = lpj_dict_map.get(chosen_lpj_id, {})

                rab_p = json.loads(rab_obj.get("PAYLOAD_JSON", "{}")) if rab_obj.get("PAYLOAD_JSON") else collect_rab_payload(df_rab_now)
                lpj_p = json.loads(lpj_obj.get("PAYLOAD_JSON", "{}")) if lpj_obj.get("PAYLOAD_JSON") else collect_lpj_payload(df_lpj_now)

                df_c_rab = normalize_items_df(rab_p.get("items_rab", []))
                df_c_lpj = normalize_items_df(lpj_p.get("items_dana", []))
                tot_c_rab = int(df_c_rab["HARGA TOTAL"].sum()) if not df_c_rab.empty else 0
                tot_c_lpj = int(df_c_lpj["HARGA TOTAL"].sum()) if not df_c_lpj.empty else 0
                diff_c = tot_c_lpj - tot_c_rab

                if lpj_obj.get("ID_RAB_REF") != chosen_rab_id:
                    st.warning(
                        f"⚠️ **Perhatian:** LPJ `{chosen_lpj_id}` ({lpj_obj.get('TEMPAT_ACARA')}) tidak dibuat dari RAB `{chosen_rab_id}` ({rab_obj.get('LOKASI_ACARA')}). "
                        "Pastikan Anda membandingkan RAB dan LPJ dari kegiatan/kota yang sama."
                    )
                else:
                    st.success(f"✅ Pasangan Kegiatan Cocok: LPJ `{chosen_lpj_id}` merupakan realisasi langsung dari RAB `{chosen_rab_id}`.")

                m1, m2, m3 = st.columns(3)
                with m1:
                    st.markdown(
                        f"""
                        <div class="fily-pill pill-emerald">
                            <div class="fily-pill-icon">📋</div>
                            <div style="text-align:right;">
                                <div style="font-size:12px;font-weight:700;">TOTAL RENCANA (RAB)</div>
                                <div style="font-size:19px;font-weight:800;">{fmt_rp_cell(tot_c_rab, False)}</div>
                                <div style="font-size:11.5px;">{chosen_rab_id}</div>
                            </div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
                with m2:
                    st.markdown(
                        f"""
                        <div class="fily-pill pill-blue">
                            <div class="fily-pill-icon">📝</div>
                            <div style="text-align:right;">
                                <div style="font-size:12px;font-weight:700;">TOTAL REALISASI (LPJ)</div>
                                <div style="font-size:19px;font-weight:800;">{fmt_rp_cell(tot_c_lpj, False)}</div>
                                <div style="font-size:11.5px;">{chosen_lpj_id}</div>
                            </div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )
                with m3:
                    diff_lbl = ("+" if diff_c > 0 else "") + fmt_ribuan(diff_c)
                    st.markdown(
                        f"""
                        <div class="fily-pill pill-rose">
                            <div class="fily-pill-icon">🎯</div>
                            <div style="text-align:right;">
                                <div style="font-size:12px;font-weight:700;">SELISIH (LPJ - RAB)</div>
                                <div style="font-size:19px;font-weight:800;">Rp {diff_lbl}</div>
                                <div style="font-size:11.5px;">Komparasi Akhir</div>
                            </div>
                        </div>
                        """,
                        unsafe_allow_html=True,
                    )

                max_len = max(len(df_c_rab), len(df_c_lpj))
                comp_rows = []
                for i in range(max_len):
                    r_rab = df_c_rab.iloc[i] if i < len(df_c_rab) else None
                    r_lpj = df_c_lpj.iloc[i] if i < len(df_c_lpj) else None
                    v_rab = int(r_rab["HARGA TOTAL"]) if r_rab is not None else 0
                    v_lpj = int(r_lpj["HARGA TOTAL"]) if r_lpj is not None else 0
                    sel_item = v_lpj - v_rab
                    comp_rows.append({
                        "NO": i + 1,
                        "KEBUTUHAN (RAB)": str(r_rab["KEBUTUHAN KEGIATAN"]) if r_rab is not None else "-",
                        "ANGGARAN RAB": fmt_rp_cell(v_rab, True) if r_rab is not None else "-",
                        "KEBUTUHAN (LPJ)": str(r_lpj["KEBUTUHAN KEGIATAN"]) if r_lpj is not None else "-",
                        "REALISASI LPJ": fmt_rp_cell(v_lpj, True) if r_lpj is not None else "-",
                        "SELISIH": ("+" if sel_item > 0 else "") + fmt_ribuan(sel_item) if (r_rab is not None or r_lpj is not None) else "-",
                    })
                st.dataframe(pd.DataFrame(comp_rows), use_container_width=True, hide_index=True)


if __name__ == "__main__":
    main()